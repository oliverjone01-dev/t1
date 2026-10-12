# Таблица «где можно снизить цену» из floor_rows.json (floor.py). Файл xlsx пишется вне репозитория.
# Запуск: python3 -I table.py <путь.xlsx>
import json,sys
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
D='/home/user/t1/analytics-mvp/data/'
rows=json.load(open('/home/user/t1/knowledge/episodes/2026-10/ozon-price-floor/floor_rows.json'))
name={}
for l in open(D+'sku_views.ndjson'):
    r=json.loads(l); name[r['sku']]=r['name']
def rnd(x,s=10): return int(round(x/s)*s)
out=[]
for r in rows:
    p0=r['p0']; z=r['p_zero']; t=r['p_10']
    head10=(1-t/p0)*100 if t else -999
    if head10<5: continue                      # снижать некуда: меньше 5% до 10% рентабельности
    step=min(10, head10)                        # первый шаг: не больше 10% и не ниже цены 10% рентабельности
    group='продаётся (от 0,5 шт/нед)' if r['wk']>=0.5 else 'медленный: сначала реклама'
    out.append([r['art'],name.get(r['sku'],'')[:70],group,round(r['wk'],2),r['u60'],rnd(p0),round(r['rent_now']*100),
                rnd(z) if z else None,rnd(t),round(head10),round(step),rnd(p0*(1-step/100)),
                round(r['s']*100,1),'да' if r['ads'] else 'нет'])
out.sort(key=lambda x:(x[2]!='продаётся (от 0,5 шт/нед)',-x[3]))
wb=Workbook(); ws=wb.active; ws.title='Можно снизить цену'
hdr=['Артикул','Товар','Группа','Заказов в неделю','Заказов за 60 дн','Наша цена сейчас, ₽','Рентабельность сейчас, %',
     'Цена при нулевой прибыли, ₽','Цена при 10% рентабельности, ₽','Запас до 10%, %','Первый шаг, %','Цена после шага, ₽',
     'Покупатель платит, % нашей цены','Реклама за клик сейчас']
ws.append(hdr)
for c in ws[1]: c.font=Font(bold=True); c.fill=PatternFill('solid',fgColor='DDE7F3'); c.alignment=Alignment(wrap_text=True,vertical='top')
for x in out: ws.append(x)
for col,w in zip('ABCDEFGHIJKLMN',[20,48,26,10,9,12,11,13,13,9,8,12,11,9]): ws.column_dimensions[col].width=w
ws.freeze_panes='C2'
n=wb.create_sheet('Как считали')
for line in ['Прибыль со штуки как в «Отчете»: К выплате (доля выплаты OZON товара по доставленным с 01.06) - ставка за заказ 5% - реклама за клик на заказ - себестоимость - наша доставка за вычетом оплаты покупателем - АДМ 30% от (К выплате - реклама) - налог 15% от платежа покупателя.',
 'Рентабельность = прибыль / (К выплате - реклама), как в колонке «Рентаб.» «Отчета».',
 'Без акции цену задаём сами; доплата OZON считается долей нашей цены и при снижении сохраняется (тест 4, 3 товара - направление).',
 'Медленным товарам без рекламы добавлена дешёвая реклама 1 500 ₽/нед на текущие заказы (доплата OZON +6,5 п.). У товаров с рекламой расход уже в текущих.',
 'В таблице - товары, у которых до 10% рентабельности не меньше 5% снижения. Первый шаг - не больше 10% и не ниже цены 10% рентабельности; через 2 недели проверить заказы.',
 'Сколько заказов даст снижение, тестами не измерено. Чтобы прибыль не упала: -5% цены требует около +16% заказов, -10% около +37% (середина по товарам).',
 'Не учтено: шкала комиссии OZON от цены, отмены и возвраты, хранение. Предварительно: модель на проверке ФЕНИКСА, сверка с «Отчетом» за сентябрь не сделана.']:
    n.append([line])
n.column_dimensions['A'].width=140
wb.save(sys.argv[1])
print('строк',len(out),'из них продаётся',sum(x[2].startswith('продаётся') for x in out))
