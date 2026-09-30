# Конвертер выгрузки листа «Copy of СС GEN - OZON» (xlsx) в формат fixtures/cogs_*_sku.csv.
# Запуск в папке с cc.xlsx: python3 cogs_sheet_to_csv.py -> new_ym.csv, new_oz.csv. Сверять с fixtures по ключу, не заменять вслепую.
import openpyxl,csv,io,sys
wb=openpyxl.load_workbook('cc.xlsx',data_only=True)
def w(rows,hdr,path):
    b=io.StringIO(); c=csv.writer(b,lineterminator='\n'); c.writerow(hdr); c.writerows(rows); open(path,'w').write(b.getvalue())
# ЯМ
ym=[];sec=None
for r in wb['ЯМ'].iter_rows(min_row=3,values_only=True):
    m,o,cost=r[0],r[1],r[2]
    if m and not o and cost is None: sec=str(m).strip(); continue
    if not o or not sec or not sec.startswith('В ПРОДАЖЕ'): continue
    if not isinstance(cost,(int,float)) or cost<=0: continue
    ym.append([str(o).strip(),str(m or '').strip().replace('*','\\*'),round(cost),sec])
w(ym,['offer','model','cost_prod','section'],'new_ym.csv')
oz=[]
for r in wb['Лист1'].iter_rows(min_row=2,values_only=True):
    sku,off,name,model,cost=r[1:6]
    if sku is None or not isinstance(cost,(int,float)): continue
    oz.append([str(int(sku)) if isinstance(sku,float) else str(sku).strip(),str(off or '').strip(),str(model or '').strip(),round(cost)])
w(oz,['sku','offer','model','cost_prod'],'new_oz.csv')
