#!/usr/bin/env python3
"""Отчёты по реализации из кабинета Маркета (united_statistics_report, по месяцу, все магазины кабинета) против 1С по магазинам, по заказам. Только чтение.
Запуск: python3 -I sverka-kabinet-realizacii.py <корень репо> <папка с отчётами кабинета> <папка 1С 19.10> <папка 1С 20.10> <выход.xlsx>
Спека (все цифры [ДАННЫЕ]):
  Маркет (файл кабинета): лист «Доставленные товары», столбец «Стоимость всех доставленных штук с НДС с учётом всех скидок» (строки доставки = 0);
      лист «Возвращенные товары», столбец «Стоимость всех возвращённых штук с НДС с учётом всех скидок». Нетто = доставлено − возвращено. Магазин по «Названия магазинов».
  1С: ДБС-магазины - отчёты комиссионера месяца (пара по сумме: равен нетто отчёта, равен доставлено отчёта, равен доставленному по выгрузке заказов за месяц сбора, ближайший);
      ФБС-магазин FBS OOO - реализации «Яндекс маркет фбс» по дате документа (кабинет по договору из выгрузки 19.10).
  Перекрёстная проверка файла кабинета с данными API дашборда (realization_orders, где есть) и с выгрузкой заказов (доставлено по месяцу сбора).
  Разница по заказам: заказы, которые есть в отчёте кабинета за месяц, но в выгрузке заказов доставлены в другом месяце, и наоборот - это и есть переходящие.
"""
import json, sys, re, glob, collections
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

root, mdir, d19, d20, out = sys.argv[1:6]
CAB = {'1023124': 'зеркала', '74986385': 'мебель'}; CONTRACT = {'54641824/26': 'зеркала', '54542918/26': 'мебель'}
NAME2CAMP = {'DBS ООО новый': '148992469', 'FBS OOO': '149001087', 'DBS OOO 15 дней': '149002232', 'DBS ООО 7 дней': '149032438', 'DBS 7 дней ООО': '149032438', 'ДБС ООО 1 день Москва и МО': '149229378', 'GEN GROUP': '66056856'}
MODEL = {'149001087': 'ФБС', '66056856': 'ФБС'}
rd = lambda f: [json.loads(l) for l in open(f'{root}/analytics-mvp/data-ym/{f}', encoding='utf-8') if l.strip()]
ymd = lambda s: s[6:10] + '-' + s[3:5]; r2 = lambda x: round(x + 0.0, 2)
orders = [r for r in rd('orders.ndjson') if not r.get('fake')]; ro = rd('realization_orders.ndjson')
ok = {(r['order'], r['sku']): r for r in orders if not r['service']}
goods = collections.defaultdict(float); dash_orders = collections.defaultdict(dict)
for r in orders:
    if not r['service'] and r['status'] == 'DELIVERED' and r.get('fin'):
        k = (r['fin'][:7], r['campaign']); goods[k] += r['p_buyer'] * r['count']; dash_orders[k][r['order']] = dash_orders[k].get(r['order'], 0.0) + r['p_buyer'] * r['count']
api = collections.defaultdict(float)
for r in ro:
    x = ok.get((r['order'], r['sku']))
    if x: api[(r['ym'], r['campaign'])] += x['p_buyer'] * r['count'] * (1 if r['role'] == 'delivered' else -1)
# --- файлы кабинета
MK = {}  # (ym, camp) -> dict
for f in sorted(glob.glob(f'{mdir}/united_statistics*.xlsx')):
    head = pd.read_excel(f, sheet_name='Сводный отчёт', header=None)
    per = re.search(r'(\d{1,2})\.(\d{2})\.(\d{2,4})[–-]', str(head.iloc[0, 0])); y = per.group(3); ym = ('20' + y if len(y) == 2 else y) + '-' + per.group(2)
    for sheet_name, qty_col, sum_col, role in (('Доставленные товары', 10, 23, 'дост'), ('Возвращенные товары', 10, 24, 'возвр')):
        v = pd.read_excel(f, sheet_name=sheet_name, header=None)
        for _, w in v.iloc[7:].iterrows():
            if pd.isna(w[1]) or str(w[0]).startswith('Итого'): continue
            c = NAME2CAMP.get(str(w[2]).strip(), '?' + str(w[2])); k = (ym, c); d = MK.setdefault(k, dict(дост=0.0, возвр=0.0, дост_шт=0, возвр_шт=0, orders={}, file=f.split('/')[-1]))
            a = float(w[sum_col]) if pd.notna(w[sum_col]) else 0.0; q = int(w[qty_col]) if pd.notna(w[qty_col]) else 0
            d[role] += a; d[role + '_шт'] += q
            if role == 'дост' and a: o = str(int(w[3])); d['orders'][o] = d['orders'].get(o, 0.0) + a
# --- 1С
kom = []
for ff in glob.glob(f'{d19}/*.xlsx'):
    v = pd.read_excel(ff, header=None)
    if 'комиссионера' not in str(v.iloc[1, 0]): continue
    for _, w in v.iloc[6:].iterrows():
        if isinstance(w[1], str) and re.match(r'\d\d\.\d\d\.\d{4}', w[1]): kom.append(dict(ym=ymd(w[1]), num=str(w[3]), amount=float(w[6]), camp=None, how=''))
used = set()
netto = {k: v['дост'] - v['возвр'] for k, v in MK.items()}; dost = {k: v['дост'] for k, v in MK.items()}
for pname, base in (('равен нетто отчёта кабинета', netto), ('равен «доставлено» отчёта кабинета', dost), ('равен отчёту по API', api), ('равен доставленному по выгрузке заказов (месяц сбора)', goods)):
    for d in sorted(kom, key=lambda d: (d['ym'], -d['amount'])):
        if d['camp']: continue
        h = [c for (m, c) in base if m == d['ym'] and (m, c) not in used and abs(base[(m, c)] - d['amount']) <= 2]
        if h: d['camp'] = h[0]; d['how'] = pname; used.add((d['ym'], h[0]))
for d in sorted(kom, key=lambda d: (d['ym'], -d['amount'])):
    if d['camp']: continue
    cands = [c for (m, c) in goods if m == d['ym'] and (m, c) not in used]
    if cands:
        c = min(cands, key=lambda c: abs(goods[(d['ym'], c)] - d['amount']))
        if abs(goods[(d['ym'], c)] - d['amount']) < 0.25 * d['amount']: d['camp'] = c; d['how'] = 'ближайший по сумме'; used.add((d['ym'], c))
new = pd.read_excel(glob.glob(f'{d20}/*.xls')[0], header=0); old = pd.read_excel(glob.glob(f'{d19}/*.xls')[0], header=0)
contract = {str(w['Номер']): CONTRACT[str(w['Контрагент'])[-11:]] for _, w in old.iterrows() if str(w['Контрагент']).startswith('ФИЗИЧЕСКОЕ ЛИЦО')}
fbs1c = collections.defaultdict(float); fbs_docs = collections.defaultdict(list)
for _, w in new.iterrows():
    k = (ymd(str(w['Дата'])), contract.get(str(w['Номер']), 'не определён')); fbs1c[k] += float(w['Сумма']); fbs_docs[k].append(str(w['Номер']) + ' ' + f'{float(w["Сумма"]):,.0f}'.replace(',', ' '))
# --- таблица по магазинам
NAMES = {v: k for k, v in NAME2CAMP.items() if k != 'DBS 7 дней ООО'}
rows = []; rows_o = []
for (ym, c), v in sorted(MK.items()):
    net = v['дост'] - v['возвр']; fbs = MODEL.get(c) == 'ФБС'
    if fbs and c == '149001087': s1c = fbs1c.get((ym, 'мебель'), 0.0); doc = 'реализации ФБС: ' + '; '.join(fbs_docs[(ym, 'мебель')])
    elif fbs: s1c = None; doc = ''
    else:
        dd = [d for d in kom if d['ym'] == ym and d['camp'] == c]; s1c = sum(d['amount'] for d in dd) if dd else None; doc = '; '.join('№ ' + d['num'] + ' (' + d['how'] + ')' for d in dd)
    g = goods.get((ym, c), 0.0); a = api.get((ym, c))
    rows.append([ym, c, NAMES.get(c, c), 'ФБС' if fbs else 'ДБС', r2(v['дост']), v['дост_шт'], r2(v['возвр']), v['возвр_шт'], r2(net), r2(a) if a is not None else None, r2(a - net) if a is not None else None, r2(g), r2(g - v['дост']),
                 r2(s1c) if s1c is not None else None, doc, r2(s1c - net) if s1c is not None else None, r2(s1c - v['дост']) if s1c is not None else None])
    # заказы: в отчёте кабинета, но не в выгрузке за этот месяц, и наоборот
    mo = v['orders']; do = dash_orders.get((ym, c), {})
    for o, amt in sorted(mo.items()):
        if o not in do:
            x = next((r for r in orders if r['order'] == o and not r['service']), None)
            rows_o.append([ym, NAMES.get(c, c), o, r2(amt), 'в отчёте кабинета, в выгрузке заказов доставлен в ' + (x['fin'][:7] if x and x.get('fin') else 'нет / статус ' + (x['status'] if x else '?')), ])
    for o, amt in sorted(do.items()):
        if o not in mo: rows_o.append([ym, NAMES.get(c, c), o, r2(-amt), 'в выгрузке заказов доставлен в этом месяце, в отчёте кабинета за месяц нет'])
wb = Workbook(); bold = Font(bold=True); hdr = PatternFill('solid', fgColor='DDE7F0'); red = Font(color='B42318', bold=True); green = Font(color='067647', bold=True)
def sheet(title, header, rows, widths, note, chk=()):
    ws = wb.create_sheet(title); ws.cell(1, 1, note).alignment = Alignment(wrap_text=True, vertical='top'); ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(header)); ws.row_dimensions[1].height = 15 * (note.count('\n') + 2)
    for j, h in enumerate(header, 1): c = ws.cell(3, j, h); c.font = bold; c.fill = hdr; c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[3].height = 75
    for i, row in enumerate(rows, 4):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v)
            if isinstance(v, float): c.number_format = '#,##0;[Red]-#,##0'
            if j in chk and isinstance(v, float): c.font = green if abs(v) <= 2 else red
    for j in range(1, len(header) + 1): ws.column_dimensions[get_column_letter(j)].width = widths[j - 1] if j - 1 < len(widths) else 14
    ws.freeze_panes = ws.cell(4, 4); return ws
wb.remove(wb['Sheet'])
ws = wb.create_sheet('Как читать'); ws.column_dimensions['A'].width = 200
files = sorted(set(v['file'] for v in MK.values())); months = sorted(set(k[0] for k in MK))
txt = [f'Отчёты по реализации из кабинета Маркета за {", ".join(months)} (кабинет мебель, все магазины) против 1С по магазинам. Все цифры [ДАННЫЕ].', '',
       'Лист «По магазинам»: по каждому магазину и месяцу - «доставлено» и «возвращено» из файла кабинета (стоимость с НДС с учётом всех скидок, строки доставки в файле нулевые), нетто, перекрёстно - то же по API дашборда (где собрано) и доставленное по выгрузке заказов за месяц сбора; 1С - документ и сумма; разница 1С − нетто и 1С − доставлено.',
       '   ДБС-магазины сравниваются с отчётами комиссионера, FBS OOO - с реализациями «Яндекс маркет фбс» за месяц (основной документ плюс корректировки, датированные этим месяцем).',
       'Лист «Переходящие заказы»: заказы, которые в отчёте кабинета стоят в этом месяце, а в выгрузке заказов дашборда доставлены в другом, и наоборот. Это объясняет разницу между документами 1С, собранными «по дате сбора за продажу», и отчётом кабинета.']
for i, t in enumerate(txt, 1): ws.cell(i, 1, t)
sheet('По магазинам', ['месяц', 'кампания', 'магазин', 'модель', 'кабинет: доставлено, ₽', 'доставлено, шт', 'кабинет: возвращено, ₽', 'возвращено, шт', 'кабинет: нетто, ₽', 'API дашборда: нетто (где собрано)', 'API − кабинет', 'выгрузка заказов: доставлено за месяц сбора, ₽', 'выгрузка − кабинет доставлено', '1С: сумма', '1С: документы', '1С − нетто кабинета', '1С − доставлено кабинета'],
      rows, [9, 11, 18, 7, 14, 9, 13, 9, 14, 15, 11, 16, 13, 14, 44, 14, 14], 'Отчёт кабинета против 1С и против данных дашборда по магазинам. Зелёным - равно до 2 ₽.', chk=(11, 16, 17))
sheet('Переходящие заказы', ['месяц отчёта', 'магазин', 'заказ', 'сумма, ₽ (минус = нет в отчёте)', 'где заказ'], rows_o, [11, 18, 14, 16, 70], 'Разница по заказам между отчётом кабинета за месяц и выгрузкой заказов дашборда за месяц сбора.')
wb.save(out)
print('ok', out, files)
for r in rows: print(r[:4], [int(x) if isinstance(x, float) else x for x in r[4:14]], r[15:17])
print('переходящих строк:', len(rows_o))
