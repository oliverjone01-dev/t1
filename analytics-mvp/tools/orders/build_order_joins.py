#!/usr/bin/env python3
# Per-order досбор для блока «Аналитика по заказам»: реклама (CPO) и доставка (ведомость) ПО НОМЕРУ
# ЗАКАЗА, чтобы сесть на orders_daily (backbone OZON) в build-katya.
#
# Выход:
#   data/cpo_orders.ndjson       - {order, sp}     реклама «за заказ» по номеру заказа (base, без -N)
#   data/delivery_orders.ndjson  - {order, ship, deliv}  наша/клиентская доставка по номеру постинга
#
# Реклама (ads) в OZON accrual/postings по заказу = 0 (CPO API не отдаёт), поэтому берём из CPO-файлов.
# Доставка (наша+клиентская) в OZON по заказу разрежена - берём из ведомости, как в таблице по артикулам.
# Запуск из analytics-mvp: python3 tools/orders/build_order_joins.py
import openpyxl, glob, os, json, re, collections

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
CPO_RAW = os.path.join(ROOT, "tools", "cpo", "raw")
DL_RAW = os.path.join(ROOT, "tools", "delivery", "raw")
OUT_CPO = os.path.join(ROOT, "data", "cpo_orders.ndjson")
OUT_DL = os.path.join(ROOT, "data", "delivery_orders.ndjson")


def num(x):
    try:
        return float(str(x).replace(",", ".").split("\n")[0].strip() or 0)
    except Exception:
        return 0.0


def norm(s):
    return re.sub(r"\s+", " ", str(s).strip()) if s else ""


def main():
    # --- CPO по номеру заказа (base) ---
    cpo = collections.defaultdict(float)
    for f in sorted(glob.glob(os.path.join(CPO_RAW, "*.xlsx"))):
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        ws = wb["Statistics"] if "Statistics" in wb.sheetnames else wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        hi = None
        for i, r in enumerate(rows[:5]):
            if any((str(c) if c else "") == "ID заказа" for c in r):
                hi = i
                break
        if hi is None:
            wb.close()
            continue
        hdr = [str(c).strip() if c else "" for c in rows[hi]]
        ci = {h: j for j, h in enumerate(hdr)}
        for r in rows[hi + 1:]:
            if r is None or all(c is None for c in r):
                continue
            no = str(r[ci["Номер заказа"]] or "").strip()
            if no:
                cpo[no] += num(r[ci["Расход, ₽"]])
        wb.close()
    with open(OUT_CPO, "w", encoding="utf-8") as w:
        for no, sp in sorted(cpo.items()):
            w.write(json.dumps({"order": no, "sp": round(sp, 2)}, ensure_ascii=False) + "\n")
    print(f"CPO заказов: {len(cpo)} | сумма рекламы: {sum(cpo.values()):,.0f} -> {OUT_CPO}")

    # --- Доставка по номеру постинга (ведомость), дедуп по уникальной отправке ---
    dl = collections.defaultdict(lambda: [0.0, 0.0])  # posting -> [ship, deliv]
    seen = set()
    for f in sorted(glob.glob(os.path.join(DL_RAW, "*.xlsx"))):
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            hdr = [norm(c) for c in rows[0]]
            ci = {h: j for j, h in enumerate(hdr)}
            if "Номер заказа" not in ci or "Стоимость отправки" not in ci:
                continue
            for r in rows[1:]:
                if r is None or all(c is None for c in r):
                    continue
                if "OZON" not in norm(r[ci["Площадка"]]).upper():
                    continue
                no = norm(r[ci["Номер заказа"]])
                if not no:
                    continue
                sp = num(r[ci["Стоимость отправки"]])
                dv = num(r[ci["Стоимость доставки"]])
                dshort = norm(r[ci["Дата отгрузки"]]) if "Дата отгрузки" in ci else ""
                key = (no, dshort, round(sp, 2))
                if key in seen:      # схлопываем повтор строк одной отправки
                    continue
                seen.add(key)
                dl[no][0] += sp
                dl[no][1] += dv
        wb.close()
    with open(OUT_DL, "w", encoding="utf-8") as w:
        for no, (sh, dv) in sorted(dl.items()):
            w.write(json.dumps({"order": no, "ship": round(sh, 2), "deliv": round(dv, 2)}, ensure_ascii=False) + "\n")
    print(f"доставка заказов: {len(dl)} | наша: {sum(v[0] for v in dl.values()):,.0f} | клиент: {sum(v[1] for v in dl.values()):,.0f} -> {OUT_DL}")


if __name__ == "__main__":
    main()
