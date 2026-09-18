#!/usr/bin/env python3
# Разнос рекламы «Оплата за заказ» (CPO) по SKU/день из ручных per-order выгрузок кабинета OZON.
#
# Зачем: OZON НЕ отдаёт атрибуцию CPO по SKU через API, поэтому в дашборде CPO целиком лежал в
# «Общих расходах». Ручной отчёт «Оплата за заказ (все товары)» (Статистика по заказам) даёт строку
# на каждый заказ с SKU, артикулом и расходом - из него собираем дневной срез по SKU, который
# build-katya.ts разносит в колонку «Реклама» и на ту же сумму уменьшает «Общие».
#
# Вход:  tools/cpo/raw/*.xlsx  - выгрузки кабинета (в git не хранятся, см. .gitignore).
#        Берём ТОЛЬКО файлы с заголовком-строкой, где есть «ID заказа» (per-order отчёт).
#        Кампанийные отчёты (по SKU, с колонкой «Инструмент») игнорируются - в них нет разбивки
#        CPO по заказу, только агрегат.
# Выход: data/cpo_sku_daily.ndjson  - {d, sku, offer, sp, n} по (sku, дню).
#
# Запуск из analytics-mvp:  python3 tools/cpo/build_cpo_sku_daily.py
# Требуется: python3 + openpyxl.
import openpyxl, glob, os, json, datetime, collections, sys

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
RAW = os.path.join(ROOT, "tools", "cpo", "raw")
OUT = os.path.join(ROOT, "data", "cpo_sku_daily.ndjson")


def num(x):
    try:
        return float(x)
    except Exception:
        return 0.0


def main():
    files = sorted(glob.glob(os.path.join(RAW, "*.xlsx")))
    if not files:
        sys.exit("нет ни одного .xlsx в " + RAW + " (положи per-order выгрузки CPO)")
    daily = collections.defaultdict(lambda: [0.0, 0])  # (sku,d)->[spend,orders]
    sku_off = {}
    permon = collections.defaultdict(float)
    used = []
    for f in files:
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        ws = wb["Statistics"] if "Statistics" in wb.sheetnames else wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        hi = None
        for i, r in enumerate(rows[:4]):
            rr = [str(c) if c is not None else "" for c in r]
            if "ID заказа" in rr:
                hi = i
                break
        if hi is None:
            wb.close()
            print("  пропуск (не per-order отчёт):", os.path.basename(f))
            continue
        hdr = [str(c) if c is not None else "" for c in rows[hi]]
        ci = {h: j for j, h in enumerate(hdr)}
        cD, cS, cA, cSp = ci["Дата"], ci["SKU"], ci["Артикул"], ci["Расход, ₽"]
        used.append(os.path.basename(f))
        for r in rows[hi + 1:]:
            if r is None or all(c is None for c in r):
                continue
            try:
                d = datetime.datetime.strptime(str(r[cD])[:10], "%d.%m.%Y").strftime("%Y-%m-%d")
            except Exception:
                continue
            sku = str(r[cS]).split(".")[0] if r[cS] is not None else ""
            if not sku or sku == "0":
                continue
            sp = num(r[cSp])
            daily[(sku, d)][0] += sp
            daily[(sku, d)][1] += 1
            if r[cA] is not None:
                sku_off.setdefault(sku, str(r[cA]))
            permon[d[:7]] += sp
        wb.close()
    out = []
    for (sku, d), (sp, n) in sorted(daily.items()):
        out.append({"d": d, "sku": sku, "offer": sku_off.get(sku, sku), "sp": round(sp, 2), "n": n})
    with open(OUT, "w", encoding="utf-8") as w:
        for o in out:
            w.write(json.dumps(o, ensure_ascii=False) + "\n")
    print("файлов:", len(used), used)
    print("CPO по месяцам:", {k: round(v) for k, v in sorted(permon.items())})
    print("строк:", len(out), "| уник SKU:", len(set(k[0] for k in daily)))
    print("записано:", OUT)


if __name__ == "__main__":
    main()
