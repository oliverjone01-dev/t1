# ФЕНИКС 14.10, итерация 2: пересборка книги скриптом автора (76d8244) на моём рендере марта-октября и сравнение
# с доставленной ячейка в ячейку (формулы и значения после recalc). python3 -I -B rebuild2.py <каталог rp> <xlsx> <recalc.py>
import shutil, subprocess, os, sys, glob
import openpyxl
src, delivered, recalc = sys.argv[1:4]
out = '/tmp/feniks-v6/rebuild2/'
os.makedirs(out, exist_ok=True)
for f in glob.glob(os.path.join(src, 'rp_2026-*.txt')):
    shutil.copy(f, out + os.path.basename(f))
p = subprocess.run(['python3', '-I', '-B', '/home/user/t1/knowledge/episodes/2026-10/ozon-testy-scripts/v6/forecast30.py', out], capture_output=True, text=True)
print('build rc', p.returncode, p.stderr[-400:])
fn = out + 'OZON-prognoz-30mln-2026-10-14.xlsx'
r = subprocess.run(['python3', recalc, fn], capture_output=True, text=True)
print('recalc', r.stdout.replace('\n', ' ')[:200])
cells = df = dv = 0; shown = 0
for mode in (False, True):
    wa, wb = openpyxl.load_workbook(delivered, data_only=mode), openpyxl.load_workbook(fn, data_only=mode)
    assert wa.sheetnames == wb.sheetnames, (wa.sheetnames, wb.sheetnames)
    for n in wa.sheetnames:
        for row in wa[n].iter_rows():
            for c in row:
                v1, v2 = c.value, wb[n][c.coordinate].value
                cells += not mode
                same = v1 == v2 or (isinstance(v1, (int, float)) and isinstance(v2, (int, float)) and abs(v1 - v2) < 1e-6)
                if not same:
                    dv += mode; df += not mode
                    if shown < 5: print('DIFF', n, c.coordinate, repr(v1)[:80], repr(v2)[:80]); shown += 1
wv = openpyxl.load_workbook(delivered, data_only=True)
errs = [(n, c.coordinate, c.value) for n in wv.sheetnames for row in wv[n].iter_rows() for c in row
        if isinstance(c.value, str) and c.value.startswith('#') and len(c.value) > 1]
print(f'листы {wa.sheetnames}; ячеек {cells}, расхождений формул {df}, значений {dv}; ошибок в доставленной {len(errs)} {errs[:3]}')
