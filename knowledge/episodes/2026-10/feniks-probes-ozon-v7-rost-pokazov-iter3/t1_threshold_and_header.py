# ФЕНИКС iter3, проба T1: порог в «Карточках» = cards7.threshold, заголовок счётчика, главный товар в «Карточках» против зелёных строк «Все кампании»
# python3 -I -B t1_threshold_and_header.py <поданный.xlsx>
import sys
from openpyxl import load_workbook
V7 = '/home/user/t1/knowledge/episodes/2026-10/ozon-testy-scripts/v7'
sys.path.insert(0, V7); sys.path.insert(0, V7 + '/../v5')
import cards7 as K, forecast as F, plan7 as P
wb = load_workbook(sys.argv[1]); ws = wb['Карточки']
hdr = [c.value for c in ws[2]]
print('заголовок J:', hdr[9])
bad = 0
for i, c in enumerate(K.cards[:10]):
    r = 3 + i; b = ws.cell(r, 2).value; j = ws.cell(r, 10).value
    ok = (b == c['main_art'] and j == K.threshold(c)); bad += not ok
    print(f'  r{r} {b:22} порог в книге {j!s:>4} threshold() {K.threshold(c):>3} {"OK" if ok else "РАСХОЖДЕНИЕ"}')
print('расхождений порога:', bad)
# главные 10 карточек: есть ли своя кампания (off = главный артикул) среди идущих
mains = {c['main']: c for c in K.cards[:10]}
own = {}
for cid in P.ACTIVE:
    off = F.camp[cid]['off'] if cid in F.camp else ''
    s = F.art2sku.get(off)
    if s in mains: own.setdefault(s, []).append(cid)
print('главные без своей кампании:', [mains[s]['main_art'] + ' (' + mains[s]['main_name'][:30] + ')' for s in mains if s not in own])
print('главные со своей кампанией:', {mains[s]['main_art']: v for s, v in own.items()})
print('главных с текущим кликом >= 20 000:', [mains[s]['main_art'] for s in mains if mains[s]['main_sp_wk'] >= 20000])
