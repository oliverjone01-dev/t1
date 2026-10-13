#!/usr/bin/env python3
"""Фикстура конвертера CPO (ФЕНИКС G6 13.10): повторная выгрузка не задваивает заказы.
1) тот же сентябрь, в одной строке Расход -> 0: 191 строка, выручка прежняя, расход меньше ровно на эту строку;
2) поверх - выгрузка 16-30.09 из тех же строк: заменяет только свои дни, итог тот же.
Запуск из analytics-mvp: python3 tools/tests/cpo_conv_fixture.py  (код 1 при провале)."""
import json, os, shutil, subprocess, sys, tempfile
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
MVP = os.path.dirname(os.path.dirname(HERE))
SRC_X = os.path.join(MVP, "data/raw/ozon-cpo/2026-09_cpo_orders.xlsx")
CONV = os.path.join(MVP, "tools/cpo/build_cpo_card_orders.py")


def run(root):
    r = subprocess.run([sys.executable, os.path.join(root, "tools/cpo/build_cpo_card_orders.py")], capture_output=True, text=True)
    if r.returncode:
        sys.exit("конвертер упал: " + r.stdout + r.stderr)
    rows = [json.loads(l) for l in open(os.path.join(root, "data/cpo_card_orders.ndjson"), encoding="utf-8")]
    return len(rows), round(sum(x["rev"] for x in rows), 2), round(sum(x["sp"] for x in rows), 2)


with tempfile.TemporaryDirectory() as t:
    for sub in ("tools/cpo", "data/raw/ozon-cpo"):
        os.makedirs(os.path.join(t, sub))
    shutil.copy(CONV, os.path.join(t, "tools/cpo/"))
    x1 = os.path.join(t, "data/raw/ozon-cpo/2026-09_cpo_orders.xlsx")
    shutil.copy(SRC_X, x1)
    base = run(t)
    wb = openpyxl.load_workbook(x1)
    ws = wb.worksheets[0]
    hdr = [c.value for c in ws[3]]
    col = hdr.index("Расход, ₽") + 1
    cut = ws.cell(row=10, column=col).value
    ws.cell(row=10, column=col).value = 0
    wb.save(os.path.join(t, "data/raw/ozon-cpo/2026-09_cpo_orders_v2.xlsx"))
    re1 = run(t)
    # выгрузка второй половины месяца: строки 16-30.09 из v2, окно в «Период:»
    wb2 = openpyxl.load_workbook(os.path.join(t, "data/raw/ozon-cpo/2026-09_cpo_orders_v2.xlsx"))
    w2 = wb2.worksheets[0]
    w2.cell(row=1, column=1).value = "Период: 16.09.2026 - 30.09.2026"
    for i in range(w2.max_row, 3, -1):
        d = str(w2.cell(row=i, column=1).value or "")
        if d and int(d[:2]) < 16:
            w2.delete_rows(i)
    wb2.save(os.path.join(t, "data/raw/ozon-cpo/2026-09_cpo_orders_v3.xlsx"))
    re2 = run(t)

exp1 = (base[0], base[1], round(base[2] - cut, 2))
print("одна выгрузка", base, "| повтор с правкой", re1, "ожидание", exp1, "| + половина месяца", re2)
ok = base[:2] == (191, 15680512.0) and re1 == exp1 and re2 == exp1
print("OK" if ok else "FAIL")
sys.exit(0 if ok else 1)
