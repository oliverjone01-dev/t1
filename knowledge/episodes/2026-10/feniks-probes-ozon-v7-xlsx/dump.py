import sys
from openpyxl import load_workbook
wb = load_workbook(sys.argv[1])
for ws in wb.worksheets:
    print('=====', ws.title, ws.max_row, 'x', ws.max_column)
    for row in ws.iter_rows(values_only=True):
        if any(v is not None and v != '' for v in row):
            print(' | '.join('' if v is None else str(v) for v in row))
