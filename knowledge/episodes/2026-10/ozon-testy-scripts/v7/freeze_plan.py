# v11: заморозить план 15.10 из Excel, который получил Иван, в JSON для вкладки «План с 15.10: как идём к цели».
# Вход - собранный по рецепту handoff v11 §5 файл OZON-plan-2026-10-15.xlsx (вне репо). Выход - JSON без прибыли по артикулам
# (репо публичный): карточки, главные, пороги, бюджеты, кампании на минимум, 26 артикулов с датой и целевой ценой, цели месяца.
# python3 -I freeze_plan.py <plan.xlsx> <out.json>
import sys, json, re
from openpyxl import load_workbook
SRC, OUT = sys.argv[1], sys.argv[2]
wb = load_workbook(SRC, read_only=True)
rows = lambda n: [r for r in wb[n].iter_rows(values_only=True)]
# короткие имена карточек - как в handoff и в переписке с Иваном
SHORT = {'GGT-03-3-3-E-16080': 'TRUBIS 160', 'GGT-35-5-1-100-180': 'KONUX', 'GGM-18-2-2': 'KUVINO', 'GGM-02-1-2': 'EVELIX',
         'GGT-47-2-2-80': 'GENTRIS', 'GGL-09-2': 'OSOLIS', 'GGT-03-3-3-L-120': 'TRUBIS L120', 'GGM-20-2': 'OZEVIS',
         'GGK-02-2-2-80': 'ARFEO', 'GGM-24-1': 'ARKIS'}
# порог с поправкой на снижение цены (handoff v11 §3 п.6) - бледная линия; у GENTRIS, TRUBIS L120, OZEVIS цена не меняется в окне
THR_PRICE = {'GGT-03-3-3-E-16080': 35, 'GGM-02-1-2': 17, 'GGT-35-5-1-100-180': 26, 'GGM-18-2-2': 22, 'GGL-09-2': 14,
             'GGK-02-2-2-80': 10, 'GGM-24-1': 9}
cards = []
for r in rows('Карточки')[2:]:
    if not r[0] or not str(r[0]).startswith('волна 1'): continue
    cards.append(dict(main_art=r[1], short=SHORT[r[1]], name=r[2], u_wk_base=r[3], v_wk_base=r[4], click_now=r[5], click_new=r[6],
                      n_min=r[7], threshold=r[9], threshold_price=THR_PRICE.get(r[1])))
assert len(cards) == 10, len(cards)
mins = {}
for r in rows('Минимальный клик')[2:]:
    if r[0] and r[4] == '8 ₽ за клик': mins.setdefault(r[0], []).append(dict(art=r[1], budget=r[5]))
for c in cards:
    c['min_arts'] = [m['art'] for m in mins.get(c['main_art'], [])]
    assert len(c['min_arts']) == c['n_min'], (c['main_art'], len(c['min_arts']), c['n_min'])
camps = []
for r in rows('Все кампании')[3:]:
    if not r[0] or not r[1]: continue
    camps.append(dict(action=r[0], cid=str(r[1]), art=r[2], tempo=r[4], budget=r[9]))
to_min = [c for c in camps if str(c['action']).startswith('на минимум')]
assert len(to_min) == 11, len(to_min)
price = []
for r in rows('Цена вниз')[2:]:
    if not r[1] or not re.match(r'^(GG|B-)', str(r[1])): continue
    price.append(dict(art=r[1], name=r[2], when=str(r[0])[:5], when_note=str(r[0]), price_now=r[4], step=r[5], price_new=r[6],
                      site_now=r[7], site_new=r[8]))
assert len(price) == 26, len(price)
# цели месяца: ячейки вида «355 (300 ... 429)»
def tri(s):
    m = re.match(r'([\d ,]+) \(([\d ,]+) \.\.\. ([\d ,]+)\)', s)
    f = lambda x: float(x.replace(' ', '').replace(',', '.'))
    return dict(mid=f(m[1]), lo=f(m[2]), hi=f(m[3]))
goal = {}
for r in rows('Цель'):
    if not r[0]: continue
    for key, lbl in (('units', 'Заказов, шт (без'), ('turn_mln', 'Оборот, млн'), ('pay_mln', 'К выплате, млн ₽ (база')):
        if str(r[0]).startswith(lbl) and r[1] and '...' in str(r[1]): goal[key] = dict(oct=tri(r[1]), nov=tri(r[2]))
    if str(r[0]).startswith('Ноябрь средне с учётом откатов'):
        u, p = str(r[2] or r[1]).split(' / '); goal['nov_rollback'] = dict(units=float(u), pay_mln=float(p.replace(',', '.')))
assert set(goal) == {'units', 'turn_mln', 'pay_mln', 'nov_rollback'}, goal.keys()
out = dict(source='OZON-plan-2026-10-15.xlsx (отправлен Ивану 15.10), freeze_plan.py', start='2026-10-15', check='2026-10-29',
           window=['2026-10-15', '2026-10-28'], main_budget=20000, min_budget_sku=1500, min_campaign_budget=2000,
           effect_from='2026-10-17', base_month_units=293, cards=cards, to_min=to_min, price=price, goal=goal)
json.dump(out, open(OUT, 'w'), ensure_ascii=False, indent=1)
print('карточек', len(cards), 'мин.клик', sum(len(c['min_arts']) for c in cards), 'кампаний на минимум', len(to_min),
      'цена', len(price), '(15.10:', sum(p['when'] == '15.10' for p in price), ')', 'цели', goal['units']['oct'])
