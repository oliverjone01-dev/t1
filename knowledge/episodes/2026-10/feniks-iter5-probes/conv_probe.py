#!/usr/bin/env python3
"""FENIKS iter5 G6: adversarial fixtures for build_cpo_card_orders.py @ 814a861d, each case in a fresh sandbox dir."""
import os, subprocess, json, shutil, sys
import openpyxl

S = "/tmp/claude-0/-home-user-t1/14242375-b0a2-5b30-b172-6a38ff5cb418/scratchpad/conv"
REPO = "/home/user/t1"
REV = "814a861d"
CONV = subprocess.run(["git", "-C", REPO, "show", f"{REV}:analytics-mvp/tools/cpo/build_cpo_card_orders.py"], capture_output=True).stdout
XL = subprocess.run(["git", "-C", REPO, "show", f"{REV}:analytics-mvp/data/raw/ozon-cpo/2026-09_cpo_orders.xlsx"], capture_output=True).stdout


def box(name):
    d = os.path.join(S, name, "analytics-mvp")
    os.makedirs(os.path.join(d, "tools/cpo"), exist_ok=True)
    os.makedirs(os.path.join(d, "data/raw/ozon-cpo"), exist_ok=True)
    open(os.path.join(d, "tools/cpo/build_cpo_card_orders.py"), "wb").write(CONV)
    return d


def put(d, fname, mut=None):
    p = os.path.join(d, "data/raw/ozon-cpo", fname)
    open(p, "wb").write(XL)
    if mut:
        wb = openpyxl.load_workbook(p)
        mut(wb.worksheets[0])
        wb.save(p)
    return p


def run(d):
    r = subprocess.run([sys.executable, os.path.join(d, "tools/cpo/build_cpo_card_orders.py")], capture_output=True, text=True)
    out = os.path.join(d, "data/cpo_card_orders.ndjson")
    rows = [json.loads(l) for l in open(out, encoding="utf-8")] if os.path.exists(out) and not r.returncode else []
    w = json.load(open(os.path.join(d, "data/cpo_card_windows.json"))) if os.path.exists(os.path.join(d, "data/cpo_card_windows.json")) and not r.returncode else None
    return {"rc": r.returncode, "rows": len(rows), "rev": round(sum(x["rev"] for x in rows), 2), "sp": round(sum(x["sp"] for x in rows), 2),
            "n": sum(x["n"] for x in rows), "win": w, "msg": (r.stdout + r.stderr).strip().splitlines()[-1][:160] if (r.stdout + r.stderr).strip() else ""}


def hdr(ws):
    for i in range(1, 7):
        vals = [c.value for c in ws[i]]
        if "Номер заказа" in [str(v or "") for v in vals]:
            return i, {str(v): j + 1 for j, v in enumerate(vals) if v}
    raise SystemExit("no header")


def set_cell(col, row_off, val):
    def m(ws):
        hi, H = hdr(ws)
        ws.cell(row=hi + row_off, column=H[col]).value = val
    return m


res = {}
d = box("c0"); put(d, "2026-09_cpo_orders.xlsx"); res["c0 base"] = run(d)
hi_info = {}
wb = openpyxl.load_workbook(os.path.join(d, "data/raw/ozon-cpo/2026-09_cpo_orders.xlsx")); hi, H = hdr(wb.worksheets[0])
cut = wb.worksheets[0].cell(row=hi + 7, column=H["Расход, ₽"]).value
res["c0 base"]["cut_row7_sp"] = cut
# C1 re-export same window, one cell changed, sorts AFTER original
d = box("c1"); put(d, "2026-09_cpo_orders.xlsx"); put(d, "2026-09_cpo_orders_v2.xlsx", set_cell("Расход, ₽", 7, 0)); res["c1 reexport after"] = run(d)
# C2 the same, but the fresh re-export is named so it sorts BEFORE the old file -> old file wins
d = box("c2"); put(d, "2026-09_cpo_orders.xlsx"); put(d, "2026-09-reexport.xlsx", set_cell("Расход, ₽", 7, 0)); res["c2 reexport sorts before"] = run(d)
# C3 a row outside the declared window
d = box("c3"); put(d, "2026-09_cpo_orders.xlsx", set_cell("Дата", 5, "05.10.2026")); res["c3 row outside window"] = run(d)
# C4 Стоимость != price x qty
def bad_line(ws):
    hi, H = hdr(ws)
    c = ws.cell(row=hi + 3, column=H["Стоимость, ₽"]); c.value = float(c.value) + 500
d = box("c4"); put(d, "2026-09_cpo_orders.xlsx", bad_line); res["c4 line != price*qty"] = run(d)
# C5 no «Период:» line
def no_per(ws):
    for i in range(1, 4):
        if "Период" in str(ws.cell(row=i, column=1).value or ""):
            ws.cell(row=i, column=1).value = "Отчёт по заказам"
d = box("c5"); put(d, "2026-09_cpo_orders.xlsx", no_per); res["c5 no period line"] = run(d)
# C6 quantity changed but «Стоимость продажи» and «Стоимость» consistent (price x qty) -> passes; units change visible
def qty2(ws):
    hi, H = hdr(ws)
    r = hi + 2
    p = float(ws.cell(row=r, column=H["Стоимость продажи, ₽"]).value); ws.cell(row=r, column=H["Количество"]).value = 2; ws.cell(row=r, column=H["Стоимость, ₽"]).value = p * 2
d = box("c6"); put(d, "2026-09_cpo_orders.xlsx", qty2); res["c6 consistent qty edit"] = run(d)
# C7 exact duplicate row inside one file (cabinet repeats a line)
def dup(ws):
    hi, H = hdr(ws)
    vals = [c.value for c in ws[hi + 1]]
    ws.insert_rows(hi + 2)
    for j, v in enumerate(vals):
        ws.cell(row=hi + 2, column=j + 1).value = v
d = box("c7"); put(d, "2026-09_cpo_orders.xlsx", dup); res["c7 duplicate row in one file"] = run(d)
# C8 second, adjacent window file (October, one row re-dated into October, period line Oct)
def octo(ws):
    for i in range(1, 4):
        if "Период" in str(ws.cell(row=i, column=1).value or ""):
            ws.cell(row=i, column=1).value = "Период: 01.10.2026 - 12.10.2026"
    hi, H = hdr(ws)
    last = ws.max_row
    for i in range(last, hi + 1, -1):
        ws.delete_rows(i)
    ws.cell(row=hi + 1, column=H["Дата"]).value = "03.10.2026"
d = box("c8"); put(d, "2026-09_cpo_orders.xlsx"); put(d, "2026-10_cpo_orders.xlsx", octo); res["c8 adjacent october window"] = run(d)
for k, v in res.items():
    print(k, json.dumps(v, ensure_ascii=False))
