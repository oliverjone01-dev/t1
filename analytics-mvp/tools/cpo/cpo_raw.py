#!/usr/bin/env python3
# Чтение ручных per-order выгрузок «Оплата за заказ» (tools/cpo/raw/*.xlsx) - один разбор на оба сборщика:
# tools/cpo/build_cpo_sku_daily.py (по SKU и дню) и tools/orders/build_order_joins.py (по номеру заказа).
#
# Правило (ФЕНИКС O2 13.10, Иван 14.10 «делай по своим рекомендациям»): каждый файл ЗАМЕНЯЕТ свой период из
# строки «Период: дд.мм.гггг - дд.мм.гггг». Файлы читаются по имени, при пересечении периодов побеждает поздний; второе
# скачивание браузера «x (1).xlsx» идёт после «x.xlsx» (ФЕНИКС G4 14.10: раньше сортировалось раньше и проигрывало).
# Раньше строки всех файлов складывались, и повтор выгрузки за месяц удваивал списания (сен 818 149 -> 1 636 298).
# Отказ (CpoRawError, сборщик выходит с кодом 1): в per-order файле нет строки «Период:», строка с датой вне
# периода своего файла, строка с номером заказа и без даты.
import datetime
import glob
import os
import re


class CpoRawError(Exception):
    pass


def _num(x):
    try:
        return float(str(x).replace(",", ".").strip() or 0)
    except Exception:
        return 0.0


def _date(x):
    if isinstance(x, (datetime.datetime, datetime.date)):
        return x.strftime("%Y-%m-%d")
    try:
        return datetime.datetime.strptime(str(x).strip()[:10], "%d.%m.%Y").strftime("%Y-%m-%d")
    except Exception:
        return None


def read_file(path):
    """Один файл -> {file, win, rows} или None (не per-order отчёт: нет шапки с «ID заказа»)."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb["Statistics"] if "Statistics" in wb.sheetnames else wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    name = os.path.basename(path)
    hi = next((i for i, r in enumerate(rows[:5]) if r and "ID заказа" in [str(c).strip() if c is not None else "" for c in r]), None)
    if hi is None:
        return None
    win = None
    for r in rows[:hi]:
        line = " ".join(str(c) for c in (r or ()) if c is not None)
        m = re.search(r"Период.*?(\d{2}\.\d{2}\.\d{4})\s*-\s*(\d{2}\.\d{2}\.\d{4})", line)
        if m:
            win = tuple(_date(x) for x in m.groups())
            break
    if not win or not all(win) or win[0] > win[1]:
        raise CpoRawError(f"{name}: нет строки «Период: дд.мм.гггг - дд.мм.гггг» над шапкой - окно выгрузки неизвестно")
    hdr = [str(c).strip() if c is not None else "" for c in rows[hi]]
    ci = {h: j for j, h in enumerate(hdr)}
    for col in ("Дата", "Номер заказа", "SKU", "Артикул", "Расход, ₽"):
        if col not in ci:
            raise CpoRawError(f"{name}: в шапке нет колонки «{col}»")
    out = []
    for k, r in enumerate(rows[hi + 1:], start=hi + 2):
        if r is None or all(c is None for c in r):
            continue
        order = str(r[ci["Номер заказа"]] or "").strip()
        d = _date(r[ci["Дата"]])
        if d is None:
            if order:
                raise CpoRawError(f"{name}, строка {k}: заказ {order} без даты")
            continue  # строка без даты и без заказа (итог, пустая) - не данные
        if not (win[0] <= d <= win[1]):
            raise CpoRawError(f"{name}, строка {k}: дата {d} вне периода файла {win[0]}..{win[1]}")
        sku = str(r[ci["SKU"]]).split(".")[0] if r[ci["SKU"]] is not None else ""
        out.append({"d": d, "order": order, "sku": sku, "offer": str(r[ci["Артикул"]]) if r[ci["Артикул"]] is not None else "",
                    "sp": _num(r[ci["Расход, ₽"]])})
    return {"file": name, "win": win, "rows": out}


# Суффиксы копий одного файла (ФЕНИКС N4 14.10): браузер «x (1)», «x(1)», ОС «x 2», «x-1», «x copy», «x copy 2», «x - копия».
_COPY = re.compile(r"^(.*?)(?:\s?\((\d+)\)|-(\d+)| (\d+)| copy(?: (\d+))?| - копия(?: \((\d+)\))?)$", re.I)


def name_key(path):
    """Ключ порядка файлов: (основа имени, номер копии). Копия идёт после оригинала и побеждает его."""
    b = os.path.basename(path)
    stem, ext = os.path.splitext(b)
    m = _COPY.match(stem)
    if not m or not m.group(1):
        return (b.lower(), 0)
    nums = [g for g in m.groups()[1:] if g]
    n = int(nums[0]) if nums else 1  # «x copy», «x - копия» без номера - первая копия
    if m.group(3) and not re.search(r"[^\d]", m.group(1)[-1:]):
        return (b.lower(), 0)  # «2026-09»: дефис с числом после цифры - часть даты, не копия
    return ((m.group(1) + ext).lower(), n)


def _rk(r):
    return (r["d"], r["order"], r["sku"], round(r["sp"], 2))


def read_raw(raw_dir):
    """Все per-order файлы папки -> (строки после замены периодов, [(файл, окно, строк взято, строк заменено)], пропущенные файлы)."""
    files, skipped = [], []
    for f in sorted(glob.glob(os.path.join(raw_dir, "*.xlsx")), key=name_key):
        x = read_file(f)
        if x is None:
            skipped.append(os.path.basename(f))
        else:
            x["base"] = name_key(f)[0]
            files.append(x)
    # Два файла с одним «Период:», разным содержимым и не копии друг друга - чей файл верный, не ясно: отказ (ФЕНИКС N4).
    # Копии одного файла (одна основа имени) - поздняя побеждает, это исправленное повторное скачивание.
    for i, x in enumerate(files):
        for y in files[i + 1:]:
            if x["win"] == y["win"] and x["base"] != y["base"] and sorted(map(_rk, x["rows"])) != sorted(map(_rk, y["rows"])):
                raise CpoRawError(f"{x['file']} и {y['file']}: один период {x['win'][0]}..{x['win'][1]}, разное содержимое, имена не копии одного файла - оставь один файл (имя YYYY-MM_cpo_orders.xlsx)")
    rows, used = [], []
    for i, x in enumerate(files):
        later = [y["win"] for y in files[i + 1:]]
        keep = [r for r in x["rows"] if not any(a <= r["d"] <= b for a, b in later)]
        rows.extend(keep)
        used.append((x["file"], x["win"], len(keep), len(x["rows"]) - len(keep)))
    return rows, used, skipped
