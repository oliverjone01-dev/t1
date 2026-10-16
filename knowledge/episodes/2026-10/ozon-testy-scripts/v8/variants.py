# v8: варианты «на сколько снизить цену, на сколько поднять рекламу, какие ещё пути» (запрос Ивана 14.10:
# «все варианты расписывай в файле ексель на сколько снизить цену, на сколько поднять рекламу, может ещё какие-то пути
# развития рекомендуешь»). Цели: рост оборота («Начислено») и «К выплате» к 20 млн/мес, при этом видеть чистую «Отчета».
# Модель «Отчета» - та же, что в v6/forecast30.py (go ФЕНИКСА 8,50) и v6/paths20.py (go 8,6): доли структуры расходов
# от «Начислено», АДМ 30% от базы (К выплате + 0,86 × реклама + b × Начислено), налог 15% от «Реализовано»,
# себестоимость от оборота по ценам сентября. Разбор и проверки «Отчета» скопированы из forecast30.py (пометка КОПИЯ),
# потому что forecast30.py при импорте сам собирает свою книгу.
# Запуск: OUT=<папка> node ../v6/report_text.js 2026-03 ... 2026-10; python3 -I variants.py <папка> [analytics-mvp/data]
# Выход: <папка>/OZON-varianty-cena-reklama-<дата>.xlsx (вне репозитория: репо публичный). Если задан RECALC=<путь к
# recalc.py>, книга пересчитывается и каждая ячейка сетки сверяется с расчётом этого скрипта на Python (до рубля).
import re, sys, os, json, math, statistics, subprocess, datetime as dt, collections as C
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter as CL

OUT = sys.argv[1] if len(sys.argv) > 1 else '/tmp/ozon-v8'
ROOT = os.path.abspath(__file__).split('knowledge/')[0]
DATA = sys.argv[2] if len(sys.argv) > 2 else os.path.join(ROOT, 'analytics-mvp/data')
DAY = dt.date(2026, 10, 14)
MON = ['2026-03', '2026-04', '2026-05', '2026-06', '2026-07', '2026-08', '2026-09']
MNAME = ['март', 'апрель', 'май', 'июнь', 'июль', 'август', 'сентябрь']

# ===================================================================== КОПИЯ из v6/forecast30.py: разбор текста «Отчета»
NB = lambda t: t.replace('\xa0', ' ').replace(' ', ' ')

def num(s):
    s = s.replace('−', '-').replace(' ', '').replace('₽', '').replace(',', '.')
    if s.endswith('%'): return float(s[:-1]) / 100
    return float(s) if re.fullmatch(r'[-+]?\d+(\.\d+)?', s) else None

def rows_of(txt, header_start, n):
    out, on = {}, False
    for line in txt.splitlines():
        if line.startswith(header_start + '\t'): on = True; continue
        if on:
            parts = line.split('\t')
            if len(parts) < n + 1: break
            out[parts[0]] = [num(p) for p in parts[1:n + 1]]
    return out

TXT = {m: NB(open(os.path.join(OUT, f'rp_{m}.txt'), encoding='utf-8').read()) for m in MON + ['2026-10']}
sep, octo = TXT['2026-09'], TXT['2026-10']
assert 'Данные OZON по 11.10.2026' in sep, 'снимок «Отчета» не тот: расчёт собран на данных по 11.10'
T3 = rows_of(sep, 'Статья\tмарт 2026', 7); T4 = rows_of(sep, 'Показатель\tмарт 2026', 7); O3 = rows_of(octo, 'Статья\tмарт 2026', 8)
need = ['Начислено (оборот)', 'Комиссия', 'Логистика (с нашей доставкой и доставкой покупателя)', 'Эквайринг', 'Хранение',
        'Прочие (без рекламы)', 'Реклама (весь расход по начислениям OZON)', 'Всего сборов = затраты площадки', 'К выплате',
        'К выплате (с учетом Доставки покупателя)', 'СС произв.', 'Валовая прибыль', 'АДМ 30%', 'Реализовано (база налога)',
        'Налоги 15%', 'Чистая прибыль']
for k in need: assert k in T3 and None not in T3[k], k
UNITS, MIRU, MIRN = [], [], []
for m in MON:
    t = TXT[m]
    a = re.search(r'\nВсего\n([\d ]+) ₽\n.*\nреализовано (\d+) шт', t); b = re.search(r'\nЗеркала\n([\d ]+) ₽\n.*\nреализовано (\d+) шт', t)
    assert a and b, m
    assert abs(num(a.group(1)) - T3['Начислено (оборот)'][MON.index(m)]) <= 1, ('блок 1 против блока 3', m)
    UNITS.append(int(a.group(2))); MIRU.append(int(b.group(2))); MIRN.append(num(b.group(1)))
for j, m in enumerate(MON):                       # инварианты на входе: «Отчет» сходится сам с собой (как в forecast30.py)
    s = {k: T3[k][j] for k in need}
    dinc = s['К выплате (с учетом Доставки покупателя)'] - s['К выплате']
    ship = s['К выплате'] - s['СС произв.'] + dinc - s['Валовая прибыль']
    logoz = s['Всего сборов = затраты площадки'] - sum(s[k] for k in need[1:7] if not k.startswith('Логистика'))
    assert abs(logoz + ship - dinc - s['Логистика (с нашей доставкой и доставкой покупателя)']) <= 20, m
    assert abs(s['Начислено (оборот)'] - s['Всего сборов = затраты площадки'] - s['К выплате']) <= 3, m
    assert abs(s['Валовая прибыль'] - s['АДМ 30%'] - s['Налоги 15%'] - s['Чистая прибыль']) <= 3, m
    assert abs(0.15 * s['Реализовано (база налога)'] - s['Налоги 15%']) <= 3, m
    assert abs(T4['за клик'][j] + T4['за заказ'][j] - s['Реклама (весь расход по начислениям OZON)']) <= 0.01 * s['Реклама (весь расход по начислениям OZON)'], m
# конец КОПИИ
# «Отчет», блок 5: товары с показами и без заказов 60 дн
def nozero(t, cat):
    m = re.search(cat + r' - всего товаров с показами и без заказов 60 дн: (\d+)', t); assert m, cat; return int(m.group(1))
NZ = {'sep_m': nozero(sep, 'Зеркала'), 'sep_f': nozero(sep, 'Мебель'), 'oct_m': nozero(octo, 'Зеркала'), 'oct_f': nozero(octo, 'Мебель')}

# ===================================================================== данные analytics-mvp/data (сентябрь, заказы без отмен)
def nd(f): return [json.loads(l) for l in open(os.path.join(DATA, f), encoding='utf-8')]
ORD = [o for o in nd('orders_daily.ndjson') if o['status'] != 'cancelled']
TAX = json.load(open(os.path.join(DATA, 'sku_taxonomy.json'), encoding='utf-8'))
COGS = json.load(open(os.path.join(DATA, 'sku_cogs.json'), encoding='utf-8'))
assert max(o['d'] for o in ORD) >= '2026-10-11', 'orders_daily короче 11.10'
SEPO = [o for o in ORD if o['d'][:7] == '2026-09']
N_sep_o = sum(o['revenue'] for o in SEPO); U_sep_o = sum(o['units'] for o in SEPO)
CC_sep_o = sum(COGS.get(o['sku'], 0) * o['units'] for o in SEPO)
# ходовые = SKU с 3+ шт без отмен за 60 дней 13.08-11.10 (дата заказа). Порог 3, а не 1: с порогом 1 это все продавшиеся
# товары (100% оборота сентября), то есть снижение цены всего магазина. Пороги 1-5 - в «Допущениях».
U60 = C.Counter()
for o in ORD:
    if '2026-08-13' <= o['d'] <= '2026-10-11': U60[o['sku']] += o['units']
HOD = {}
for thr in (1, 2, 3, 4, 5):
    S = {s for s, u in U60.items() if u >= thr}
    rs = sum(o['revenue'] for o in SEPO if o['sku'] in S); us = sum(o['units'] for o in SEPO if o['sku'] in S)
    cs = sum(COGS.get(o['sku'], 0) * o['units'] for o in SEPO if o['sku'] in S)
    HOD[thr] = dict(n=len(S), mir=sum(1 for s in S if TAX.get(s, {}).get('category') == 'Зеркала'), srev=rs / N_sep_o, su=us / U_sep_o,
                    ccrel=(cs / rs) / (CC_sep_o / N_sep_o), skus=S)
    assert 0 < rs <= N_sep_o + 1 and 0 < us <= U_sep_o, thr
THR = 3; H = HOD[THR]
assert 0.5 < H['srev'] < 0.95 and 0.9 < H['ccrel'] < 1.05, H
# комиссия и логистика OZON по схеме, июль-сентябрь, доставленные заказы (orders_daily: доля от revenue)
SCH = {}
for o in ORD:
    if '2026-07' <= o['d'][:7] <= '2026-09' and o['status'] == 'delivered':
        a = SCH.setdefault(o['scheme'], C.Counter()); a['u'] += o['units']
        for k in ('revenue', 'commission', 'delivery', 'partner', 'other', 'payout'): a[k] += o[k]
for k, a in SCH.items():
    a['comm'] = -a['commission'] / a['revenue']; a['logi'] = -(a['delivery'] + a['partner']) / a['revenue']; a['p'] = a['revenue'] / a['u']
OCTS = C.Counter(o['scheme'] for o in ORD if o['d'][:7] == '2026-10' for _ in range(o['units']))
# цена витрины сентября к июню у ходовых (coinv_daily, observed, site_listed; медиана по дням), вес - оборот сентября
# «обычная» цена как в разборе v7 = медиана SKU за окно панели 15.06-20.09; июнь = 12-30.06 (как индекс витрины разбора)
PJ, PS, PW = C.defaultdict(list), C.defaultdict(list), C.defaultdict(list)
for r in nd('coinv_daily.ndjson'):
    if not (r.get('observed') and r.get('site_listed')) or r.get('sku') not in H['skus']: continue
    if '2026-06-12' <= r['date'] <= '2026-06-30': PJ[r['sku']].append(r['site_listed'])
    if '2026-06-15' <= r['date'] <= '2026-09-20': PW[r['sku']].append(r['site_listed'])
    if r['date'][:7] == '2026-09': PS[r['sku']].append(r['site_listed'])
REVS = C.Counter()
for o in SEPO:
    if o['sku'] in H['skus']: REVS[o['sku']] += o['revenue']
both = [s for s in H['skus'] if PJ[s] and PS[s]]
IDX = {s: statistics.median(PS[s]) / statistics.median(PJ[s]) for s in both}
w_all = sum(REVS[s] for s in both)
BORDER = dict(n=len(both), n_all=len(H['skus']), cover=w_all / sum(REVS.values()),
              over20=sum(REVS[s] for s in both if IDX[s] > 1.20) / w_all, over10=sum(REVS[s] for s in both if IDX[s] > 1.10) / w_all,
              n20=sum(1 for s in both if IDX[s] > 1.20), med=statistics.median(IDX.values()))
IDW = {s: statistics.median(PS[s]) / statistics.median(PW[s]) for s in both}
BORDER.update(wmed=statistics.median(IDW.values()), w_over10=sum(REVS[s] for s in both if IDW[s] > 1.10) / w_all,
              w_below=sum(REVS[s] for s in both if IDW[s] < 0.97) / w_all)
assert 0 < BORDER['over20'] <= 1 and 0 < BORDER['cover'] <= 1
# показы сентября на товарах без заказов (sku_views; без заказов = 0 шт без отмен 02.08-30.09, как «60 дн» блока 5)
U60b = C.Counter()
for o in ORD:
    if '2026-08-02' <= o['d'] <= '2026-09-30': U60b[o['sku']] += o['units']
VS = C.Counter()
for r in nd('sku_views.ndjson'):
    if r['date'][:7] == '2026-09': VS[r['sku']] += r['views']
VZ = dict(n=sum(1 for s, v in VS.items() if v > 0 and U60b[s] == 0), share=sum(v for s, v in VS.items() if U60b[s] == 0) / sum(VS.values()))

# ===================================================================== константы из прошлых расчётов (источник рядом)
TEST1 = dict(u=2.1, days=23, cw=1500, pool=23, expo=0.81, p=57000)  # тест 1 и тест 3 (итоги тестов 12.10; handoff v8 §4)
KS = {'B': 0.3, 'C': 0.5, 'D': 1.0}; EPS = {'B': -1.0, 'C': -2.0, 'D': -3.0}; RO = {'B': 'G', 'C': 'F', 'D': 'H'}  # R клика «Отчета»: авг/июль/сент
DLIST = [0.0, 0.05, 0.10, 0.15, 0.20]; BLIST = [0, 250_000, 500_000, 1_000_000, 1_500_000, 2_000_000]
LLIST = [0.0, -0.15, -0.33, -0.50]
WEEKS = 30.4 / 7

# ===================================================================== Python-зеркало модели (для инвариантов и сверки книги)
FCOL = {'B': (4, 6), 'C': (6, 6)}           # структура: B = июль-сентябрь (индексы 4..6), C = сентябрь (6)
def ssum(k, sc, tab=None):
    a, b = FCOL[sc]; tab = tab or T3; return sum(tab[k][a:b + 1])
def struct(sc):
    a, b = FCOL[sc]
    N = ssum('Начислено (оборот)', sc); U = sum(UNITS[a:b + 1]); adacc = ssum('Реклама (весь расход по начислениям OZON)', sc)
    other = sum(ssum(k, sc) for k in ['Комиссия', 'Эквайринг', 'Хранение', 'Прочие (без рекламы)'])
    logoz = ssum('Всего сборов = затраты площадки', sc) - other - adacc
    dinc = ssum('К выплате (с учетом Доставки покупателя)', sc) - ssum('К выплате', sc)
    ship = ssum('К выплате', sc) - ssum('СС произв.', sc) + dinc - ssum('Валовая прибыль', sc)
    s = dict(comm=ssum('Комиссия', sc) / N, logi=logoz / N, acq=ssum('Эквайринг', sc) / N, sto=ssum('Хранение', sc) / N,
             oth=ssum('Прочие (без рекламы)', sc) / N, cpo=ssum('за заказ', sc, T4) / N, clk=ssum('за клик', sc, T4) / N,
             cc=ssum('СС произв.', sc) / N, dinc=dinc / U, ship=ship / U, real=ssum('Реализовано (база налога)', sc) / N, adm=0.30, kadm=0.86, tax=0.15)
    s['badm'] = (ssum('АДМ 30%', sc) / 0.3 - ssum('К выплате', sc) - 0.86 * adacc) / N
    return s
ST = {sc: struct(sc) for sc in 'BC'}
N0, U0 = T3['Начислено (оборот)'][6], UNITS[6]; P0 = N0 / U0

def pl(s, N, U, cogs, adcpo, adclk):
    ad = adcpo + adclk; fees = N * (s['comm'] + s['logi'] + s['acq'] + s['sto'] + s['oth']) + ad; kv = N - fees
    gross = kv - cogs - U * s['ship'] + U * s['dinc']; admv = s['adm'] * (kv + s['kadm'] * ad + s['badm'] * N)
    return dict(N=N, U=U, ad=ad, kv=kv, net=gross - admv - s['tax'] * s['real'] * N, adsh=ad / N)

# отдача клика по артикулам рекламы «Отчета», блок 4 (зеркала + мебель): низкий - август, средний - июль, высокий - сентябрь
ROBS = {sc: (T4['Зеркала, за клик: выручка'][j] + T4['Мебель, за клик: выручка'][j]) / (T4['Зеркала, за клик: расход'][j] + T4['Мебель, за клик: расход'][j])
        for sc, j in (('B', 5), ('C', 4), ('D', 6))}
assert all(1 < v < 4 for v in ROBS.values()), ROBS
# правка ФЕНИКСА V1+V2 (go 8,75): база - сверх пула 23 товаров отдача не выше нынешнего клика по «Отчету» (mode='cap');
# показы ~ расход^0,81 - отдельная таблица «верхняя граница» (mode='up')
def click_units(B, k, scen='C', mode='cap'):
    ui = TEST1['u'] * 30.4 / TEST1['days'] * k; cpi = TEST1['cw'] * WEEKS; Bp = TEST1['pool'] * cpi
    if B <= Bp: return B / cpi * ui
    if mode == 'cap': return TEST1['pool'] * ui + (B - Bp) * ROBS[scen] / TEST1['p']
    return TEST1['pool'] * ui * (B / Bp) ** TEST1['expo']

def grid_m(sc, scen, d, B, mode='cap'):
    s = ST[sc]; g = (1 - d) ** EPS[scen]
    dU = U0 * H['su'] * (g - 1); dNS = N0 * H['srev'] * (g - 1); dN = N0 * H['srev'] * ((1 - d) * g - 1)
    uc = click_units(B, KS[scen], scen, mode); Nc = uc * TEST1['p']
    N = N0 + dN + Nc
    return pl(s, N, U0 + dU + uc, s['cc'] * (N0 + H['ccrel'] * dNS + Nc), s['cpo'] * N, s['clk'] * N0 + B)

# инвариант (data-guard шаг 4): «0% цены, 0 клика» в структуре сентября повторяет сентябрь «Отчета» до рубля
b0 = grid_m('C', 'C', 0, 0)
assert abs(b0['net'] - T3['Чистая прибыль'][6]) <= 1, ('чистая сентября', b0['net'], T3['Чистая прибыль'][6])
assert abs(b0['kv'] - T3['К выплате'][6]) <= 1 and abs(b0['ad'] - T3['Реклама (весь расход по начислениям OZON)'][6]) <= 1
for sc in 'BC':                                   # монотонность: клик растит оборот, цена при ε<-1 растит штуки
    for scen in 'BCD':
        assert all(grid_m(sc, scen, 0, BLIST[i])['N'] < grid_m(sc, scen, 0, BLIST[i + 1])['N'] for i in range(5))
        assert all(grid_m(sc, scen, DLIST[i], 0)['U'] <= grid_m(sc, scen, DLIST[i + 1], 0)['U'] + 1e-9 for i in range(4))
GRID = {(sc, scen, d, B): grid_m(sc, scen, d, B) for sc in 'BC' for scen in 'BCD' for d in DLIST for B in BLIST}
GRIDU = {(sc, scen, d, B): grid_m(sc, scen, d, B, 'up') for sc in 'BC' for scen in 'BCD' for d in DLIST for B in BLIST}

if __name__ == '__main__' and os.environ.get('MIRROR_ONLY'):
    print('ходовые', {t: (v['n'], v['mir'], round(v['srev'], 3), round(v['su'], 3), round(v['ccrel'], 3)) for t, v in HOD.items()})
    print('схемы', {k: (a['u'], round(a['comm'], 4), round(a['logi'], 4), round(a['p'])) for k, a in SCH.items()}, 'окт', OCTS)
    print('граница', BORDER); print('без заказов', VZ, NZ)
    for sc in 'BC':
        print('структура', sc, {k: round(v, 4) for k, v in ST[sc].items()})
        for scen in 'BCD':
            print(' сценарий', scen)
            for d in DLIST:
                print('  d', d, ' '.join(f"{GRID[(sc, scen, d, B)]['kv'] / 1e6:5.2f}/{GRID[(sc, scen, d, B)]['net'] / 1e6:+5.2f}" for B in BLIST))
    sys.exit()

# ===================================================================== книга: стиль (как в forecast30.py)
wb = Workbook(); F = 'Arial'
fN = Font(name=F, size=10); fB = Font(name=F, size=10, bold=True); fH = Font(name=F, size=13, bold=True)
fIn = Font(name=F, size=10, color='0000FF'); fLink = Font(name=F, size=10, color='008000')
fGrey = Font(name=F, size=9, color='666666'); fWhite = Font(name=F, size=10, bold=True, color='FFFFFF')
fillHead = PatternFill('solid', fgColor='1F3864'); fillKey = PatternFill('solid', fgColor='FFFF00')
fillSub = PatternFill('solid', fgColor='D9E1F2'); fillRes = PatternFill('solid', fgColor='E2EFDA'); fillBad = PatternFill('solid', fgColor='FCE4D6')
fillG = PatternFill('solid', fgColor='C6EFCE'); fillR = PatternFill('solid', fgColor='FFC7CE'); fillT = PatternFill('solid', fgColor='BDD7EE')
thin = Side(style='thin', color='BFBFBF'); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
RUB = '#,##0;-#,##0;"-"'; PCT = '0.0%;-0.0%;"-"'; X1 = '0.0;-0.0;"-"'; X2 = '0.00;-0.00;"-"'; SHT = '#,##0;-#,##0;"-"'; DPCT = '"-"0%;"+"0%;"0%"'

def put(ws, ref, v, font=fN, fmt=None, fill=None, wrap=False, bold=False):
    c = ws[ref]; c.value = v; c.font = fB if bold and font is fN else font
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if wrap: c.alignment = Alignment(wrap_text=True, vertical='top')
    return c

def head(ws, r, labels, widths=None, c0=1):
    for i, t in enumerate(labels):
        c = ws.cell(row=r, column=c0 + i, value=t); c.font = fWhite; c.fill = fillHead
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center'); c.border = BOX
    if widths:
        for i, w in enumerate(widths): ws.column_dimensions[CL(c0 + i)].width = w

def note(ws, r, t, ncols=6, h=30, font=None):
    put(ws, f'A{r}', t, font or fGrey, wrap=True); ws.merge_cells(f'A{r}:{CL(ncols)}{r}'); ws.row_dimensions[r].height = h

M = lambda x: f"SUBSTITUTE(TEXT(({x})/1000000,\"0.00\"),\".\",\",\")"           # млн с запятой
P1 = lambda x: f"SUBSTITUTE(TEXT(({x})*100,\"0.0\"),\".\",\",\")"             # проценты с запятой
X_ = lambda x: f"SUBSTITUTE(TEXT({x},\"0.0\"),\".\",\",\")"
X2_ = lambda x: f"SUBSTITUTE(TEXT({x},\"0.00\"),\".\",\",\")"

# ----------------------------------------------------------------- лист «Факт Отчет» (КОПИЯ из forecast30.py, без бэктеста)
wf = wb.active; wf.title = 'Факт Отчет'
put(wf, 'A1', 'Факт «Отчета» OZON по месяцам (страница «Отчет», блоки 1-4; данные OZON по 11.10.2026) [ДАННЫЕ]', fH)
put(wf, 'A2', 'Строки со страницы перенесены как есть, ниже - производные строки формулами. Октябрь - только 01-11.10.', fGrey)
head(wf, 4, ['Статья, ₽'] + [m + ' 2026' for m in MNAME] + ['октябрь 01-11'], [56] + [14] * 8)
FROW = {}; r = 5
for k in need:
    put(wf, f'A{r}', k, bold=k in ('Начислено (оборот)', 'К выплате', 'Чистая прибыль'))
    for j in range(7): put(wf, f'{CL(j + 2)}{r}', T3[k][j], fmt=RUB)
    put(wf, f'I{r}', O3[k][7] if k in O3 else None, fmt=RUB)
    FROW[k] = r; r += 1
for lab, key in [('Реклама за клик (рекламный кабинет)', 'за клик'), ('Реклама за заказ (рекламный кабинет)', 'за заказ'),
                 ('Зеркала, за клик: расход', 'Зеркала, за клик: расход'), ('Зеркала, за клик: выручка', 'Зеркала, за клик: выручка'),
                 ('Мебель, за клик: расход', 'Мебель, за клик: расход'), ('Мебель, за клик: выручка', 'Мебель, за клик: выручка')]:
    put(wf, f'A{r}', lab)
    for j in range(7): put(wf, f'{CL(j + 2)}{r}', T4[key][j], fmt=RUB)
    FROW[key] = r; r += 1
for lab, key, vals in [('Реализовано, шт (блок 1)', 'шт', UNITS), ('Зеркала: реализовано, шт (блок 1)', 'зерк', MIRU), ('Зеркала: Начислено, ₽ (блок 1)', 'зерк₽', MIRN)]:
    put(wf, f'A{r}', lab)
    for j in range(7): put(wf, f'{CL(j + 2)}{r}', vals[j], fmt=SHT)
    FROW[key] = r; r += 1
r += 1; put(wf, f'A{r}', 'Производные строки (формулы)', fB, fill=fillSub); r += 1
V = lambda k, c: f'{c}{FROW[k]}'
derived = [
    ('logoz', 'Логистика OZON без нашей доставки (= Всего сборов − остальные сборы; как блок 2)',
     lambda c: f"={V('Всего сборов = затраты площадки', c)}-{V('Комиссия', c)}-{V('Эквайринг', c)}-{V('Хранение', c)}-{V('Прочие (без рекламы)', c)}-{V('Реклама (весь расход по начислениям OZON)', c)}"),
    ('dinc', 'Доставка покупателя (доход) = К выплате с доставкой − К выплате', lambda c: f"={V('К выплате (с учетом Доставки покупателя)', c)}-{V('К выплате', c)}"),
    ('ship', 'Наша доставка = К выплате − СС + доставка покупателя − Валовая', lambda c: f"={V('К выплате', c)}-{V('СС произв.', c)}+{V('dinc', c)}-{V('Валовая прибыль', c)}"),
    ('admb', 'База АДМ = АДМ / 0,3', lambda c: f"={V('АДМ 30%', c)}/0.3"),
    ('clkpb', 'Отдача клика по артикулам рекламы, ₽ на 1 ₽ (зеркала + мебель)',
     lambda c: f"=IFERROR(({V('Зеркала, за клик: выручка', c)}+{V('Мебель, за клик: выручка', c)})/({V('Зеркала, за клик: расход', c)}+{V('Мебель, за клик: расход', c)}),0)"),
    ('commsh', 'Комиссия / Начислено', lambda c: f"={V('Комиссия', c)}/{V('Начислено (оборот)', c)}"),
    ('kvsh', 'К выплате / Начислено', lambda c: f"={V('К выплате', c)}/{V('Начислено (оборот)', c)}"),
]
for key, lab, fn in derived:
    put(wf, f'A{r}', lab); FROW[key] = r
    for j in range(7):
        c = CL(j + 2); put(wf, f'{c}{r}', fn(c), fmt=X2 if key == 'clkpb' else (PCT if key in ('commsh', 'kvsh') else RUB))
    if key in ('commsh', 'kvsh'): put(wf, f'I{r}', fn('I'), fmt=PCT)
    r += 1
for c in 'BCD': wf[f'{c}{FROW["clkpb"]}'].value = None
wf.freeze_panes = 'B5'
FR = lambda k, col='H': f"'Факт Отчет'!${col}${FROW[k]}"

# ----------------------------------------------------------------- лист «Допущения»
wa = wb.create_sheet('Допущения')
put(wa, 'A1', 'Допущения: синие - ввод (можно менять), зелёные - из «Отчета» или данных, чёрные - формула', fH)
put(wa, 'A2', 'Колонки B-D в каждом разделе значат разное - смотрите строку-заголовок раздела. Все цифры с меткой и источником.', fGrey)
for c, w in zip('ABCDEF', [60, 15, 15, 15, 8, 120]): wa.column_dimensions[c].width = w
A = {}; r = 4
def sect(t, cols=('', '', '')):
    global r
    r += 1; put(wa, f'A{r}', t, fB, fill=fillSub)
    for col, lab in zip('BCD', cols): put(wa, f'{col}{r}', lab, fB, fill=fillSub)
    for col in 'EF': wa[f'{col}{r}'].fill = fillSub
    r += 1
def inp(name, label, vals, unit, src, fmt=RUB, kind='in', key=False, ncol=3):
    global r
    put(wa, f'A{r}', label, wrap=len(label) > 70)
    if callable(vals): vals = [vals(sc) for sc in 'BCD'[:ncol]]
    if not isinstance(vals, (list, tuple)): vals = [vals]
    font = fLink if kind == 'link' else (fN if kind == 'f' else fIn)
    if len(vals) == 1:
        put(wa, f'B{r}', vals[0], font, fmt, fillKey if key else None)
        for col in 'CD'[:ncol - 1]: put(wa, f'{col}{r}', f'=$B{r}', fN, fmt)
    else:
        for col, v in zip('BCD', vals): put(wa, f'{col}{r}', v, font, fmt, fillKey if key else None)
    put(wa, f'E{r}', unit, fGrey); put(wa, f'F{r}', src, fGrey, wrap=True)
    A[name] = r; r += 1
AR = lambda name, col='B': f"'Допущения'!${col}${A[name]}"

SC = {'B': ('F', 'H'), 'C': ('H', 'H')}           # структура: B = июль-сентябрь, C = сентябрь (колонки «Факт Отчет»)
def SUMR(k, sc): a, b = SC[sc]; return f"SUM('Факт Отчет'!${a}${FROW[k]}:${b}${FROW[k]})"
def share(k, sc): return f"={SUMR(k, sc)}/{SUMR('Начислено (оборот)', sc)}"
def perunit(k, sc): return f"={SUMR(k, sc)}/{SUMR('шт', sc)}"

sect('1. Стартовая точка: сентябрь 2026 по «Отчету» [ДАННЫЕ]', ('значение', '', ''))
inp('N0', 'Начислено (оборот), сентябрь', f"={FR('Начислено (оборот)')}", '₽', '«Отчет», блок 3', kind='link', ncol=1)
inp('U0', 'Реализовано, шт, сентябрь', f"={FR('шт')}", 'шт', '«Отчет», блок 1', SHT, kind='link', ncol=1)
inp('P0', 'Начислено на 1 шт, сентябрь', f"=B{A['N0']}/B{A['U0']}", '₽', 'формула', kind='f', ncol=1)
inp('KVT', 'Цель: К выплате в месяц', 20_000_000, '₽', 'Иван 14.10: «К выплате = 20 млн»', key=True, ncol=1)

sect('2. Структура расходов, доли от Начислено [ДАННЫЕ «Отчет»]', ('июль-сентябрь', 'сентябрь', ''))
for nm, lab, f_, fmt, src in [
        ('comm', 'Комиссия', lambda sc: share('Комиссия', sc), PCT, 'в Начислено входят баллы OZON за скидки, поэтому комиссия около половины. Октябрь 01-11 - 50,8%'),
        ('logi', 'Логистика OZON (без нашей доставки)', lambda sc: share('logoz', sc), PCT, 'как блок 2 «Отчета»'),
        ('acq', 'Эквайринг', lambda sc: share('Эквайринг', sc), PCT, ''),
        ('sto', 'Хранение', lambda sc: share('Хранение', sc), PCT, ''),
        ('oth', 'Прочие (без рекламы)', lambda sc: share('Прочие (без рекламы)', sc), PCT, ''),
        ('cpo', 'Реклама за заказ (ставка 5%)', lambda sc: share('за заказ', sc), PCT, 'растёт вместе с оборотом; 9% не окупается (v5, go ФЕНИКСА 8,45)'),
        ('clk', 'Реклама за клик сейчас, доля от Начислено', lambda sc: share('за клик', sc), PCT, 'в рублях фиксируется на уровне «доля × Начислено сентября»; дополнительный клик сетки - сверху'),
        ('cc', 'Себестоимость', lambda sc: share('СС произв.', sc), PCT, 'от оборота по ценам сентября: при снижении цены себестоимость штуки не падает'),
        ('dinc', 'Доставка покупателя (доход) на 1 шт', lambda sc: perunit('dinc', sc), RUB, 'К выплате с доставкой − К выплате'),
        ('ship', 'Наша доставка (расход) на 1 шт', lambda sc: perunit('ship', sc), RUB, 'Валовая = К выплате − СС − наша доставка + доставка покупателя'),
        ('real', 'Реализовано (база налога)', lambda sc: share('Реализовано (база налога)', sc), PCT, 'баллы OZON в базу налога не входят')]:
    inp(nm, lab, f_, '%' if fmt == PCT else '₽', src, fmt, kind='link', ncol=2)
inp('adm', 'АДМ, ставка', 0.30, '%', '«Отчет»: 30% от базы АДМ (Иван 25.09)', PCT, ncol=2)
inp('kadm', 'База АДМ: доля рекламы, которая базу НЕ уменьшает', 0.86, '%', 'ФЕНИКС 12.10: 86% рекламы сентября базу АДМ не уменьшили', PCT, ncol=2)
KROW = r - 1
inp('badm', 'База АДМ: прочая добавка, доля от Начислено',
    lambda sc: f"=({SUMR('admb', sc)}-{SUMR('К выплате', sc)}-$B${KROW}*{SUMR('Реклама (весь расход по начислениям OZON)', sc)})/{SUMR('Начислено (оборот)', sc)}",
    '%', 'калибровка: база = К выплате + 0,86 × реклама + доля × Начислено; месяц структуры сходится до рубля [РАСЧЁТ]', PCT, kind='link', ncol=2)
inp('tax', 'Налог, ставка от «Реализовано»', 0.15, '%', '«Отчет»', PCT, ncol=2)

sect('3. Сценарии силы рычагов [ГИПОТЕЗА]', ('низкий', 'средний', 'высокий'))
inp('eps', 'Эластичность заказов ходовых SKU к цене (на весь набор сразу)', [EPS[c] for c in 'BCD'], 'x',
    'Разбор v7: -4 - относительная недельная эластичность (карточка против самой себя в другие недели на фоне магазина). В неё входит перетекание спроса между нашими товарами и перенос покупки на дешёвую неделю. '
    'Когда дешевеют сразу 69 ходовых SKU (76% оборота), перетекать почти некуда, а перенос покупки на месяц не работает, поэтому для магазина целиком берём -1 / -2 / -3: модуль меньше 4 [ГИПОТЕЗА, долгосрочная эластичность в данных не оценена]. Формула: шт = шт × (1 − снижение)^ε', X1, key=True)
inp('kclk', 'Отдача дешёвого клика: доля от осторожной оценки теста 1', [KS[c] for c in 'BCD'], 'x', 'тест 1: +2,1 шт на товар за 23 дня (осторожный счёт, верх). Низкий ×0,3 (P9), средний ×0,5 (одна волна, 23 дня), высокий ×1', X2, key=True)
inp('Robs', 'Отдача клика сверх пула (база сетки) = нынешний клик по «Отчету», ₽ выручки на 1 ₽', [f"={FR('clkpb', RO[c])}" for c in 'BCD'], '₽/₽',
    '«Отчет», блок 4: август / июль / сентябрь. Выручка заказов после клика, а не прирост - верхняя граница для нынешних кампаний. Правка ФЕНИКСА V1+V2: в базе сетки клик сверх пула 23 товаров отдаёт не больше этого', X2, kind='link', key=True)

sect('4. Ходовые SKU (на них снижаем цену) [ДАННЫЕ orders_daily, sku_cogs]', ('значение', '', ''))
inp('thr', 'Определение: SKU с продажами 3+ шт без отмен за 60 дней 13.08-11.10', THR, 'шт',
    'Порог 3 выбран, чтобы набор был «ходовым», а не всем, что продалось (порог 1 = 100% оборота). Другие пороги, SKU / доля оборота сентября: ' +
    '; '.join(f"{t} шт - {v['n']} SKU / {v['srev'] * 100:.0f}%".replace('.', ',') for t, v in HOD.items()), SHT, ncol=1)
inp('nhod', 'Ходовых SKU', H['n'], 'шт', f"из них зеркал {H['mir']} [ДАННЫЕ]", SHT, kind='link', ncol=1)
inp('srev', 'Доля ходовых в обороте сентября', H['srev'], '%', f"orders_daily, сентябрь, заказы без отмен, дата заказа: оборот {N_sep_o:,.0f} ₽ (в «Отчете» 21 010 545 ₽) [ДАННЫЕ]".replace(',', ' '), PCT, kind='link', ncol=1)
inp('su', 'Доля ходовых в штуках сентября', H['su'], '%', f'orders_daily, {U_sep_o} шт без отмен [ДАННЫЕ]', PCT, kind='link', ncol=1)
inp('ccrel', 'Себестоимость ходовых относительно магазина', H['ccrel'], 'x', f"sku_cogs × шт сентября: {H['ccrel'] * CC_sep_o / N_sep_o * 100:.2f}% у ходовых против {CC_sep_o / N_sep_o * 100:.2f}% у магазина [ДАННЫЕ]".replace('.', ','), X2, kind='link', ncol=1)
inp('bj', 'Цена витрины ходовых в сентябре к июню, медиана', round(BORDER['med'] - 1, 4), '%',
    f"coinv_daily, observed, site_listed; {BORDER['n']} из {BORDER['n_all']} SKU с ценой в обоих периодах, {BORDER['cover'] * 100:.0f}% оборота ходовых. Дороже июня больше чем на 20%: {BORDER['over20'] * 100:.0f}% оборота ({BORDER['n20']} SKU) [ДАННЫЕ]".replace('.', ','), PCT, kind='link', ncol=1)
inp('bwo', 'Ходовые: доля оборота, где витрина сентября дороже «обычной» больше чем на 10%', BORDER['w_over10'], '%', 'coinv_daily, как строка ниже [ДАННЫЕ]', PCT, kind='link', ncol=1)
inp('bw', 'Цена витрины ходовых в сентябре к «обычной» (медиана SKU 15.06-20.09, как в разборе v7)', round(BORDER['wmed'] - 1, 4), '%',
    f"дороже обычной больше чем на 10%: {BORDER['w_over10'] * 100:.0f}% оборота; дешевле обычной больше чем на 3%: {BORDER['w_below'] * 100:.0f}% [ДАННЫЕ]. Обычная цена уже летняя, поднятая".replace('.', ','), PCT, kind='link', ncol=1)

sect('5. Дешёвый клик 8 ₽ на товарах без рекламы или со снятой рекламой [ГИПОТЕЗА, ШИРОКИЙ ДИАПАЗОН - НЕПРОВЕРЕНО]', ('низкий', 'средний', 'высокий'))
inp('ut', 'Тест 1: прирост на 1 товар за срок теста', TEST1['u'], 'шт', 'тест 1, осторожный счёт: за 23 дня +12,6 шт на 6 товарах (без KONUX и шумного GGL-09-2) = 2,1 [ДАННЫЕ, n мал]', X2)
inp('dt', 'Тест 1: дней', TEST1['days'], 'дн', 'тест 1', SHT)
inp('dm', 'Дней в месяце', 30.4, 'дн', '', X1)
inp('ui', 'Прирост на 1 товар в месяц', lambda c: f"=B{A['ut']}*B{A['dm']}/B{A['dt']}*{c}{A['kclk']}", 'шт', 'формула: 2,1 × 30,4 / 23 × доля сценария', X2, kind='f')
inp('cw', 'Расход на 1 товар в неделю', TEST1['cw'], '₽', 'тест 1: факт 1 523 ₽/нед, лимит кампании 1 500 ₽')
inp('wk', 'Недель в месяце', '=30.4/7', '', '30,4 / 7', X2, kind='f')
inp('cpi', 'Расход на 1 товар в месяц', f"=B{A['cw']}*B{A['wk']}", '₽', 'формула', kind='f')
inp('pool', 'Пул товаров по условиям теста 1', TEST1['pool'], 'шт', 'итоги тестов 12.10 (условия ФЕНИКСА): 23 товара в 16 карточках. Ещё 725 товаров с показами без заказов (блок 5 «Отчета») - на них механика не проверялась', SHT)
inp('Bp', 'Бюджет пула в месяц', f"=B{A['pool']}*B{A['cpi']}", '₽', 'пул × расход на товар; до этой суммы прирост линейный', kind='f')
inp('expo', 'Только для таблицы «верхняя граница»: показатель убывающей отдачи сверх пула', TEST1['expo'], 'x',
    'тест 3: показы ~ расход^0,81. Допущение верхней границы: сверх пула штуки растут как показы, то есть шт = шт пула × (бюджет / бюджет пула)^0,81. Это ОЦЕНКА СВЕРХУ: тест 3 показал, что дорогой клик даёт показы, а заказов почти нет; '
    'на товарах без заказов механика не проверялась. Средняя отдача при таком допущении выше, чем у нынешних кампаний по «Отчету», поэтому в базу не берётся', X2)
inp('indep', 'Цена и клик в сетке сложены независимо', 'да', '', 'Ограничение модели (ФЕНИКС V3): клик «на товарах без рекламы или со снятой рекламой» пересекается с 69 ходовыми SKU (летом клик сняли с ходовых карточек), поэтому совместный эффект цены и клика на одних товарах может быть и больше, и меньше суммы', None, kind='f', ncol=1)
inp('pc', 'Начислено на 1 шт от клика', TEST1['p'], '₽', 'тест 1: средняя предельная цена тестовых товаров без KONUX 57 054 ₽')

sect('6. Сетка: шаги цены и бюджета (можно менять)', ('значение', '', ''))
for i, d in enumerate(DLIST): inp(f'd{i}', f'Снижение цены ходовых, шаг {i + 1}', d, '%', 'решение 13.10: шаг 5-10%; больше 5% - HITL (Иван); больше 10% - новое решение Ивана' if i == 0 else '', PCT, ncol=1)
for i, b in enumerate(BLIST): inp(f'b{i}', f'Дополнительный дешёвый клик в месяц, шаг {i + 1}', b, '₽', 'сверх нынешнего клика; больше 100 тыс. - HITL (Иван)' if i == 0 else '', ncol=1)

sect('7. Ставка за заказ: что будет, если выключить [ДАННЫЕ окна 29.08-03.09; эффект - ГИПОТЕЗА]', ('значение', '', ''))
for i, l in enumerate(LLIST): inp(f'L{i}', f'Вариант «0%»: изменение штук, шаг {i + 1}', l, '%', ['без потери штук (только показы)', 'половина потери окна', 'как в окне 29.08-03.09: 6,3 шт/день против 9,6 и 9,3 в соседние две недели (-33%)', 'штуки падают вместе с показами (-50%)'][i], PCT, ncol=1)
inp('cpo9', 'Ставка 9%: множитель расхода к 5%', 1.8, 'x', 'v5: ставка 9% заказов не добавила (-2,8% на пункт, ДИ -7,6...+2,3); расход ×9/5 [ДАННЫЕ v5]', X2, ncol=1)

sect('8. Другие пути: входы', ('значение', '', ''))
inp('Usp', 'Весенний объём: шт в апреле (лучший месяц)', f"={FR('шт', 'C')}", 'шт', '«Отчет», блок 1: апрель 581 шт', SHT, kind='link', ncol=1)
inp('Mjun', 'Зеркала: шт в июне', f"={FR('зерк', 'E')}", 'шт', '«Отчет», блок 1', SHT, kind='link', ncol=1)
inp('Msep', 'Зеркала: шт в сентябре', f"={FR('зерк')}", 'шт', '«Отчет», блок 1', SHT, kind='link', ncol=1)
inp('Mp', 'Зеркала: Начислено на 1 шт, сентябрь', f"={FR('зерк₽')}/B{A['Msep']}", '₽', 'формула', kind='f', ncol=1)
inp('cut', 'Срез дорогого клика (две кампании с 0-1 заказом и KONUX, без KUVINO)', 250_000, '₽', 'записка «К выплате 20 млн», темп 01-11.10 (проба ФЕНИКСА p_campaigns); KUVINO не трогаем (решение 13.10)', ncol=1)
inp('wk11', 'Распродажа 11.11: доля месяца (одна неделя)', '=7/30.4', '%', '7 дней из 30,4', PCT, kind='f', ncol=1)

sect('9. Пороги HITL (Protocol 6)', ('значение', '', ''))
inp('hp', 'Изменение цены без решения Ивана - не больше', 0.05, '%', 'CLAUDE.md, HITL', PCT, ncol=1)
inp('ha', 'Изменение рекламы без решения Ивана - не больше', 100_000, '₽', 'CLAUDE.md, HITL', ncol=1)
wa.freeze_panes = 'B4'

# ----------------------------------------------------------------- лист «Расчёт»: одна строка = один вариант, все строки «Отчета»
wr = wb.create_sheet('Расчёт')
put(wr, 'A1', 'Расчёт: одна строка - один вариант, все строки «Отчета» в месяц (формулы от «Допущений»)', fH)
put(wr, 'A2', 'Колонки G-M - только для сетки цена × клик. Ставка за заказ и другие пути считаются в тех же колонках N-AF своими формулами объёма.', fGrey)
RC = ['№', 'Структура', 'Сценарий', 'Снижение цены', 'Доп. клик, ₽', 'Вариант', 'g = (1−снижение)^ε', 'Δ шт от цены', 'Δ оборота от цены по ценам сентября', 'Δ оборота от цены',
      'Товаров в клике (экв. 1 500 ₽/нед)', 'шт от клика', 'Δ оборота от клика', 'Начислено (оборот)', 'Реализовано, шт', 'СС произв.', 'Реклама за заказ', 'Реклама за клик',
      'Реклама всего', 'Сборы без рекламы', 'К выплате', 'Доставка покупателя (доход)', 'Наша доставка', 'Валовая прибыль', 'База АДМ', 'АДМ 30%', 'Налоги 15%', 'Чистая прибыль',
      'Реклама / Начислено', 'Δ К выплате к базе', 'Δ чистой к базе', 'Отдача клика, ₽ оборота на 1 ₽']
head(wr, 4, RC, [5, 14, 10, 9, 12, 26] + [13] * 26)
wr.row_dimensions[4].height = 58
SNAME = {'B': 'июль-сентябрь', 'C': 'сентябрь'}; ENAME = {'B': 'низкий', 'C': 'средний', 'D': 'высокий'}
fmtd = lambda d: f"{-d * 100:.0f}%" if d else '0%'
fmtb = lambda b: f"+{b / 1e6:.2f}".rstrip('0').rstrip('.').replace('.', ',') + ' млн' if b else '0'
GROW, BASE, GROWU = {}, {}, {}
def common(r, sc, base_row):
    a = lambda n: AR(n, sc)
    f = {'S': f"=Q{r}+R{r}", 'T': f"=N{r}*({a('comm')}+{a('logi')}+{a('acq')}+{a('sto')}+{a('oth')})", 'U': f"=N{r}-T{r}-S{r}",
         'V': f"=O{r}*{a('dinc')}", 'W': f"=O{r}*{a('ship')}", 'X': f"=U{r}-P{r}-W{r}+V{r}", 'Y': f"=U{r}+{a('kadm')}*S{r}+{a('badm')}*N{r}",
         'Z': f"=Y{r}*{a('adm')}", 'AA': f"=N{r}*{a('real')}*{a('tax')}", 'AB': f"=X{r}-Z{r}-AA{r}", 'AC': f"=IFERROR(S{r}/N{r},0)",
         'AD': f"=U{r}-$U${base_row}", 'AE': f"=AB{r}-$AB${base_row}"}
    for col, v in f.items():
        put(wr, f'{col}{r}', v, fmt=PCT if col == 'AC' else RUB, bold=col in ('U', 'AB'), fill=fillRes if col == 'AB' else None)
r = 5; i = 0
for mode in ('cap', 'up'):
  if mode == 'up':
    r += 1; put(wr, f'A{r}', 'Верхняя граница клика: сверх пула штуки растут как показы, расход^0,81 (оценка сверху)', fB, fill=fillSub); r += 1
  for sc in 'BC':
    for scen in 'BCD':
        for di, d in enumerate(DLIST):
            for bi, B in enumerate(BLIST):
                i += 1
                if di == 0 and bi == 0 and scen == 'B' and mode == 'cap': BASE[sc] = r
                (GROW if mode == 'cap' else GROWU)[(sc, scen, di, bi)] = r
                e = AR('eps', scen); n0 = AR('N0'); u0 = AR('U0')
                put(wr, f'A{r}', i); put(wr, f'B{r}', SNAME[sc]); put(wr, f'C{r}', ENAME[scen])
                put(wr, f'D{r}', f"={AR(f'd{di}')}", fmt=DPCT); put(wr, f'E{r}', f"={AR(f'b{bi}')}", fmt=RUB)
                put(wr, f'F{r}', f'="цена "&IF(D{r}=0,"0%","-"&TEXT(D{r}*100,"0")&"%")&", клик "&IF(E{r}=0,"0","+"&SUBSTITUTE(TEXT(E{r}/1000000,"0.00"),".",",")&" млн")' + ('&" (верх)"' if mode == 'up' else ''))
                put(wr, f'G{r}', f"=(1-D{r})^{e}", fmt='0.000')
                put(wr, f'H{r}', f"={u0}*{AR('su')}*(G{r}-1)", fmt=X1)
                put(wr, f'I{r}', f"={n0}*{AR('srev')}*(G{r}-1)", fmt=RUB)
                put(wr, f'J{r}', f"={n0}*{AR('srev')}*((1-D{r})*G{r}-1)", fmt=RUB)
                put(wr, f'K{r}', f"=E{r}/{AR('cpi')}", fmt=X1)
                over = (f"{AR('pool')}*{AR('ui', scen)}+(E{r}-{AR('Bp')})*{AR('Robs', scen)}/{AR('pc')}" if mode == 'cap'
                        else f"{AR('pool')}*{AR('ui', scen)}*(E{r}/{AR('Bp')})^{AR('expo')}")
                put(wr, f'L{r}', f"=IF(E{r}<={AR('Bp')},K{r}*{AR('ui', scen)},{over})", fmt=X1)
                put(wr, f'M{r}', f"=L{r}*{AR('pc')}", fmt=RUB)
                put(wr, f'N{r}', f"={n0}+J{r}+M{r}", fmt=RUB, bold=True); put(wr, f'O{r}', f"={u0}+H{r}+L{r}", fmt=X1)
                put(wr, f'P{r}', f"={AR('cc', sc)}*({n0}+{AR('ccrel')}*I{r}+M{r})", fmt=RUB)
                put(wr, f'Q{r}', f"=N{r}*{AR('cpo', sc)}", fmt=RUB); put(wr, f'R{r}', f"={AR('clk', sc)}*{n0}+E{r}", fmt=RUB)
                common(r, sc, BASE[sc])
                put(wr, f'AF{r}', f'=IF(E{r}>0,M{r}/E{r},"")', fmt=X2)
                r += 1
GEND = r - 1
r += 1; put(wr, f'A{r}', 'Ставка за заказ', fB, fill=fillSub); r += 1
CROW = {}
for sc in 'BC':
    for kind, li in [('5', None), ('9', None)] + [('0', k) for k in range(len(LLIST))]:
        n0 = AR('N0'); u0 = AR('U0'); lr = AR(f'L{li}') if li is not None else '0'
        put(wr, f'B{r}', SNAME[sc]); put(wr, f'C{r}', '-')
        lab = {'5': 'ставка 5% (сейчас)', '9': 'ставка 9%'}.get(kind)
        put(wr, f'F{r}', lab if lab else f'="ставка 0%, штуки "&TEXT({lr}*100,"0")&"%"')
        put(wr, f'N{r}', f"={n0}*(1+{lr})", fmt=RUB, bold=True); put(wr, f'O{r}', f"={u0}*(1+{lr})", fmt=X1)
        put(wr, f'P{r}', f"={AR('cc', sc)}*N{r}", fmt=RUB)
        put(wr, f'Q{r}', {'5': f"=N{r}*{AR('cpo', sc)}", '9': f"=N{r}*{AR('cpo', sc)}*{AR('cpo9')}", '0': '=0'}[kind], fmt=RUB)
        put(wr, f'R{r}', f"={AR('clk', sc)}*{n0}", fmt=RUB)
        common(r, sc, BASE[sc]); CROW[(sc, kind, li)] = r; r += 1
r += 1; put(wr, f'A{r}', 'Другие пути', fB, fill=fillSub); r += 1
PROW = {}
n0 = AR('N0'); u0 = AR('U0')
PATHS = [('spring', 'весенние шт (апрель) по цене штуки сентября', f"={n0}+({AR('Usp')}-{u0})*{AR('P0')}", f"={AR('Usp')}", None, None),
         ('mirror', 'зеркала к уровню июня', f"={n0}+({AR('Mjun')}-{AR('Msep')})*{AR('Mp')}", f"={u0}+{AR('Mjun')}-{AR('Msep')}", None, None),
         ('cut', 'срез дорогого клика', f"={n0}", f"={u0}", None, f"-{AR('cut')}")]
for sc in 'BC':
    for key, lab, fN_, fU_, _, dclk in PATHS:
        put(wr, f'B{r}', SNAME[sc]); put(wr, f'C{r}', '-'); put(wr, f'F{r}', lab)
        put(wr, f'N{r}', fN_, fmt=RUB, bold=True); put(wr, f'O{r}', fU_, fmt=X1); put(wr, f'P{r}', f"={AR('cc', sc)}*N{r}", fmt=RUB)
        put(wr, f'Q{r}', f"=N{r}*{AR('cpo', sc)}", fmt=RUB); put(wr, f'R{r}', f"={AR('clk', sc)}*{n0}" + (dclk or ''), fmt=RUB)
        common(r, sc, BASE[sc]); PROW[(sc, key)] = r; r += 1
r += 1; CHK = r
put(wr, f'A{r}', 'Проверка (инвариант): вариант «цена 0%, клик 0» в структуре сентября минус «Отчет» сентября (чистая + К выплате + валовая + реклама), ₽ - должно быть 0', fB)
b0r = GROW[('C', 'C', 0, 0)]
put(wr, f'F{r}', f"=ROUND(AB{b0r}-{FR('Чистая прибыль')},0)+ROUND(U{b0r}-{FR('К выплате')},0)+ROUND(X{b0r}-{FR('Валовая прибыль')},0)+ROUND(S{b0r}-{FR('Реклама (весь расход по начислениям OZON)')},0)", fmt='0', fill=fillKey)
wr.freeze_panes = 'G5'
GR = lambda col, sc, scen, di, bi, mode='cap': f"'Расчёт'!${col}${(GROW if mode == 'cap' else GROWU)[(sc, scen, di, bi)]}"
for rr in range(5, GEND + 1):
    pass
wr.conditional_formatting.add(f'AB5:AB{r}', CellIsRule(operator='greaterThanOrEqual', formula=['0'], fill=fillG))
wr.conditional_formatting.add(f'AB5:AB{r}', CellIsRule(operator='lessThan', formula=['0'], fill=fillR))

# ----------------------------------------------------------------- Python-зеркало ставки за заказ и путей (для проверки и текста)
def cpo_m(sc, kind, L=0.0):
    s = ST[sc]; N = N0 * (1 + L); U = U0 * (1 + L)
    return pl(s, N, U, s['cc'] * N, {'5': s['cpo'] * N, '9': s['cpo'] * N * 1.8, '0': 0}[kind], s['clk'] * N0)
def path_m(sc, key):
    s = ST[sc]
    if key == 'spring': N = N0 + (UNITS[1] - U0) * P0; U = UNITS[1]; dc = 0
    elif key == 'mirror': N = N0 + (MIRU[3] - MIRU[6]) * MIRN[6] / MIRU[6]; U = U0 + MIRU[3] - MIRU[6]; dc = 0
    else: N, U, dc = N0, U0, -250_000
    return pl(s, N, U, s['cc'] * N, s['cpo'] * N, s['clk'] * N0 + dc)

# ----------------------------------------------------------------- лист «Сетка цена×реклама»
wg = wb.create_sheet('Сетка цена×реклама')
put(wg, 'A1', 'Сетка: снижение цены ходовых SKU × дополнительный дешёвый клик, ₽ в месяц (строки «Отчета»)', fH)
note(wg, 2, 'Строки - снижение цены на ходовых SKU (69 SKU, 76% оборота сентября; лист «Допущения» п.4). Колонки - дополнительный дешёвый клик сверх нынешнего (п.5). '
     'Сценарий связывает силу цены и клика: низкий = эластичность -1 и клик ×0,3; средний = -2 и ×0,5; высокий = -3 и ×1. Структура расходов - как в июле-сентябре (опорная) или как в сентябре (единственный месяц с плюсом). '
     'Клик в базе: до бюджета пула (23 товара) - по тесту 1, сверх пула - не больше отдачи нынешнего клика по «Отчету» (1,24 / 2,56 / 2,85 ₽ на 1 ₽). Внизу листа - верхняя граница (расход^0,81), оценка сверху. Цена и клик сложены независимо. '
     'Чистая: зелёный - не меньше 0, красный - меньше 0. «К выплате» не меньше 20 млн - синим. Всё [ГИПОТЕЗА]: расчёт, не факт.', 8, 72)
for c, w in zip('ABCDEFGH', [34, 15, 15, 15, 15, 15, 15, 4]): wg.column_dimensions[c].width = w
METRICS = [('AB', 'Чистая прибыль «Отчета»', RUB), ('U', 'К выплате', RUB), ('N', 'Начислено (оборот)', RUB), ('O', 'Реализовано, шт', SHT), ('S', 'Реклама всего', RUB), ('AF', 'Отдача клика, ₽ оборота на 1 ₽', X2)]
r = 4; GT = {}
for scen in 'CBD':
    for sc in 'BC':
        put(wg, f'A{r}', f"Сценарий {ENAME[scen]} (ε {EPS[scen]:.0f}, клик ×{KS[scen]:.1f}), расходы как в {'июле-сентябре' if sc == 'B' else 'сентябре'}".replace('.', ','), fWhite, fill=fillHead)
        for col in 'BCDEFG': wg[f'{col}{r}'].fill = fillHead
        r += 1
        for col, lab, fmt in METRICS:
            put(wg, f'A{r}', f'{lab}: цена ↓ / клик →', fB, fill=fillSub)
            for bi in range(6): put(wg, f'{CL(bi + 2)}{r}', f"={AR(f'b{bi}')}", fB, '"+"#,##0;-#,##0;"0"', fill=fillSub)
            r += 1; r0 = r
            for di in range(5):
                put(wg, f'A{r}', f"={AR(f'd{di}')}", fN, '"цена -"0%;"цена +"0%;"цена 0%"')
                for bi in range(6): put(wg, f'{CL(bi + 2)}{r}', f"={GR(col, sc, scen, di, bi)}", fmt=fmt, bold=col == 'AB')
                r += 1
            rng = f'B{r0}:G{r - 1}'
            if col == 'AB':
                wg.conditional_formatting.add(rng, CellIsRule(operator='greaterThanOrEqual', formula=['0'], fill=fillG))
                wg.conditional_formatting.add(rng, CellIsRule(operator='lessThan', formula=['0'], fill=fillR))
            if col == 'U':
                wg.conditional_formatting.add(rng, CellIsRule(operator='greaterThanOrEqual', formula=[AR('KVT').replace("'Допущения'!", "'Допущения'!")], fill=fillT))
            GT[(scen, sc, col)] = rng
        r += 1
r += 1; put(wg, f'A{r}', 'ВЕРХНЯЯ ГРАНИЦА клика (оценка сверху): сверх пула штуки растут как показы, расход^0,81. Тест 3 этого не подтверждает: дорогой клик даёт показы, а заказов почти нет', fB, fill=fillBad)
for col in 'BCDEFG': wg[f'{col}{r}'].fill = fillBad
r += 2
for scen in 'CD':
    for sc in 'BC':
        put(wg, f"A{r}", f"Верхняя граница: сценарий {ENAME[scen]}, расходы как в {'июле-сентябре' if sc == 'B' else 'сентябре'}", fWhite, fill=fillHead)
        for col in 'BCDEFG': wg[f'{col}{r}'].fill = fillHead
        r += 1
        for col, lab, fmt in [m for m in METRICS if m[0] in ('AB', 'U', 'AF')]:
            put(wg, f'A{r}', f'{lab} (верх): цена ↓ / клик →', fB, fill=fillSub)
            for bi in range(6): put(wg, f'{CL(bi + 2)}{r}', f"={AR(f'b{bi}')}", fB, '"+"#,##0;-#,##0;"0"', fill=fillSub)
            r += 1; r0 = r
            for di in range(5):
                put(wg, f'A{r}', f"={AR(f'd{di}')}", fN, '"цена -"0%;"цена +"0%;"цена 0%"')
                for bi in range(6): put(wg, f'{CL(bi + 2)}{r}', f"={GR(col, sc, scen, di, bi, 'up')}", fmt=fmt, bold=col == 'AB')
                r += 1
            rng_ = f'B{r0}:G{r - 1}'
            if col == 'AB':
                wg.conditional_formatting.add(rng_, CellIsRule(operator='greaterThanOrEqual', formula=['0'], fill=fillG))
                wg.conditional_formatting.add(rng_, CellIsRule(operator='lessThan', formula=['0'], fill=fillR))
            if col == 'U':
                wg.conditional_formatting.add(rng_, CellIsRule(operator='greaterThanOrEqual', formula=[AR('KVT')], fill=fillT))
        r += 1
wg.freeze_panes = 'B4'

# ----------------------------------------------------------------- лист «Ставка за заказ»
wc = wb.create_sheet('Ставка за заказ')
put(wc, 'A1', 'Ставка за заказ (оплата за заказ на все товары, кампания 28753239): 5% сейчас, 9% или выключить', fH)
note(wc, 2, 'Выключение (0%) убирает расход на оплату за заказ, но и около половины показов магазина. Сколько при этом теряется штук, данными не разделено: в окне 29.08-03.09, когда оплату за заказ выключали, '
     'одновременно на 7,5% подорожала витрина всего магазина, поэтому -33% штук в окне - это цена и реклама вместе (ФЕНИКС, K1). Поэтому ниже несколько вариантов потери и порог, при котором выключение окупается. '
     'Вывод модели: по «К выплате» оплата за заказ нужна (выключение хуже уже при потере около 10% штук), по чистой «Отчета» - нет (выключение лучше, пока теряем меньше 41-63% штук). Какая цель важнее - решает Иван.', 10, 58)
for c, w in zip('ABCDEFGHIJ', [34, 12, 15, 10, 15, 15, 15, 15, 15, 4]): wc.column_dimensions[c].width = w
r = 4; CB = {}
for sc in 'BC':
    put(wc, f'A{r}', f"Расходы как в {'июле-сентябре (опорная)' if sc == 'B' else 'сентябре'}", fWhite, fill=fillHead)
    for col in 'BCDEFGHI': wc[f'{col}{r}'].fill = fillHead
    r += 1
    head(wc, r, ['Вариант', 'Изменение штук', 'Начислено', 'шт', 'Реклама всего', 'К выплате', 'Чистая «Отчета»', 'Δ К выплате к 5%', 'Δ чистой к 5%']); wc.row_dimensions[r].height = 30; r += 1
    b5 = CROW[(sc, '5', None)]; rows_ = [(sc, '5', None), (sc, '9', None)] + [(sc, '0', k) for k in range(len(LLIST))]
    for key in rows_:
        cr = CROW[key]; R_ = lambda col: f"='Расчёт'!{col}{cr}"
        put(wc, f'A{r}', R_('F')); put(wc, f'B{r}', f"={AR(f'L{key[2]}')}" if key[2] is not None else 0, fmt=PCT)
        for col, src, fmt in [('C', 'N', RUB), ('D', 'O', X1), ('E', 'S', RUB), ('F', 'U', RUB), ('G', 'AB', RUB)]: put(wc, f'{col}{r}', R_(src), fmt=fmt, bold=col == 'G')
        put(wc, f'H{r}', f"='Расчёт'!U{cr}-'Расчёт'!U{b5}", fmt=RUB); put(wc, f'I{r}', f"='Расчёт'!AB{cr}-'Расчёт'!AB{b5}", fmt=RUB, bold=True)
        CB[key] = r; r += 1
    wc.conditional_formatting.add(f'G{r - 6}:G{r - 1}', CellIsRule(operator='greaterThanOrEqual', formula=['0'], fill=fillG))
    wc.conditional_formatting.add(f'G{r - 6}:G{r - 1}', CellIsRule(operator='lessThan', formula=['0'], fill=fillR))
    a0, a3 = CB[(sc, '0', 0)], CB[(sc, '0', 3)]
    put(wc, f'A{r}', 'Порог: выключение не хуже 5% по чистой, только если штук теряем не больше', fB, fill=fillRes)
    put(wc, f'B{r}', f"=B{a0}+(G{CB[(sc, '5', None)]}-G{a0})/((G{a3}-G{a0})/(B{a3}-B{a0}))", fmt=PCT, bold=True, fill=fillRes); CB[(sc, 'be_net')] = r; r += 1
    put(wc, f'A{r}', 'Порог: выключение не хуже 5% по «К выплате», только если штук теряем не больше', fB, fill=fillRes)
    put(wc, f'B{r}', f"=B{a0}+(F{CB[(sc, '5', None)]}-F{a0})/((F{a3}-F{a0})/(B{a3}-B{a0}))", fmt=PCT, bold=True, fill=fillRes); CB[(sc, 'be_kv')] = r; r += 2
put(wc, f'A{r}', 'Факты [ДАННЫЕ, разбор v7]', fB, fill=fillSub); r += 1
for t in ['Окно 29.08-03.09 (оплата за заказ выключена) против 22-28.08, в день: итоговые показы 184 458 -> 91 418 (-50%), рекламные 114 993 -> 26 623; органика (итоговые минус рекламные) 69 466 -> 64 796.',
          'Штуки: 9,6 шт/день (15-28.08) -> 6,3 (29.08-03.09, 38 шт за 6 дней) -> 9,3 (04-17.09). Витрина магазина в окне +7,5% (проба ФЕНИКСА) - вклад рекламы и цены не разделить.',
          'Через оплату за заказ прошли 54-73% штук июня-сентября (cpo_orders × orders_daily, атрибуция OZON, пересекается с органикой). Расход 0,75-1,18 млн в месяц.',
          'Ставка 9% против 5% (v5): заказов не добавила, -2,8% на пункт ставки, ДИ -7,6...+2,3 [ДАННЫЕ]. Держим 5%.',
          'Если всё же проверять выключение: чередовать включение и выключение на весь магазин по 2 недели, минимум 2 цикла, только с решением Ивана (изменение рекламы около 0,8-1,0 млн в месяц, HITL).']:
    note(wc, r, t, 9, 30, fN); r += 1

# ----------------------------------------------------------------- лист «Другие пути»
wo = wb.create_sheet('Другие пути')
put(wo, 'A1', 'Другие пути развития: что даёт каждый в месяц к сентябрю (₽), сколько стоит и насколько уверены', fH)
note(wo, 2, 'Числовые оценки - той же моделью «Отчета» (лист «Расчёт»), средняя структура = расходы как в июле-сентябре. «Не оценено» - эффект в данных не измерен, цифру не выдумываем. Пути 1 и 2 не складывать: весенний объём уже содержит зеркала.', 10, 32)
head(wo, 4, ['Путь', 'Что делаем', 'Δ оборота, ₽', 'Δ К выплате, ₽ (июль-сент)', 'Δ чистой, ₽ (июль-сент)', 'Δ чистой, ₽ (сентябрь)', 'Стоимость', 'Срок', 'Уверенность', 'Метка и источник'],
     [30, 46, 14, 15, 15, 15, 22, 14, 12, 70])
wo.row_dimensions[4].height = 44
PRx = lambda sc, key, col: f"'Расчёт'!{col}{PROW[(sc, key)]}"
mid10 = GROW[('B', 'C', 2, 0)]; mid10c = GROW[('C', 'C', 2, 0)]
so = sorted(SCH.items(), key=lambda kv: -kv[1]['u'])
sch_txt = '; '.join(f"{k}: {a['u']} шт, комиссия {a['comm'] * 100:.1f}%, логистика OZON {a['logi'] * 100:.1f}%, {a['p'] / 1000:.0f} тыс. ₽/шт" for k, a in so).replace('.', ',')
OTH = [
    ('1. Вернуть весенний объём', 'Разобрать и снять причины падения заказов (разбор v7: цена +24-38% и снятый клик). Строка = 581 шт (апрель) по цене штуки сентября',
     f"={PRx('B', 'spring', 'N')}-{AR('N0')}", f"={PRx('B', 'spring', 'AD')}", f"={PRx('B', 'spring', 'AE')}", f"={PRx('C', 'spring', 'AE')}",
     'цена и клик (сетка)', '1-3 мес', 'низкая', '[ГИПОТЕЗА] Такого месяца не было: весной штука стоила 42,9-46,8 тыс., осенью 77,5 тыс. При весенней цене и весеннем клике чистая была в минусе (март-июнь -0,01...-1,35 млн, «Отчет»). Это условие, а не план'),
    ('2. Зеркала к уровню июня', 'Дешёвый клик 8-12 ₽ на ходовых зеркалах (GGL-09-2, GGM-24-2, GGM-02, GGM-16; KUVINO GGM-18-2-2 не трогаем - решение 13.10) с контролем',
     f"={PRx('B', 'mirror', 'N')}-{AR('N0')}", f"={PRx('B', 'mirror', 'AD')}", f"={PRx('B', 'mirror', 'AE')}", f"={PRx('C', 'mirror', 'AE')}",
     'клик ~1 500 ₽/нед на товар', '2-4 нед', 'низкая', '[ГИПОТЕЗА] «Отчет», блок 1: зеркала июнь 173 шт -> сентябрь 60. Разбор v7: показы зеркал ушли вместе с кликом (эластичность показов к цене +0,11 после контроля клика). Клик в оценке не вычтен - смотреть сетку'),
    ('3. Срез дорогого клика', 'Две кампании с 0-1 заказом за 4 недели (40015169, 40047060) и KONUX - на минимум; KUVINO не трогаем (решение 13.10)',
     0, f"={PRx('B', 'cut', 'AD')}", f"={PRx('B', 'cut', 'AE')}", f"={PRx('C', 'cut', 'AE')}",
     '0 (экономия)', 'сразу после «да»', 'средняя', '[ГИПОТЕЗА] тест 3: дорогой клик даёт показы, а заказов почти нет. Что заказов не потеряем - не доказано. Изменение рекламы 0,25 млн - HITL'),
    ('4. Перезалив карточек без заказов', f"Перезалив фото и заголовков: в сентябре {NZ['sep_f']} мебели + {NZ['sep_m']} зеркал с показами и без заказов 60 дн (октябрь 01-11: {NZ['oct_f']} + {NZ['oct_m']}). Пилот на 10 карточках с самыми большими показами",
     'не оценено', 'не оценено', 'не оценено', 'не оценено', 'работа контента', '2 нед на пилот', 'низкая',
     f"[ДАННЫЕ] «Отчет», блок 5; sku_views: на товары без заказов 02.08-30.09 пришлось {VZ['share'] * 100:.0f}% показов сентября ({VZ['n']} SKU). Эффект перезалива не измерен - замер по заходам в карточку за 2 недели против похожих карточек".replace('.', ',')),
    ('5. Ассортимент', 'Выяснить, почему товаров с заметными показами стало меньше (архив, остатки, блокировки), и вернуть продававшиеся',
     'не оценено', 'не оценено', 'не оценено', 'не оценено', 'работа команды', '2 нед', 'низкая',
     '[ДАННЫЕ] разбор v7: SKU с показами от 100 за 30 дней 1 002 (июнь) -> 629 (сентябрь) -> 488 (октябрь 01-11); SKU с заказом 199 -> 122. Сжатие с августа, падение заказов - с июля: фактор вторичный. Причина не проверена'),
    ('6. Распродажа 11.11', 'Участие ходовых SKU в распродаже неделю; условия OZON (глубина скидки) узнать заранее. Цена больше 5% - решение Ивана',
     f"=('Расчёт'!N{mid10}-{AR('N0')})*{AR('wk11')}", f"='Расчёт'!AD{mid10}*{AR('wk11')}", f"='Расчёт'!AE{mid10}*{AR('wk11')}", f"='Расчёт'!AE{mid10c}*{AR('wk11')}",
     'недополученная цена', '06-12.11', 'низкая', '[ГИПОТЕЗА] Строка = сетка «цена -10%, клик 0», средний сценарий, на одну неделю. Участия в акциях до 23.09 в данных нет (promo_daily - 1 день, acts_daily и eb_daily - с 23.09). Для недели распродажи часть прироста - перенос покупок с соседних недель (поэтому -4 «относительная» тут ближе, но месяц в целом получит меньше)'),
    ('7. Комиссия: схема FBS / rFBS / FBO (в колонках - цена 1 пункта комиссии, а не эффект пути)', 'Сравнить комиссию по схемам и не переводить товары ради комиссии без данных. Цифры в колонках - сколько стоит 1 пункт комиссии в месяц, эффект самого пути не оценён',
     0, f"={AR('N0')}*0.01", f"={AR('N0')}*0.01*(1-{AR('adm', 'B')})", f"={AR('N0')}*0.01*(1-{AR('adm', 'C')})",
     '-', '-', 'нет данных', f"[ДАННЫЕ orders_daily, доставленные июль-сентябрь, доля от цены продажи] {sch_txt}. Наша доставка rFBS по схемам не разнесена. Разница FBS и rFBS 0,8 п. - в пределах разницы ассортимента; FBO - 6 шт, данных нет. Строка = цена 1 пункта комиссии в месяц при обороте сентября [РАСЧЁТ]. Комиссия / Начислено: весна 46,1%, сентябрь 49,5%, октябрь 01-11 50,8% («Отчет»)"),
    ('8. Ставка за заказ', 'Держать 5%, 9% не ставить. Если цель - чистая: проверить выключение чередованием 2 недели вкл / 2 выкл на весь магазин (минимум 2 цикла), решение Ивана', 0,
     f"='Ставка за заказ'!H{CB[('B', '0', 2)]}", f"='Ставка за заказ'!I{CB[('B', '0', 2)]}", f"='Ставка за заказ'!I{CB[('C', '0', 2)]}", '0 (экономия 0,8-1,0 млн)', 'после замера цены', 'низкая',
     '[ГИПОТЕЗА] Строка = выключение при потере штук как в окне 29.08-03.09 (-33%, вместе с ростом витрины). По «К выплате» минус, по чистой плюс: лист «Ставка за заказ». 9% против 5% заказов не добавила (v5) [ДАННЫЕ]'),
    ('9. Ежедневные снимки цен, остатков, акций, сроков', 'Сохранять каждый день нашу цену, остатки, участие в акциях и обещанный срок доставки; выяснить, кто поднял цены 29.06-19.07',
     'не оценено', 'не оценено', 'не оценено', 'не оценено', 'время аналитики', '1 нед', 'высокая (что нужно)', '[ДАННЫЕ] prices_daily с 19.09 и в первые дни скачет; остатков по дням нет; акции с 23.09. Без этого эффект цены не отделить от рекламы и акций - любые решения по сетке остаются гипотезой'),
    ('10. Что дал переход на FBS в октябре', 'Проверить, связан ли октябрьский переход на FBS со сроком доставки и конверсией',
     'не оценено', 'не оценено', 'не оценено', 'не оценено', 'время аналитики', '1 нед', 'низкая',
     f"[ДАННЫЕ orders_daily] октябрь 01-12: FBS {OCTS['FBS']} шт из {sum(OCTS.values())} (в июле-сентябре 43-48%); заказ -> начисление продажи медиана 5 дней против 9-11 (n=44, прокси срока). Связь с заказами не проверена"),
]
r = 5
for path, what, dn, dkv, dnet, dnetc, cost, when, conf, src in OTH:
    put(wo, f'A{r}', path, fB, wrap=True); put(wo, f'B{r}', what, wrap=True)
    for col, v in zip('CDEF', (dn, dkv, dnet, dnetc)): put(wo, f'{col}{r}', v, fmt=RUB, bold=col == 'E')
    put(wo, f'G{r}', cost, wrap=True); put(wo, f'H{r}', when, wrap=True); put(wo, f'I{r}', conf, wrap=True); put(wo, f'J{r}', src, fGrey, wrap=True)
    wo.row_dimensions[r].height = 92; r += 1
wo.conditional_formatting.add(f'E5:F{r - 1}', CellIsRule(operator='greaterThan', formula=['0'], fill=fillG))
wo.conditional_formatting.add(f'E5:F{r - 1}', CellIsRule(operator='lessThan', formula=['0'], fill=fillR))
OTH_END = r - 1
note(wo, r + 1, 'Дорогой клик (большие ставки и бюджеты на продающих товарах) не наращивать: тест 3, KONUX - перетягивание заказов у соседей по карточке. Это не то же, что дешёвый клик сетки.', 10, 30)

# ----------------------------------------------------------------- лист «План»
wp = wb.create_sheet('План')
put(wp, 'A1', 'План 16.10-30.11: что делать по неделям, где стоп и кто решает', fH)
note(wp, 2, 'HITL (Protocol 6): цена больше 5% и реклама больше 100 тыс. ₽ в месяц - только после «да» Ивана. KUVINO не трогаем (решение 13.10). Ставка за заказ - 5%. Даты - если Иван ответит до 16.10.', 7, 30)
head(wp, 4, ['Неделя', 'Что делаем', 'Ориентир (средний сценарий, июль-сентябрь)', 'Стоп-правило', 'Исполнитель', 'Решение / HITL'], [14, 58, 30, 52, 18, 26])
ORI = lambda sc, scen, di, bi, col: f"'Расчёт'!{col}{GROW[(sc, scen, di, bi)]}"
PLAN = [
    ('16-22.10', 'Иван решает четыре вопроса: (1) цена -5% на ходовых SKU как замер эластичности; (2) волна 2 дешёвого клика: 20 товаров по условиям теста 1 + 10 товаров «показы без заказов», с контролем; (3) срез двух кампаний с 0-1 заказом и KONUX; (4) ежедневные снимки цен, остатков, акций, сроков',
     f'="клик волны 2 на месяц: "&{M(AR("cpi") + "*30")}&" млн"', 'без «да» Ивана по пункту - пункт не начинаем', 'аналитика -> Иван', 'Иван: п.2 и п.3 (реклама >100 тыс. в месяц), п.1 - в пределах 5%'),
    ('16-22.10', 'Запустить ежедневные снимки: наша цена, остатки, участие в акциях, обещанный срок доставки. Выяснить у команды, кто поднял цены 29.06-19.07 и почему', '-', 'снимок не собрался 2 дня подряд - чинить до любых решений по цене', 'аналитика', '-'),
    ('23-29.10', 'Цена -5% на ходовых: чередовать по неделям (неделя -5%, неделя обычная) на одних и тех же моделях; TRUBIS Oval 160 и 200 - отдельной группой (половина штук)',
     f'="К выплате "&{M(ORI("B", "C", 1, 0, "AD"))}&" млн, чистая "&{M(ORI("B", "C", 1, 0, "AE"))}&" млн в месяц к сейчас"', 'рентабельность SKU ниже 10% (решение 13.10) - вернуть цену. Итог - только после всего окна (4-8 недель), без откатов по шуму', 'команда OZON', 'в пределах 5% - без HITL'),
    ('23-29.10', 'Волна 2 клика, 5-й день: проверить расход и CTR', '-', 'расход больше 1 500 ₽/нед на товар или CTR вдвое ниже теста 1 - поправить ставку', 'команда OZON', '-'),
    ('30.10-05.11', 'Итог волны 2 (14 дней) против контроля', f'="ориентир +0,25 млн клика: К выплате "&{M(ORI("B", "C", 0, 1, "AD"))}&" млн, чистая "&{M(ORI("B", "C", 0, 1, "AE"))}&" млн в месяц"',
     'ниже 2,4 ₽ оборота на 1 ₽ клика или заказы карточек не выше контроля - выключить. Продление выше 14 дней - только если отдача не ниже ×0,5 теста 1', 'аналитика', 'Иван: рост до +0,25 млн в месяц (HITL)'),
    ('30.10-05.11', 'Подготовить 11.11: условия OZON (глубина скидки), список ходовых, решение по глубине', f'="неделя -10%: чистая "&{M("\'Другие пути\'!E10")}&" млн"', 'скидка больше 10% - без участия, пока нет решения Ивана', 'команда OZON', 'Иван (цена >5%)'),
    ('06-12.11', 'Распродажа 11.11 (если «да»): мерить шт/день против двух соседних недель и против прошлого месяца', '-', 'чистая недели «Отчета» ниже минус 0,1 млн - больше в акции такой глубины не идём', 'аналитика', '-'),
    ('13-19.11', 'Октябрь в «Отчете» закрыт: факт против сетки. Какая эластичность ближе к факту замера -5% - такой сценарий и берём дальше. Решение по шагу -10%',
     f'="-10%: К выплате "&{M(ORI("B", "C", 2, 0, "AD"))}&" млн, чистая "&{M(ORI("B", "C", 2, 0, "AE"))}&" млн в месяц"', 'замер дал эластичность слабее порога (лист «Итог») - цену не снижаем дальше, возвращаем', 'аналитика -> Иван', 'Иван (цена >5%)'),
    ('13-19.11', 'Иван выбирает главную цель: «К выплате» или чистая. Если чистая - чередование оплаты за заказ 2 недели вкл / 2 выкл на весь магазин после окна замера цены (не на тех же неделях)',
     f'="выключение при -33% штук: К выплате "&{M("'Ставка за заказ'!H" + str(CB[('B', '0', 2)]))}&" млн, чистая "&{M("'Ставка за заказ'!I" + str(CB[('B', '0', 2)]))}&" млн"', 'штук в неделю выключения меньше, чем в неделю включения, больше чем на порог листа «Ставка за заказ» - включить обратно', 'аналитика -> Иван', 'Иван (реклама >100 тыс.)'),
    ('20-26.11', 'Клик сверх +0,25 млн в месяц - только если волны 2 и 3 покажут отдачу выше нынешних кампаний (2,56 ₽ на 1 ₽); иначе «К выплате» почти не растёт, а чистая падает', f'="+0,5 млн (база): К выплате "&{M(ORI("B", "C", 0, 2, "AD"))}&" млн, чистая "&{M(ORI("B", "C", 0, 2, "AE"))}&" млн"',
     'чистая «Отчета» за скользящие 4 недели ниже 0 или шт на 1000 показов 2 недели подряд ниже средней 4 прошлых недель больше чем на 20% - урезать клик', 'команда OZON', 'Иван (реклама >100 тыс.)'),
    ('27-30.11', 'Итог ноября: факт против сетки, обновить допущения (эластичность, отдача клика), ретро в реестр ошибок', '-', '-', 'аналитика', '-'),
    ('каждый пн', 'Факт против плана по «Отчету»: оборот, К выплате, реклама / Начислено, чистая за скользящие 4 недели', '-', 'чистая за 4 недели ниже 0 - новых денег в клик не добавляем', 'аналитика -> Иван', '-'),
]
r = 5
for wk, what, ori, stop, who, hitl in PLAN:
    put(wp, f'A{r}', wk, fB, wrap=True); put(wp, f'B{r}', what, wrap=True); put(wp, f'C{r}', ori, wrap=True); put(wp, f'D{r}', stop, wrap=True)
    put(wp, f'E{r}', who, wrap=True); put(wp, f'F{r}', hitl, wrap=True); wp.row_dimensions[r].height = 76; r += 1

CB_KV_B, CB_NET_B, CB_NET_C, CB9_B = CB[('B', 'be_kv')], CB[('B', 'be_net')], CB[('C', 'be_net')], CB[('B', '9', None)]
# ----------------------------------------------------------------- лист «Итог»
wi = wb.create_sheet('Итог', 0)
wi.column_dimensions['A'].width = 44
for col in 'BCDEFG': wi.column_dimensions[col].width = 15
put(wi, 'A1', 'OZON: на сколько снизить цену, на сколько поднять рекламу и какие ещё пути', fH)
note(wi, 2, f'Сборка {DAY:%d.%m.%Y}. Запрос Ивана 14.10: «все варианты расписывай в файле ексель на сколько снизить цену, на сколько поднять рекламу, может ещё какие-то пути развития рекомендуешь». '
     'Цели: оборот («Начислено»), «К выплате» 20 млн в месяц, и видеть чистую «Отчета». Всё ниже - расчёт моделью «Отчета» от сентября [ГИПОТЕЗА], а не факт. Основной ответ - средний сценарий (эластичность -2, клик ×0,5).', 7, 58)
def sub(sc, scen):
    a, b = GROW[(sc, scen, 0, 0)], GROW[(sc, scen, 4, 5)]; return a, b
def rng(col, sc, scen): a, b = sub(sc, scen); return f"'Расчёт'!${col}${a}:${col}${b}"
def rngu(col, sc, scen): a, b = GROWU[(sc, scen, 0, 0)], GROWU[(sc, scen, 4, 5)]; return f"'Расчёт'!${col}${a}:${col}${b}"
ALLU = ','.join(rng('U', sc, scen) for sc in 'BC' for scen in 'BCD')
def best(col, sc, scen): return f"INDEX({rng('F', sc, scen)},MATCH(MAX({rng(col, sc, scen)}),{rng(col, sc, scen)},0))"
def at(col, valcol, sc, scen): return f"INDEX({rng(valcol, sc, scen)},MATCH(MAX({rng(col, sc, scen)}),{rng(col, sc, scen)},0))"
b0 = lambda sc, col: f"'Расчёт'!{col}{BASE[sc]}"
r = 4; put(wi, f'A{r}', 'Ответ (средний сценарий; цифры - формулы)', fB, fill=fillSub)
for c in 'BCDEFG': wi[f'{c}{r}'].fill = fillSub
r += 1
# пороги эластичности (таблица ниже) - строки заранее
BE0 = 39
STEP_P1 = 27; STEP_C1 = 31          # строки таблицы шагов (проверяется ниже)
ans = [
    f"=\"1. Максимум чистой. Расходы как в июле-сентябре: \"&{best('AB', 'B', 'C')}&\", чистая \"&{M(at('AB', 'AB', 'B', 'C'))}&\" млн в месяц; в плюс выводит \"&COUNTIF({rng('AB', 'B', 'C')},\">=0\")&\" из 30 вариантов. "
    f"Расходы как в сентябре: \"&{best('AB', 'C', 'C')}&\", чистая \"&{M(at('AB', 'AB', 'C', 'C'))}&\" млн; в плюсе \"&COUNTIF({rng('AB', 'C', 'C')},\">=0\")&\" из 30 (только малые шаги: цена до -15%, клик до +0,5 млн). В среднем сценарии снижение цены и дешёвый клик чистую не растят ни в одной структуре.\"",
    f"=\"2. Максимум «К выплате» (база: сверх пула 23 товаров клик отдаёт не больше нынешних кампаний по «Отчету»): \"&{best('U', 'B', 'C')}&\" - \"&{M(at('U', 'U', 'B', 'C'))}&\" млн (сейчас \"&{M(b0('B', 'U'))}&\"), чистая \"&{M(at('U', 'AB', 'B', 'C'))}&\" млн. В структуре сентября - \"&{M(at('U', 'U', 'C', 'C'))}&\" млн и чистая \"&{M(at('U', 'AB', 'C', 'C'))}&\" млн. "
    f"20 млн «К выплате» в базе нет ни в одном сценарии: вариантов с 20 млн \"&SUMPRODUCT(--(({rng('U', 'B', 'D')})>={AR('KVT')}))+SUMPRODUCT(--(({rng('U', 'C', 'D')})>={AR('KVT')}))+SUMPRODUCT(--(({rng('U', 'B', 'C')})>={AR('KVT')}))+SUMPRODUCT(--(({rng('U', 'C', 'C')})>={AR('KVT')}))&\", максимум во всей сетке \"&{M(f"MAX({ALLU})")}&\" млн (высокий сценарий, цена -20%, клик +2 млн). "
    f"Только верхняя граница клика (расход^0,81, ОЦЕНКА СВЕРХУ: штуки растут как показы, тест 3 этого не подтверждает) даёт 20 млн - в \"&(COUNTIF({rngu('U', 'B', 'D')},\">=\"&{AR('KVT')})+COUNTIF({rngu('U', 'C', 'D')},\">=\"&{AR('KVT')}))&\" вариантах высокого сценария, все с чистой в минусе.\"",
    f"=\"3. Цена. Снижение на ходовых (69 SKU, 76% оборота) растит оборот и «К выплате», но чистую уменьшает: -5% в среднем сценарии даёт \"&{M(GR('AD', 'B', 'C', 1, 0))}&\" млн «К выплате» и \"&{M(GR('AE', 'B', 'C', 1, 0))}&\" млн чистой в месяц. "
    f"Чистая не падает, только если эластичность всего набора сильнее \"&{X_(f'B{BE0}')}&\" (при -5%) и \"&{X_(f'B{BE0 + 1}')}&\" (при -10%). В структуре сентября порог \"&{X_(f'C{BE0}')}&\" и \"&{X_(f'C{BE0 + 1}')}&\". Разбор мерил -4, но это относительная недельная оценка (перетекание между нашими товарами и перенос покупки на дешёвую неделю); для всего набора сразу модуль, скорее всего, меньше. Поэтому -5% - как замер, а не как рычаг прибыли.\"",
    f"=\"4. Реклама. Первые +0,25 млн дешёвого клика: «К выплате» \"&{M(GR('AD', 'B', 'C', 0, 1))}&\" млн, чистая \"&{M(GR('AE', 'B', 'C', 0, 1))}&\" млн в месяц. Дальше клик выходит за пул (\"&{M(AR('Bp'))}&\" млн, 23 товара) и в базе отдаёт как нынешние кампании (\"&{X2_(AR('Robs', 'C'))}&\" ₽ на 1 ₽): «К выплате» почти не растёт, чистая падает. +2 млн: «К выплате» \"&{M(GR('AD', 'B', 'C', 0, 5))}&\", чистая \"&{M(GR('AE', 'B', 'C', 0, 5))}&\" млн. "
    f"Верхняя граница (расход^0,81, оценка сверху): +2 млн дали бы «К выплате» \"&{M(GR('AD', 'B', 'C', 0, 5, 'up'))}&\" и чистую \"&{M(GR('AE', 'B', 'C', 0, 5, 'up'))}&\" млн при отдаче \"&{X_(GR('AF', 'B', 'C', 0, 5, 'up'))}&\" ₽ на 1 ₽ - выше любых нынешних кампаний (до \"&{X2_(AR('Robs', 'D'))}&\").\"",
    f"=\"5. Граница «показы против цены». Показы покупаются рекламой (клик ×4,4, оплата за заказ - около половины показов), от цены они почти не зависят (-0,5% на +1% цены). Заказы от цены зависят сильно. Витрина ходовых в сентябре дороже июня на \"&{P1(AR('bj'))}&\"% (медиана), но к своей летней «обычной» - только на \"&{P1(AR('bw'))}&\"%. "
    f"Дороже обычной больше чем на 10% - \"&{P1(AR('bwo'))}&\"% их оборота. Заказы на SKU обваливаются (в 4-5 раз) выше +20% к обычной, а таких сейчас мало, поэтому дешёвые показы покупать смысл есть. Но «обычная» уже летняя, поднятая: вернуть весенние заказы клик не может, это вопрос цены (п.3). "
    f"По «К выплате на 1 ₽ потерянной чистой» первый шаг клика (\"&{X_(f'D{STEP_C1}')}&\") выгоднее первого шага цены (\"&{X_(f'D{STEP_P1}')}&\"); клик хуже цены с шага \"&IFERROR(INDEX(A{STEP_C1}:A{STEP_C1 + 4},MATCH(TRUE,INDEX(D{STEP_C1}:D{STEP_C1 + 4}<D{STEP_P1},0),0)),\"- ни на одном шаге\")&\" (таблица шагов ниже).\"",
    f"=\"6. Ставка за заказ - выбор цели. Выключение (0%) уменьшает «К выплате», если штук теряем больше \"&{P1("-'Ставка за заказ'!B" + str(CB_KV_B))}&\"% (июль-сентябрь), но по чистой окупается, пока потеря меньше \"&{P1("-'Ставка за заказ'!B" + str(CB_NET_B))}&\"% (в структуре сентября \"&{P1("-'Ставка за заказ'!B" + str(CB_NET_C))}&\"%). "
    f"В окне 29.08-03.09 штук было на 33% меньше, но там же подорожала витрина - вклад рекламы неизвестен. Для цели «К выплате» держать 5%; для цели «чистая» - проверить выключение чередованием 2 недели вкл / 2 выкл (решение Ивана). 9% не ставить: \"&{M("'Ставка за заказ'!I" + str(CB9_B))}&\" млн чистой в месяц.\"",
    "7. Рекомендация. Цена: -5% на ходовых, чередованием недель, 4-8 недель - чтобы узнать эластичность магазина; -10% и глубже - только решением Ивана и только если замер покажет эластичность сильнее порога (п.3). Реклама: волна 2 дешёвого клика (30 товаров, с контролем), расширять до +0,25 млн в месяц только при успехе; выше - только если волна покажет отдачу выше нынешних кампаний (п.4). Ставка за заказ - 5% до решения Ивана о цели (п.6). Если важнее «К выплате», чем чистая, - самый дешёвый «К выплате» дают первые шаги клика (таблица шагов).",
    "8. 20 млн «К выплате» ценой и рекламой не получить без минуса по чистой. Нужен возврат весеннего объёма (лист «Другие пути», путь 1), а он зависит от разбора причин падения заказов и данных о ценах, остатках и акциях по дням (путь 9).",
    "9. Ограничение: цена и клик сложены независимо, а клик «на товарах без рекламы или со снятой рекламой» пересекается с 69 ходовыми SKU (летом клик сняли с ходовых карточек). Совместный эффект на одних товарах может быть и больше, и меньше суммы - тесты цены и клика вести на разных карточках.",
]
for t, h in zip(ans, [58, 100, 86, 100, 112, 86, 86, 44, 44]):
    put(wi, f'A{r}', t, wrap=True); wi.merge_cells(f'A{r}:G{r}'); wi.row_dimensions[r].height = h; r += 1
r += 1
put(wi, f'A{r}', 'Ключевые варианты, ₽ в месяц (средний сценарий)', fB, fill=fillSub)
for c in 'BCDEFG': wi[f'{c}{r}'].fill = fillSub
r += 1
head(wi, r, ['Вариант', 'Начислено', 'К выплате (июль-сент)', 'Чистая (июль-сент)', 'К выплате (сентябрь)', 'Чистая (сентябрь)', 'Реклама / Начислено (июль-сент)']); wi.row_dimensions[r].height = 44; r += 1
KEYR = r
for di, bi in [(0, 0), (1, 0), (0, 1), (1, 1), (2, 2), (2, 3), (4, 5)]:
    put(wi, f'A{r}', f"={GR('F', 'B', 'C', di, bi)}")
    for col, (src, sc) in zip('BCDEFG', [('N', 'B'), ('U', 'B'), ('AB', 'B'), ('U', 'C'), ('AB', 'C'), ('AC', 'B')]):
        put(wi, f'{col}{r}', f"={GR(src, sc, 'C', di, bi)}", fmt=PCT if src == 'AC' else RUB, bold=src == 'AB')
    r += 1
for col in 'DF':
    wi.conditional_formatting.add(f'{col}{KEYR}:{col}{r - 1}', CellIsRule(operator='greaterThanOrEqual', formula=['0'], fill=fillG))
    wi.conditional_formatting.add(f'{col}{KEYR}:{col}{r - 1}', CellIsRule(operator='lessThan', formula=['0'], fill=fillR))
r += 1
put(wi, f'A{r}', 'Что даёт каждый шаг (средний сценарий, июль-сентябрь)', fB, fill=fillSub)
for c in 'BCDEFG': wi[f'{c}{r}'].fill = fillSub
r += 1
head(wi, r, ['Шаг', 'Δ «К выплате»', 'Δ чистой', '«К выплате» на 1 ₽ потерянной чистой']); wi.row_dimensions[r].height = 44; r += 1
STEPR = {}
for kind, idx in [('p', i) for i in range(1, 5)] + [('c', i) for i in range(1, 6)]:
    if kind == 'p': a_, b_ = (idx - 1, 0), (idx, 0); lab = f'="цена "&IF({AR(f"d{idx - 1}")}=0,"0%","-"&TEXT({AR(f"d{idx - 1}")}*100,"0")&"%")&" -> -"&TEXT({AR(f"d{idx}")}*100,"0")&"%, клик 0"'
    else: a_, b_ = (0, idx - 1), (0, idx); lab = f'="клик +"&SUBSTITUTE(TEXT({AR(f"b{idx - 1}")}/1000000,"0.00"),".",",")&" -> +"&SUBSTITUTE(TEXT({AR(f"b{idx}")}/1000000,"0.00"),".",",")&" млн, цена 0%"'
    put(wi, f'A{r}', lab)
    put(wi, f'B{r}', f"={GR('U', 'B', 'C', *b_)}-{GR('U', 'B', 'C', *a_)}", fmt=RUB)
    put(wi, f'C{r}', f"={GR('AB', 'B', 'C', *b_)}-{GR('AB', 'B', 'C', *a_)}", fmt=RUB)
    put(wi, f'D{r}', f'=IF(C{r}<0,-B{r}/C{r},"чистая не падает")', fmt=X2, bold=True)
    STEPR[(kind, idx)] = r; r += 1
assert STEPR[('p', 1)] == STEP_P1 and STEPR[('c', 1)] == STEP_C1, STEPR
r += 1
assert r <= BE0 - 1, ('сдвинуть BE0', r)
r = BE0 - 1
head(wi, r, ['Порог эластичности: цена не уменьшает чистую, если ε сильнее (меньше)', 'июль-сентябрь', 'сентябрь']); wi.row_dimensions[r].height = 44; r += 1
for di in range(1, 5):
    put(wi, f'A{r}', f'="цена -"&TEXT({AR(f"d{di}")}*100,"0")&"%"')
    for col, sc in zip('BC', 'BC'):
        lo, hi = GR('G', sc, 'B', di, 0), GR('G', sc, 'D', di, 0); nlo, nhi = GR('AB', sc, 'B', di, 0), GR('AB', sc, 'D', di, 0)
        gs = f"({lo}+({b0(sc, 'AB')}-{nlo})*({hi}-{lo})/({nhi}-{nlo}))"
        put(wi, f'{col}{r}', f"=IFERROR(LN({gs})/LN(1-{AR(f'd{di}')}),\"нет\")", fmt=X1, bold=True)
    r += 1
note(wi, r, 'Как считается: при одном шаге цены чистая линейна по g = (1 − снижение)^ε; порог - g, при котором чистая равна сегодняшней, по двум точкам (ε -1 и -3).', 7, 30)
r += 2
note(wi, r, 'Листы: «Сетка цена×реклама» - все 30 вариантов для каждой эластичности и структуры; «Ставка за заказ»; «Другие пути»; «Допущения» - все входы с источником; «План» - по неделям со стоп-правилами; «Расчёт» - строки «Отчета» по каждому варианту и проверка «0% и 0 клика = сентябрь до рубля».', 7, 44)
r += 1
put(wi, f'A{r}', 'Проверка модели (должно быть 0)', fGrey); put(wi, f'B{r}', f"='Расчёт'!F{CHK}", fmt='0', fill=fillKey)

for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.font and c.font.name != F: c.font = Font(name=F, size=c.font.size or 10, bold=c.font.bold, color=c.font.color)
            if isinstance(c.value, str): assert '\u2014' not in c.value, ('длинное тире', ws.title, c.coordinate)
for ws in wb.worksheets:
    ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True
wb._sheets = [wb[n] for n in ['Итог', 'Сетка цена×реклама', 'Ставка за заказ', 'Другие пути', 'Допущения', 'План', 'Расчёт', 'Факт Отчет']]
wb.active = 0
fn = os.path.join(OUT, f'OZON-varianty-cena-reklama-{DAY:%Y-%m-%d}.xlsx')
wb.save(fn); print(fn)

# ----------------------------------------------------------------- проверки утверждений «Итога» на Python-зеркале (текст не разойдётся с цифрами)
for sc in 'BC':
    sub_ = {k: v for k, v in GRID.items() if k[0] == sc and k[1] == 'C'}
    assert max(sub_, key=lambda k: sub_[k]['net'])[2:] == (0, 0), ('максимум чистой не «0% и 0»', sc)
    assert max(sub_, key=lambda k: sub_[k]['kv'])[2:] == (0.2, 2_000_000), sc
    assert all(v['kv'] < 20e6 for v in sub_.values()), 'в среднем сценарии есть 20 млн'
assert all(v['kv'] < 20e6 for v in GRID.values()), 'в базе есть 20 млн'
kmax = max(GRID, key=lambda k: GRID[k]['kv']); assert kmax[1:] == ('D', 0.2, 2_000_000), kmax
assert sum(1 for k, v in GRID.items() if k[:2] == ('B', 'C') and v['net'] >= 0) == 0
okc = [k[2:] for k, v in GRID.items() if k[:2] == ('C', 'C') and v['net'] >= 0]
assert all(d <= 0.15 and B <= 500_000 for d, B in okc), okc          # «цена до -15%, клик до +0,5 млн»
hi20 = [k for k, v in GRIDU.items() if v['kv'] >= 20e6]
assert hi20 and all(k[1] == 'D' and GRIDU[k]['net'] < 0 for k in hi20), hi20
for sc in 'BC':                                     # цена и клик чистую не растят (средний сценарий)
    assert all(GRID[(sc, 'C', d, B)]['net'] <= GRID[(sc, 'C', 0, 0)]['net'] + 1 for d in DLIST for B in BLIST)
assert BORDER['wmed'] - 1 < 0.10 < BORDER['med'] - 1                # «к обычной - меньше 10%, к июню - больше»
k1 = GRID[('B', 'C', 0, 250_000)]; p5 = GRID[('B', 'C', 0.05, 0)]; base = GRID[('B', 'C', 0, 0)]
ratio_c = (k1['kv'] - base['kv']) / (base['net'] - k1['net']); ratio_p = (p5['kv'] - base['kv']) / (base['net'] - p5['net'])
assert ratio_c > ratio_p, ('первый шаг клика дешевле цены по «К выплате»', ratio_c, ratio_p)
assert cpo_m('B', '9')['net'] < cpo_m('B', '5')['net']
print('шаги: клик', round(ratio_c, 2), 'цена', round(ratio_p, 2))

# ----------------------------------------------------------------- пересчёт формул и сверка книги с зеркалом
RECALC = os.environ.get('RECALC')
if RECALC:
    out = subprocess.run([sys.executable, RECALC, fn, '60'], capture_output=True, text=True).stdout
    res = json.loads(out[out.index('{'):]); assert res.get('total_errors') == 0, res
    wv = load_workbook(fn, data_only=True); ws_ = wv['Расчёт']
    assert ws_[f'F{CHK}'].value == 0, ('инвариант сентября', ws_[f'F{CHK}'].value)
    bad = 0
    for (mode, (sc, scen, d, B)), v in [(('cap', k), v) for k, v in GRID.items()] + [(('up', k), v) for k, v in GRIDU.items()]:
        rr = (GROW if mode == 'cap' else GROWU)[(sc, scen, DLIST.index(d), BLIST.index(B))]
        for col, k in [('N', 'N'), ('U', 'kv'), ('AB', 'net'), ('S', 'ad')]:
            if abs(ws_[f'{col}{rr}'].value - v[k]) > 1: bad += 1; print('расхождение', sc, scen, d, B, col, ws_[f'{col}{rr}'].value, v[k])
    for (sc, kind, li), rr in CROW.items():
        v = cpo_m(sc, kind, LLIST[li] if li is not None else 0)
        if abs(ws_[f'AB{rr}'].value - v['net']) > 1 or abs(ws_[f'U{rr}'].value - v['kv']) > 1: bad += 1; print('ставка', sc, kind, li)
    for (sc, key), rr in PROW.items():
        v = path_m(sc, key)
        if abs(ws_[f'AB{rr}'].value - v['net']) > 1: bad += 1; print('путь', sc, key)
    assert bad == 0, bad
    for ws in wv.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith('#'): raise SystemExit(f'ошибка в {ws.title}!{c.coordinate}: {c.value}')
    print('пересчёт:', res['total_formulas'], 'формул, 0 ошибок; сетка, ставка и пути совпали с Python до рубля; сентябрь = «Отчет»')
    wi_ = wv['Итог']
    for rr in range(5, 12): print(wi_[f'A{rr}'].value)
    for rr in range(BE0, BE0 + 4): print('порог', wi_[f'A{rr}'].value, wi_[f'B{rr}'].value, wi_[f'C{rr}'].value)
