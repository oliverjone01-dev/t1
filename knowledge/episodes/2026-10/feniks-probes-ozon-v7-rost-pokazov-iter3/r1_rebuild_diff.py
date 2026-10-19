# ФЕНИКС iter3, проба R1: пересборка по брифу против поданного файла (значения + заливки), без вывода прибыли по артикулам
# python3 -I -B r1_rebuild_diff.py <поданный.xlsx> <пересобранный.xlsx>
import sys
from openpyxl import load_workbook
def cells(p):
    wb = load_workbook(p); val = {}; fil = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None: val[(ws.title, c.coordinate)] = c.value
                fg = c.fill.fgColor.rgb if c.fill and c.fill.fill_type else None
                if fg: fil[(ws.title, c.coordinate)] = fg
    return wb.sheetnames, val, fil
na, a, fa = cells(sys.argv[1]); nb, b, fb = cells(sys.argv[2])
print('листы совпадают:', na == nb, na)
dv = [k for k in sorted(set(a) | set(b)) if a.get(k) != b.get(k)]
df = [k for k in sorted(set(fa) | set(fb)) if fa.get(k) != fb.get(k)]
print('различающихся значений:', len(dv), ' различающихся заливок:', len(df))
for k in dv[:10]: print('  V', k, repr(a.get(k))[:80], '|', repr(b.get(k))[:80])
for k in df[:10]: print('  F', k, fa.get(k), '|', fb.get(k))
