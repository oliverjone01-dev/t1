#!/usr/bin/env python3
# Связка SKU -> артикул (offer) из УПД-отчётов OZON «Отчёт о реализации».
#
# Зачем: у части SKU в дашборде артикул не разрешился (строка помечена числовым SKU вместо
# кода товара), из-за чего доставка по артикулу не сшивалась с этой строкой и уходила в «Общие».
# УПД даёт точную связку «SKU <-> Артикул» - дописываем её в маппинг каталога (offerAlt в
# build-katya), и такие строки получают настоящий артикул, а доставка садится на них.
#
# Кириллические гомоглифы в коде артикула (х/с/е/о/р/а/у/к/…) нормализуем в латиницу: реальные
# коды - латиница+цифры, кириллица там всегда опечатка (GGTP-20-3х2 -> GGTP-20-3x2).
#
# Вход:  tools/delivery/raw_upd/*.xlsx  (в git не хранятся)
# Выход: data/upd_sku_offer.json  = {sku: "артикул"}
# Запуск из analytics-mvp:  python3 tools/delivery/build_upd_sku_offer.py
import openpyxl, glob, os, json, sys

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
RAW = os.path.join(ROOT, "tools", "delivery", "raw_upd")
OUT = os.path.join(ROOT, "data", "upd_sku_offer.json")

# кириллица -> латиница (гомоглифы)
HOMO = str.maketrans({
    "А": "A", "В": "B", "С": "C", "Е": "E", "Н": "H", "К": "K", "М": "M", "О": "O",
    "Р": "P", "Т": "T", "Х": "X", "У": "Y",
    "а": "a", "в": "b", "с": "c", "е": "e", "н": "h", "к": "k", "м": "m", "о": "o",
    "р": "p", "т": "t", "х": "x", "у": "y",
})


def nrm(s):
    return str(s).strip().translate(HOMO)


def cell(r, i):
    return r[i] if r is not None and i < len(r) else None


def main():
    files = sorted(glob.glob(os.path.join(RAW, "*.xlsx")))
    if not files:
        sys.exit("нет .xlsx в " + RAW)
    sku_off = {}
    for f in files:
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        ws = wb["Отчет о реализации"] if "Отчет о реализации" in wb.sheetnames else wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        hi = None
        for i, r in enumerate(rows[:40]):
            if any((str(c) if c else "") == "Артикул" for c in (r or [])):
                hi = i
                break
        if hi is None:
            wb.close()
            continue
        for r in rows[hi + 3:]:
            off = nrm(cell(r, 2)) if cell(r, 2) else ""
            sk = cell(r, 3)
            sku = str(sk).split(".")[0] if sk else ""
            if sku.isdigit() and off and "того" not in off.lower():
                sku_off[sku] = off  # последний непустой выигрывает
        wb.close()
    with open(OUT, "w", encoding="utf-8") as w:
        json.dump(sku_off, w, ensure_ascii=False)
    print("SKU->артикул из УПД:", len(sku_off))
    print("записано:", OUT)


if __name__ == "__main__":
    main()
