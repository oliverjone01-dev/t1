#!/usr/bin/env python3
"""Чего не хватает в 1С против отчётов Маркета: разложение по месяцам и кабинетам. Только чтение.
Запуск: python3 -I sverka-itog.py <корень репо> <папка с файлами 1С> <выход.xlsx>
Спека (деньги, без баллов; месяц = месяц проводки в реестре взаиморасчётов Маркета, как блок 7 «Отчёта»):
  продажи деньгами = платежи покупателей за товар минус возвраты; доставка деньгами = платежи за доставку минус возвраты;
  прочие начисления = премия, компенсации, возврат услуг; сборы деньгами = удержания (размещение, перевод платежа, доставка Маркета, буст, штрафы, без заказа);
  к перечислению = продажи + доставка + прочие − сборы = «Перевёл» + «В пути» (сходится с блоком 7 до копейки, проверено sverka.py).
  1С: «Отчёт комиссионера» (кабинет по кампании из sverka-prodazhi), реализации «ФИЗИЧЕСКОЕ ЛИЦО <договор>», поступления в банк.
  Не хватает в 1С = продажи деньгами + доставка + прочие − отчёты комиссионера. Разложение: кампании без документа (по выгрузке заказов),
  доставка, переходящие заказы и остальное.
"""
import json, sys, re, glob, collections
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
root, bdir, out = sys.argv[1], sys.argv[2], sys.argv[3]
CAB = {'1023124': 'зеркала', '74986385': 'мебель'}; CONTRACT = {'54641824/26': 'зеркала', '54542918/26': 'мебель'}
rd = lambda f: [json.loads(l) for l in open(f'{root}/analytics-mvp/data-ym/{f}', encoding='utf-8') if l.strip()]
net = rd('netting.ndjson'); orders = [r for r in rd('orders.ndjson') if not r.get('fake')]
def cancelled(r): return bool(re.match(r'^не будет', str(r.get('status') or '').strip(), re.I))
L = collections.defaultdict(lambda: collections.defaultdict(float)); P = collections.defaultdict(lambda: collections.defaultdict(float))
for r in net:
    st = str(r.get('status') or '').strip(); k = (r['d'][:7], CAB.get(r['business'], '?')); a = r['amount']
    if cancelled(r): continue
    if st.lower().startswith('справочно') or r['type'] in ('Списание', 'Возврат списания') or (r['type'] in ('Начисление', 'Возврат') and 'балл' in r['src'].lower()):
        if r['type'] == 'Начисление' or (r['type'] == 'Возврат' and 'балл' in r['src'].lower()): P[k]['баллы начислено'] += a
        else: P[k]['баллы списано на услуги'] += a
        continue
    if not st: continue   # старая схема без статуса (декабрь): деньги от баллов не отличить, как на дашборде
    if r['type'] == 'Начисление' and r['src'] == 'Платёж покупателя': L[k]['доставка' if r['service'] == 'Доставка' else 'продажи'] += a
    elif r['type'] == 'Возврат': L[k]['доставка' if r['service'] == 'Доставка' else 'продажи'] += a
    elif r['type'] == 'Начисление': L[k]['прочие начисления'] += a
    elif r['type'] == 'Удержание': L[k]['сборы деньгами'] += a
kom = []
for f in glob.glob(f'{bdir}/*.xlsx'):
    v = pd.read_excel(f, header=None)
    if 'комиссионера' not in str(v.iloc[1, 0]): continue
    for _, w in v.iloc[6:].iterrows():
        if isinstance(w[1], str) and re.match(r'\d\d\.\d\d\.\d{4}', w[1]): kom.append(dict(ym=w[1][6:] + '-' + w[1][3:5], num=str(w[3]), amount=float(w[6])))
dash = collections.defaultdict(float); camp_cab = {}
for r in orders:
    camp_cab[r['campaign']] = CAB[r['business']]
    if not r['service'] and r['status'] == 'DELIVERED' and r.get('fin'): dash[(r['fin'][:7], r['campaign'])] += r['p_buyer'] * r['count']
used = set()
for d in sorted(kom, key=lambda d: (d['ym'], -d['amount'])):
    cands = [(c, dash[(d['ym'], c)]) for (m, c) in dash if m == d['ym'] and (d['ym'], c) not in used]
    exact = [c for c, v in cands if abs(v - d['amount']) <= 2]
    c = exact[0] if exact else (min(cands, key=lambda cv: abs(cv[1] - d['amount']))[0] if cands and abs(min(cands, key=lambda cv: abs(cv[1] - d['amount']))[1] - d['amount']) < 0.25 * d['amount'] else None)
    d['camp'] = c; d['cab'] = camp_cab.get(c, 'не определён') if c else 'не определён'
    if c: used.add((d['ym'], c))
nodoc = collections.defaultdict(float)
for (m, c), v in dash.items():
    if (m, c) not in used and m <= max(d['ym'] for d in kom): nodoc[(m, camp_cab[c])] += v
upd = collections.defaultdict(float)
for f in glob.glob(f'{bdir}/*.xls'):
    for _, w in pd.read_excel(f, header=0).iterrows():
        c = str(w['Контрагент'])
        if c.startswith('ФИЗИЧЕСКОЕ ЛИЦО'): d = str(w['Дата']); upd[(d[6:] + '-' + d[3:5], CONTRACT.get(c[-11:], '?'))] += float(w['Сумма'])
bank = collections.defaultdict(float)
for f in glob.glob(f'{bdir}/*.xlsx'):
    v = pd.read_excel(f, header=None)
    if 'банковских' not in str(v.iloc[1, 0]): continue
    for _, w in v.iloc[6:].iterrows():
        if isinstance(w[1], str) and re.match(r'\d\d\.\d\d\.\d{4}', w[1]):
            m = re.search(r'договору (\S+?)\.', str(w[7])); bank[(w[1][6:] + '-' + w[1][3:5], CONTRACT.get(m.group(1) if m else '', '?'))] += float(w[6])
months = [m for m in sorted(set(k[0] for k in L)) if '2026-01' <= m <= '2026-09']
r2 = lambda x: round(x + 0.0, 2)
rows = []; tot = collections.defaultdict(float)
for cab in ('зеркала', 'мебель'):
    for m in months:
        l = L[(m, cab)]; k1 = sum(d['amount'] for d in kom if d['ym'] == m and d['cab'] == cab); k0 = sum(d['amount'] for d in kom if d['ym'] == m and d['cab'] == 'не определён') if cab == 'мебель' else 0.0
        sales = l['продажи']; dlv = l['доставка']; oth = l['прочие начисления']; fees = -l['сборы деньгами']; pay = sales + dlv + oth - fees
        gap = sales + dlv + oth - k1 - k0; nd = nodoc.get((m, cab), 0.0); rest = gap - nd - dlv - oth
        row = [cab, m, r2(sales), r2(dlv), r2(oth), r2(fees), r2(pay), r2(bank.get((m, cab), 0.0)), r2(k1 + k0), r2(upd.get((m, cab), 0.0)), r2(gap), r2(nd), r2(dlv), r2(oth), r2(rest), r2(P[(m, cab)]['баллы начислено']), r2(P[(m, cab)]['баллы списано на услуги'])]
        rows.append(row)
        for i, v in enumerate(row[2:], 2): tot[(cab, i)] += v
    rows.append([cab, 'ИТОГО'] + [r2(tot[(cab, i)]) for i in range(2, 17)])
H = ['кабинет', 'месяц проводки', 'Маркет: продажи деньгами (товар, минус возвраты)', 'Маркет: доставка деньгами', 'Маркет: прочие начисления (премия, компенсации, возврат услуг)', 'Маркет: сборы деньгами', '= к перечислению (Перевёл + В пути)', 'Банк: поступило (по дате поступления)',
     '1С: отчёты комиссионера', '1С: реализации (УПД) по договору', 'НЕ ХВАТАЕТ в 1С = продажи + доставка + прочие − отчёты комиссионера', 'из них: кампании без документа в 1С', 'из них: доставка, оплаченная покупателем', 'из них: прочие начисления Маркета', 'из них: переходящие заказы и прочее',
     'справочно: баллы скидки Маркета начислено', 'справочно: баллы списано на услуги']
wb = Workbook(); ws = wb.active; ws.title = 'Как читать'
txt = ['Чего не хватает в 1С против отчётов Маркета. Источник Маркета - реестр взаиморасчётов (API, те же данные, что вкладка «Отчет», блок 7), деньги без баллов, по месяцу проводки. Все цифры [ДАННЫЕ].', '',
       'Почему выходит «мы должны Маркету». Маркет перечисляет нам: продажи деньгами + доставка, оплаченная покупателем + премии и компенсации − сборы деньгами. Если в 1С продажи проведены меньше, чем Маркет реально продал, то поступления и сборы оказываются больше проведённых продаж, и в учёте возникает долг перед Маркетом. Это не долг, а непроведённые продажи.', '',
       'Что именно не проведено (лист «Разложение», столбцы «из них»):',
       '  1. Отчёты комиссионера есть не по всем кампаниям (магазинам) Маркета. По кампаниям 149001087, 149002232, 66056856 (мебель) и 149154933, 142277249, 149087657 (зеркала) документов в 1С нет - это основная часть.',
       '  2. Доставка, оплаченная покупателем: Маркет перечисляет её нам, в отчётах комиссионера её нет (документы равны деньгам за товар без доставки).',
       '  3. Премии Маркета, компенсации за утерю и возвраты услуг: Маркет начисляет, в отчётах комиссионера их нет.',
       '  4. Остаток «переходящие заказы и прочее»: разница месяца проводки (реестр) и месяца доставки (отчёт комиссионера), заказы кампаний, которых нет в выгрузке дашборда (кабинет зеркал до июля), и один документ за июнь без кампании (401 995 ₽, отнесён к мебели).', '',
       'Отдельно про УПД: реализации по договорам в 1С ещё меньше отчётов комиссионера. Если выручка в 1С признаётся по УПД, разрыв с Маркетом больше, чем по отчётам комиссионера. Из чего в 1С собран документ реализации - вопрос к бухгалтерии, по данным Маркета его базу найти не удалось.',
       'Баллы скидки Маркета (справочные столбцы): Маркет начисляет их нам и ими же списывает размещение и буст. В сверке денег их нет; если в 1С услуги Маркета проведены по актам с баллами, а продажи только деньгами, это второй источник «долга» - проверить, что акты услуг в 1С содержат только денежную часть.',
       'Январь: по зеркалам и мебели в выписке банка поступлений нет (первое 20.02), отчётов комиссионера за январь тоже нет.']
for i, t in enumerate(txt, 1): ws.cell(i, 1, t)
ws.column_dimensions['A'].width = 200
ws = wb.create_sheet('Разложение'); bold = Font(bold=True); hdr = PatternFill('solid', fgColor='DDE7F0'); grey = PatternFill('solid', fgColor='EEEEEE'); red = Font(color='B42318', bold=True)
for j, h in enumerate(H, 1): c = ws.cell(1, j, h); c.font = bold; c.fill = hdr; c.alignment = Alignment(wrap_text=True, vertical='top'); ws.column_dimensions[get_column_letter(j)].width = 11 if j < 3 else 17
ws.row_dimensions[1].height = 75
for i, row in enumerate(rows, 2):
    for j, v in enumerate(row, 1):
        c = ws.cell(i, j, v)
        if isinstance(v, (int, float)): c.number_format = '#,##0;[Red]-#,##0'
        if row[1] == 'ИТОГО': c.font = bold; c.fill = grey
        if j == 11 and isinstance(v, (int, float)) and v > 1: c.font = red
ws.freeze_panes = 'C2'
# документы 1С с кабинетом
ws = wb.create_sheet('Отчёты комиссионера 1С'); hd = ['месяц', '№', 'сумма 1С', 'кампания Маркета', 'кабинет']
for j, h in enumerate(hd, 1): c = ws.cell(1, j, h); c.font = bold; c.fill = hdr; ws.column_dimensions[get_column_letter(j)].width = 14
for i, d in enumerate(sorted(kom, key=lambda d: (d['ym'], d['cab'])), 2):
    for j, v in enumerate([d['ym'], d['num'], d['amount'], d['camp'] or 'не найдена', d['cab']], 1): ws.cell(i, j, v)
wb.save(out)
print('ok', out); [print(r[:2], [round(x) for x in r[2:15]]) for r in rows if r[1] == 'ИТОГО' or r[1] in ('2026-05', '2026-09')]
