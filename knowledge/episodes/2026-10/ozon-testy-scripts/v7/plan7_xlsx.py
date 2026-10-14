# v7: книга Excel «Вариант 1 по шагам» для Ивана (14.10: «не видно какие артикулы снижаем, какие повышаем, клики кроме KONUX и KUVINO,
# в Excel коротко, простыми словами и поэтапно»). Числа - из plan7.py / v5 forecast.py (импорт), записки варианта 1 (go ФЕНИКСА 7,85).
# Волна 2 дешёвого клика - условия из sessions/2026-10-12-ozon-testy-itogi.md:78 (пул пересчитан здесь).
# Выход вне репо (прибыль по артикулам): python3 -I plan7_xlsx.py <out.xlsx>
import sys, os, json, datetime as dt, collections as C
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE); sys.path.insert(0, HERE + '/../v5')
import plan7 as P, forecast as F
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
OUT = sys.argv[1]
NAMES = json.load(open('/tmp/ozon-v7/names.json'))
res = json.load(open('/tmp/ozon-v7/plan7.json'))

wb = Workbook(); B = Font(bold=True); H = PatternFill('solid', fgColor='DDEBF7'); G = PatternFill('solid', fgColor='E2EFDA')
R_ = PatternFill('solid', fgColor='FCE4D6'); Y = PatternFill('solid', fgColor='FFF2CC'); thin = Side(style='thin', color='BBBBBB')
def sheet(ws, head, rows, widths, fills=None):
    ws.append(head)
    for c in ws[ws.max_row]: c.font = B; c.fill = H; c.alignment = Alignment(wrap_text=True, vertical='top')
    for i, r in enumerate(rows):
        ws.append(r)
        for c in ws[ws.max_row]:
            c.alignment = Alignment(wrap_text=True, vertical='top'); c.border = Border(bottom=thin)
            if fills and fills[i]: c.fill = fills[i]
    for i, w in enumerate(widths): ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = ws.cell(row=ws.max_row - len(rows) + 1, column=1)

# ---------- лист «Цена вниз»
WAIT = P.WAIT; Q = F.art2sku['GGT-48-5-1-100-180']
price_rows = []
for o in sorted(F.lst.values(), key=lambda o: -F.U[o['sku']]):
    sku = o['sku']; r = F.rows[sku]; st = o['step']; p1 = round(r['p0'] * (1 - st / 100), -1)
    when = '15.10' if sku in P.NOW else ('22.10 (после теста 4; если тест продлят - 29.10)' if sku in WAIT else '22.10 (после новой базы показов)')
    a0 = F.pi(r, r['p0'], 0.05, r['s']); a1 = F.pi(r, r['p0'] * (1 - st / 100), 0.05, r['s'])
    price_rows.append([when, o['art'], NAMES.get(o['art'], ''), round(F.U[sku] / F.WEEKS, 2), round(r['p0']), f'-{st}%', p1,
                       round(r['p0'] * r['s']), round(p1 * r['s']), round(a0), round(a1), f'+{a0 / a1 - 1:.0%}'])
# ---------- лист «Клик»: все идущие кампании
def what(k):
    c = F.camp[k]; t = P.tempo(k)
    if k in P.LAD_IDS:
        nw = {'37675958': '10 000 с 19.11 (по умолчанию)', '39931246': '5 000 с 19.11 (по умолчанию)'}[k]
        return 'не трогаем до 10.11', nw, 'клик и цена в одной карточке: урез сейчас включит стоп-правило цены. 10.11 решаем бюджет', Y
    if k in P.EXP: return 'на минимум 15.10', '2 000 (8-12 ₽ за клик)', 'дорого, заказов почти нет: окупается только если половина заказов рекламы свои', R_
    inp = k in P.IN_P
    if t > 2500 and not inp: return 'на минимум 15.10', '2 000 (8-12 ₽ за клик)', 'тратит больше минимума, заказов мало; в карточке нет снижения цены', R_
    if t > 2500: return 'оставляем до 10.11', f'{t:,.0f}'.replace(',', ' '), 'в карточке снижаем цену: до итога 10.11 рекламу не трогаем', Y
    return 'оставляем', f'{t:,.0f}'.replace(',', ' ') if t else 'как есть', 'уже минимум или около', None
camp_rows, cf = [], []
for k in sorted(P.ACTIVE, key=lambda k: -P.tempo(k)):
    c = F.camp[k]; o4 = c['o']; sp4 = c['sp']; act, nw, why, fl = what(k)
    pi_, nm_ = F.camp_pi(c); pi_ = round(pi_) if 'медиана' not in nm_ else 'общая на группу'
    camp_rows.append([act, k, c['off'], NAMES.get(c['off'], ''), round(P.tempo(k)), round(sp4 / F.WEEKS), o4,
                      round(sp4 / o4) if o4 else 'нет заказов', pi_, nw, why]); cf.append(fl)
N_MIN = sum(1 for r in camp_rows if r[0].startswith('на мин'))
# ---------- волна 2 дешёвого клика: пул по условиям итогов 12.10 (п.3)
ads_sku = C.defaultdict(float)
for a in F.L('ads_sku_daily.ndjson'):
    if a['cid'] != F.CPO_ID and a['d'] >= '2026-09-01': ads_sku[a['sku']] += a['sp']
o60 = C.Counter()
for o in F.L('orders_daily.ndjson'):
    if '2026-08-13' <= o['d'] <= '2026-10-11' and o['status'] != 'cancelled': o60[o['sku']] += o['units']
camp_cards = {P.camp_card(F.camp[k]) for k in P.ACTIVE}
excl_cards = P.PCARDS | {P.card.get(s) for s in WAIT}
pool = []
for sku, r in F.rows.items():
    cd = P.card.get(sku)
    if not cd or cd in camp_cards or cd in excl_cards or ads_sku[sku] > 0 or o60[sku] < 1: continue
    pool.append((o60[sku], r['art'], sku, cd, F.pi(r, r['p0'], 0.05, r['s']), r['p0']))
pool.sort(key=lambda x: (-x[0], -x[4]))
wave, per_card = [], C.Counter()
for x in pool:
    if per_card[x[3]] >= 3 or len(wave) >= 20: continue
    per_card[x[3]] += 1; wave.append(x)
wave_rows = [['новая кампания 16.10', x[1], NAMES.get(x[1], ''), x[0], round(x[5]), round(x[4]), '8 ₽ за клик, 1 500 ₽ в неделю',
              'товар продаётся без рекламы, в карточке нет рекламы и нет снижения цены'] for x in wave]

# ---------- лист 1 «Коротко»
ph1 = res['ph1']; ph2 = res['ph2']
def r3(key, f): return ' / '.join(f'{ph1[f"{key}|{sc}"]["d"][f] / 1e3:+,.0f}'.replace(',', ' ') for sc in F.SCN)
mo = {sc: res[f'month|False|{sc}']['m'] for sc in F.SCN}
ws = wb.active; ws.title = 'Коротко'
lines = [
    ['Вариант 1 по шагам. GENGLASS OZON, 14.10.2026. Все суммы - прогноз [ГИПОТЕЗА], три сценария: плохо / средне / хорошо.'],
    [''],
    ['Что делаем', 'Сколько артикулов', 'Когда', 'Что ждём (прибыль, тыс. ₽)'],
    ['1. Снижаем цену на 5-10%', f'{len(P.NOW)} сразу + {len(F.lst) - len(P.NOW)} позже', '15.10 и 22.10', 'одна цена: около нуля в среднем сценарии (решает, вырастут ли заказы)'],
    ['2. Дорогую рекламу, которая не даёт заказов, переводим на минимум', f'{N_MIN} кампаний', '15.10', '+98...+118 за 4 недели в любом сценарии (по 7 самым дорогим)'],
    ['3. KONUX и KUVINO пока не трогаем', '2 кампании', 'до 10.11', 'чтобы проверить цену в их карточках честно'],
    ['4. Новая дешёвая реклама на товарах без рекламы (волна 2)', f'{len(wave)} товаров', '16.10-30.10', '+10...+35 за волну, расход до 120 за 4 недели (лимитом кампаний)'],
    ['5. Итог по цене и решение по KONUX и KUVINO', '-', '10.11', 'возвращаем цену, только если заказы упали на 50% и больше'],
    ['6. KONUX 10 000, KUVINO 5 000 ₽ в неделю (вместо 30 000 и 20 000)', '2 кампании', 'с 19.11', 'по умолчанию; решаем 10.11'],
    [''],
    ['Итог за 4 недели (15.10-11.11), шаги 1-3, тыс. ₽', 'плохо / средне / хорошо', '', r3('v', 'prof')],
    ['Месяц после (шаги 1-6 без волны 2), прибыль тыс. ₽', '', '', ' / '.join(f'{mo[sc]["prof"] / 1e3:+,.0f}'.replace(',', ' ') for sc in F.SCN)],
    ['Месяц после, заказов (сейчас около 286)', '', '', ' / '.join(f'{mo[sc]["units"]:+.0f}' for sc in F.SCN)],
    ['Месяц после, реклама за клик, тыс. ₽ (сейчас 612)', '', '', f'{612 + mo["средний"]["clk"] / 1e3:,.0f}'.replace(',', ' ')],
    [''],
    ['Что решаете: 1) старт 15.10 до 10:00 - да / нет; 2) шаг цены -10% (этот план) или -5% замером (книга вариантов); 3) волна 2 дешёвой рекламы 16.10 - да / нет.'],
    ['Почему KONUX и KUVINO не режем сразу: у них почти все показы в поиске - рекламные. Урежем клик - доля показов упадёт, и стоп-правило вернёт цену, хотя цена ни при чём.'],
    ['Лесенку ставок и рекламу только в выходные не берём: в выходные заказов на 1 000 кликов столько же, а лесенка не отличает свои заказы от перетянутых у соседей.'],
]
for l in lines: ws.append(l)
ws['A1'].font = Font(bold=True, size=12)
for c in ws[3]: c.font = B; c.fill = H
for row in ws.iter_rows():
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
for col, w in zip('ABCD', (62, 22, 16, 60)): ws.column_dimensions[col].width = w

# ---------- лист 2 «Этапы»
ws = wb.create_sheet('Этапы')
stages = [
    ['15.10 чт, до 10:00', 'Цена вниз', f'{len(P.NOW)} артикулов (лист «Цена вниз», дата 15.10)', 'кабинет', 'цена на сайте ниже на 5-10% к вечеру'],
    ['15.10 чт', 'Реклама на минимум', f'{N_MIN} кампаний (лист «Клик», красные)', 'кабинет', 'расход клика меньше на 0,13 млн в неделю (по темпу 08-11.10)'],
    ['15-21.10', 'Новая база показов', 'карточки TRUBIS и QUATRIS, где урезали рекламу', 'Катя', 'база для стоп-правила'],
    ['каждый день', 'Стоп-правило цены', 'все артикулы со сниженной ценой', 'Катя', 'доля в показах ниже половины базы 3 дня или выпал из «Эластичного бустинга» - вернуть цену'],
    ['16.10 пт', 'Волна 2 дешёвой рекламы', f'{len(wave)} товаров (лист «Клик», зелёные)', 'кабинет', '8 ₽ за клик, 1 500 ₽ в неделю на товар'],
    ['22.10 чт', 'Цена вниз, часть 2', f'{len(F.lst) - len(P.NOW)} артикулов (лист «Цена вниз», дата 22.10)', 'кабинет', 'если тест 4 продлят - 29.10'],
    ['30.10 пт', 'Итог волны 2', f'{len(wave)} товаров', 'Катя + Иван', 'заказы карточек выше контроля и прибыль прироста больше расхода - продолжаем; нет - выключаем'],
    ['22.10, 29.10, 05.11', 'Промежуточный взгляд', 'карточки со сниженной ценой против остального магазина', 'Катя', 'без решений, только провал'],
    ['10.11 вт', 'Итог по цене', 'все артикулы со сниженной ценой', 'Иван', 'возвращаем цену, только если заказы упали на 50% и больше'],
    ['10.11 вт', 'Решение по KONUX и KUVINO', '2 кампании', 'Иван', 'минимум / 10 000 и 5 000 / как сейчас'],
    ['11-17.11', 'Распродажа', '-', '-', 'ничего не меняем, из замера исключаем'],
    ['19.11 чт', 'Клик KONUX и KUVINO', '2 кампании', 'кабинет', 'по решению 10.11'],
]
sheet(ws, ['Когда', 'Что', 'Какие артикулы', 'Кто', 'Что проверяем / ждём'], stages, (18, 26, 46, 14, 60))

# ---------- лист 3 «Цена вниз»
ws = wb.create_sheet('Цена вниз')
sheet(ws, ['Когда', 'Артикул', 'Товар', 'Заказов в неделю (14.09-11.10)', 'Наша цена сейчас, ₽', 'Шаг', 'Новая наша цена, ₽',
           'Цена на сайте сейчас, ₽', 'Цена на сайте после, ₽', 'Прибыль с заказа сейчас, ₽', 'Прибыль с заказа после, ₽',
           'Заказов должно стать больше на'], price_rows, (16, 22, 40, 12, 12, 7, 12, 12, 12, 12, 12, 12),
      [G if r[0] == '15.10' else Y for r in price_rows])
ws.append([]); ws.append(['25 артикулов из вашего файла не трогаем: после снижения у них рентабельность ниже 10% (у 5 - убыток).'])

# ---------- лист 4 «Клик»
ws = wb.create_sheet('Клик')
sheet(ws, ['Что делаем', 'Кампания', 'Артикул', 'Товар', 'Тратит сейчас, ₽/нед (темп 08-11.10)', 'Тратила в среднем, ₽/нед (4 нед)',
           'Заказов по рекламе за 4 нед', '₽ рекламы на заказ', 'Прибыль с заказа, ₽', 'Новый бюджет, ₽/нед', 'Почему'],
      camp_rows, (22, 11, 22, 30, 14, 14, 11, 12, 12, 22, 60), cf)
ws.append([]); ws.append(['Волна 2: новая дешёвая реклама на товарах, где рекламы не было (условия теста 1)'])
ws.cell(ws.max_row, 1).font = B
ws.append(['Что делаем', 'Артикул', 'Товар', 'Заказов за 60 дней', 'Наша цена, ₽', 'Прибыль с заказа, ₽', 'Бюджет', 'Почему этот товар'])
for c in ws[ws.max_row]: c.font = B; c.fill = H
for r in wave_rows:
    ws.append(r)
    for c in ws[ws.max_row]: c.fill = G; c.alignment = Alignment(wrap_text=True, vertical='top')

# ---------- лист 5 «Прогноз»
ws = wb.create_sheet('Прогноз')
def row(lbl, f, key='v'): return [lbl] + [round(ph1[f'{key}|{sc}']['d'][f] / (1e3 if f in ('prof', 'prof_adm') else 1e6 if f in ('turn', 'pay') else 1), 1 if f in ('turn', 'pay', 'units') else 0) for sc in F.SCN]
fr = [
    ['4 недели 15.10-11.11 (шаги 1-3), разница с «как сейчас»', '', '', ''],
    row('Прибыль без АДМ, тыс. ₽', 'prof'), row('Прибыль с АДМ 30%, тыс. ₽', 'prof_adm'), row('Заказов, шт', 'units'),
    row('Оборот по нашей цене, млн ₽', 'turn'), row('К выплате, млн ₽', 'pay'),
    ['Для сравнения: только урез рекламы, цену не трогаем, тыс. ₽', 118, 108, 98],
    ['Для сравнения: всё как сейчас, только цена, тыс. ₽', *[round(ph1[f'False|{sc}']['d']['prof'] / 1e3) for sc in F.SCN]],
    ['', '', '', ''],
    ['Месяц после (цена, урез, KONUX 10 000, KUVINO 5 000), разница с «как сейчас»', '', '', ''],
    ['Прибыль без АДМ, тыс. ₽'] + [round(mo[sc]['prof'] / 1e3) for sc in F.SCN],
    ['Заказов, шт'] + [round(mo[sc]['units']) for sc in F.SCN],
    ['Оборот, млн ₽'] + [round(mo[sc]['turn'] / 1e6, 2) for sc in F.SCN],
    ['К выплате, млн ₽'] + [round(mo[sc]['pay'] / 1e6, 2) for sc in F.SCN],
    ['', '', '', ''],
    ['KONUX и KUVINO с 19.11, 6 недель, разница с «как сейчас», тыс. ₽', '', '', ''],
] + [[{'А': 'оба на минимум', 'Б': 'как сейчас (30 000 и 20 000)', 'В': 'лесенка 5-10-15 тыс.', 'Г': 'только выходные',
       'Д': 'KONUX 10 000, KUVINO 5 000'}[v]] + [round(ph2[sc]['R'][v] / 1e3) for sc in F.SCN] for v in 'АБВГД'] + [
    ['', '', '', ''],
    ['Волна 2 дешёвой рекламы, 4 недели, тыс. ₽', '+10', '', '+35'],
    ['Сценарии: плохо - цена почти не добавляет заказов, реклама не приносит своих заказов; средне - -10% цены дают +23% заказов, половина заказов рекламы свои; хорошо - +46% и все свои.', '', '', ''],
]
sheet(ws, ['', 'Плохо', 'Средне', 'Хорошо'], fr, (70, 12, 12, 12))
wb.save(OUT)
print('цена', len(price_rows), 'кампаний', len(camp_rows), 'на минимум', sum(1 for r in camp_rows if r[0].startswith('на мин')),
      'волна 2', len(wave), 'пул', len(pool))
