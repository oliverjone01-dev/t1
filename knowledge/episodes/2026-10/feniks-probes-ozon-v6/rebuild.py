# ФЕНИКС 14.10, аудит OZON v6: пересборка книги скриптом автора и сравнение с доставленной ячейка в ячейку (формулы и значения).
# python3 -I rebuild.py <каталог с rp_2026-09/10.txt> <доставленная xlsx> <recalc.py>
import shutil, subprocess, os, sys
import openpyxl
src, delivered, recalc = sys.argv[1:4]
out = '/tmp/feniks-v6/rebuild/'
os.makedirs(out, exist_ok=True)
for f in ('rp_2026-09.txt', 'rp_2026-10.txt'):
    shutil.copy(os.path.join(src, f), out + f)
p = subprocess.run(['python3', '-I', '/home/user/t1/knowledge/episodes/2026-10/ozon-testy-scripts/v6/forecast30.py', out], capture_output=True, text=True)
print('build rc', p.returncode, p.stderr[-300:])
fn = out + 'OZON-prognoz-30mln-2026-10-14.xlsx'
r = subprocess.run(['python3', recalc, fn], capture_output=True, text=True)
print('recalc', r.stdout.replace('\n', ' ')[:200])
cells = df = dv = 0
for mode in (False, True):
    wa, wb = openpyxl.load_workbook(delivered, data_only=mode), openpyxl.load_workbook(fn, data_only=mode)
    assert wa.sheetnames == wb.sheetnames
    for n in wa.sheetnames:
        for row in wa[n].iter_rows():
            for c in row:
                v1, v2 = c.value, wb[n][c.coordinate].value
                cells += not mode
                same = v1 == v2 or (isinstance(v1, (int, float)) and isinstance(v2, (int, float)) and abs(v1 - v2) < 1e-6)
                if not same:
                    dv += mode; df += not mode
print(f'ячеек {cells}, расхождений формул {df}, значений {dv}')
