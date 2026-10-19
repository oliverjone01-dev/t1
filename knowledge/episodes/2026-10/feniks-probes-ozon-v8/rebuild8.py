# ФЕНИКС 14.10, книга вариантов v8: пересборка variants.py в /tmp и сравнение с доставленной ячейка в ячейку; тексты листа «Другие пути».
import shutil, subprocess, os, glob, sys
import openpyxl
src, delivered = sys.argv[1], sys.argv[2]
out = '/tmp/feniks-v6/rebuild8/'
os.makedirs(out, exist_ok=True)
for f in glob.glob(os.path.join(src, 'rp_2026-*.txt')): shutil.copy(f, out + os.path.basename(f))
env = dict(os.environ, RECALC='/root/.claude/skills/synced/7cb27221-5718-4fc3-bc16-76561189fcae_0fffcc3c-f3a6-44d7-8da9-74f689d08934/xlsx/scripts/recalc.py')
p = subprocess.run(['python3', '-I', '-B', '/home/user/t1/knowledge/episodes/2026-10/ozon-testy-scripts/v8/variants.py', out, '/home/user/t1/analytics-mvp/data'],
                   capture_output=True, text=True, env=env, timeout=600)
print('build rc', p.returncode, p.stdout[-200:].replace('\n', ' '), p.stderr[-300:])
fn = glob.glob(out + 'OZON-varianty-*.xlsx')[0]
cells = df = dv = 0
for mode in (False, True):
    wa, wb = openpyxl.load_workbook(delivered, data_only=mode), openpyxl.load_workbook(fn, data_only=mode)
    for n in wa.sheetnames:
        for row in wa[n].iter_rows():
            for c in row:
                v1, v2 = c.value, wb[n][c.coordinate].value
                cells += not mode
                if not (v1 == v2 or (isinstance(v1, (int, float)) and isinstance(v2, (int, float)) and abs(v1 - v2) < 1e-6)):
                    dv += mode; df += not mode
print(f'ячеек {cells}, расхождений формул {df}, значений {dv}')
wv = openpyxl.load_workbook(delivered, data_only=True)
ws = wv['Другие пути']
for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
    vals = [str(c.value)[:150] for c in row if c.value is not None]
    if vals: print(' | '.join(vals)[:420])
