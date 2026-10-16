import sys, re
from openpyxl import load_workbook
for p in sys.argv[1:]:
    wb = load_workbook(p); txt = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if isinstance(v, str): txt.append(v)
    t = '\n'.join(txt)
    print(p.rsplit('/', 1)[-1], 'emdash', t.count('—'), 'около', len(re.findall(r'около', t)), 'примерно', len(re.findall(r'примерно', t)), 'порядка', len(re.findall(r'порядка', t)))
