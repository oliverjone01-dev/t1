import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
SP = sys.argv[1]
d = json.load(open(f"{SP}/ordan.json"))
MON = {"2026-07": "Июль 2026", "2026-08": "Август 2026"}
HF = PatternFill("solid", fgColor="1F4E5A"); HFont = Font(bold=True, color="FFFFFF")
EST = Font(italic=True, color="7B61FF"); BOLD = Font(bold=True)
def num(t):
    t = t.replace("\xa0", "").replace(" ", "").replace("−", "-").replace(",", ".")
    if t in ("", "—", "-"): return None
    if t.endswith("%"):
        try: return float(t[:-1]) / 100
        except: return t
    try: return float(t)
    except: return t
def head(ws, cols):
    ws.append(cols)
    for c in ws[1]: c.fill = HF; c.font = HFont; c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "B2"; ws.row_dimensions[1].height = 42
def finish(ws, widths):
    ws.auto_filter.ref = ws.dimensions
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
def orders_of(r):
    out = []
    for x in r["rows"]:
        if x["kind"] != "order": continue
        o = x["first"].split()[0]; m = r["meta"].get(o, {})
        off = next((sp for sp in x["spans"] if not sp.startswith("(")), "") if x["spans"] else ""
        note = [sp for sp in x["spans"] if sp != off]      # «(2 отпр.)», «отменён, только расходы», «delivering», ...
        label = o + (" (" + "; ".join(n.strip("()") for n in note) + ")" if note else "")
        city = x["cells"][25]["t"]
        out.append((label, dict(m, off=off or m.get("off"), city=city, key=o), x))
    out.sort(key=lambda t: (t[1].get("d", ""), t[1]["key"]))
    return out
def method(typ, g):
    ship, prt, dl = g("Наша доставка") or 0, g("Услуги партнёров") or 0, g("Логистика") or 0
    if isinstance(ship, str): ship = 0  # «нет в ведомости»
    if typ in ("FBO", "FBS"): return "логистика OZON"
    if typ == "rFBS": return "наша перевозка" if ship else "услуги партнёров"
    if dl: return "логистика OZON"
    if ship: return "наша перевозка"
    if prt: return "услуги партнёров"
    return "нет начислений"

# 1. Аналитика по заказам
wb = Workbook(); wb.remove(wb.active)
for ym, r in d.items():
    H = r["head"]; mid = H[2:25]            # «Заказано» ... «Рентаб.»; «Города доставки» ушла в «Город (ведомость)»
    ws = wb.create_sheet(MON[ym])
    head(ws, ["Заказ", "Дата", "Тип доставки", "Город (ведомость)", "Категория", "Артикул"] + mid)
    for o, m, x in orders_of(r):
        c = x["cells"]
        ws.append([o, m.get("d"), c[1]["t"], m["city"], x["cat"], m.get("off")] + [num(c[j]["t"]) for j in range(2, 25)])
        for j in range(2, 25):
            if c[j]["est"]: ws.cell(ws.max_row, 7 + j - 2).font = EST
    for kind, label in (("other", None), ("total", "ИТОГО")):
        for x in r["rows"]:
            if x["kind"] != kind: continue
            c = x["cells"]
            ws.append([label or x["first"], "", "", "", "", ""] + [num(c[j]["t"]) for j in range(2, 25)])
            for cell in ws[ws.max_row]: cell.font = BOLD
    for row in ws.iter_rows(min_row=2):
        for cell in row[6:]:
            if isinstance(cell.value, float):
                cell.number_format = "0.0%" if mid[cell.column - 7].startswith("Рентаб") else "#,##0"
    finish(ws, [20, 11, 10, 34, 14, 22] + [13] * len(mid))
wb.save(f"{SP}/ozon_analitika_po_zakazam_2026-07_2026-08.xlsx")

# 2. Логистика по заказам
wb = Workbook(); wb.remove(wb.active)
for ym, r in d.items():
    H = r["head"]; ws = wb.create_sheet(MON[ym])
    cols = ["Заказ (постинг)", "Дата заказа", "Артикул", "Категория", "Город доставки", "Способ доставки",
            "Логистика OZON, ₽", "Услуги партнёра, ₽", "Наша доставка (счёт перевозчика), ₽",
            "Доход с покупателя за доставку, ₽", "Нетто нашей доставки (доход−расход), ₽"]
    head(ws, cols); S = [0] * 5
    for o, m, x in orders_of(r):
        c = x["cells"]; g = lambda name: num(c[H.index(name)]["t"])
        # «Наша доставка» может быть текстом «нет в ведомости» (Иван 30.09): в ячейку - текст, в сумму - 0.
        raw = [g("Логистика") or 0, g("Услуги партнёров") or 0, g("Наша доставка") or 0, g("Доставка покупателя") or 0]
        vals = [v if isinstance(v, (int, float)) else 0 for v in raw]
        net = vals[3] - vals[2]
        ws.append([o, m.get("d"), m.get("off"), x["cat"], m["city"], method(c[1]["t"], g)] + raw + [net])
        for i, v in enumerate(vals + [net]): S[i] += v
    ws.append(["ИТОГО", "", "", "", "", ""] + S)
    for cell in ws[ws.max_row]: cell.font = BOLD
    for row in ws.iter_rows(min_row=2):
        for cell in row[6:]: cell.number_format = "#,##0"
    finish(ws, [20, 11, 22, 14, 34, 18, 14, 14, 18, 18, 18])
    print(ym, "заказов", ws.max_row - 2, "лог OZON", round(S[0]), "партнёры", round(S[1]), "наша", round(S[2]), "доход", round(S[3]), "нетто", round(S[4]))
wb.save(f"{SP}/ozon_logistika_po_zakazam_2026-07_2026-08.xlsx")
