#!/usr/bin/env python3
"""Сверка 1С: реализации Маркета (отчёты по продажам, контрагент «Яндекс маркет фбс») против поступлений от Маркета в банк. Только чтение.
Запуск: python3 -I sverka-prihod-prodazhi.py <корень репо> <папка 1С 19.10 (банк, договоры)> <папка 1С 20.10 (реализации)> <выход.xlsx>
Спека (все цифры [ДАННЫЕ]):
  Реализации: выгрузка 20.10, 52 документа, колонка «Сумма», месяц по дате документа. Кабинет документа: по договору из выгрузки 19.10
      (контрагент «ФИЗИЧЕСКОЕ ЛИЦО <договор>», 54641824/26 зеркала, 54542918/26 мебель), иначе по документу-основанию («Документ отгрузки»),
      кабинет которого известен по другой строке; иначе «не определён».
  Банк: реестры банковских документов 19.10, по договору в назначении платежа, месяц по дате поступления.
  Отчёты комиссионера: реестр 19.10 (20.10 идентичен), кабинет по паре с магазином Маркета (как в sverka-otkloneniya.py).
  Маркет: реестр взаиморасчётов (деньги без баллов, месяц проводки): продажи товаром, к перечислению = продажи + доставка + прочие − сборы.
  Сверка: поступления должны быть меньше продаж на сборы Маркета. Отклонение = банк − реализации (положительное: пришло больше, чем продано).
"""
import json, sys, re, glob, collections
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

root, d19, d20, out = sys.argv[1:5]
CAB = {'1023124': 'зеркала', '74986385': 'мебель'}; CONTRACT = {'54641824/26': 'зеркала', '54542918/26': 'мебель'}
rd = lambda f: [json.loads(l) for l in open(f'{root}/analytics-mvp/data-ym/{f}', encoding='utf-8') if l.strip()]
r2 = lambda x: round(x + 0.0, 2)
ymd = lambda s: s[6:10] + '-' + s[3:5]

# --- реализации 20.10 + договор из 19.10
new = pd.read_excel(glob.glob(f'{d20}/*.xls')[0], header=0)
old = pd.read_excel(glob.glob(f'{d19}/*.xls')[0], header=0)
contract = {str(w['Номер']): CONTRACT[str(w['Контрагент'])[-11:]] for _, w in old.iterrows() if str(w['Контрагент']).startswith('ФИЗИЧЕСКОЕ ЛИЦО')}
docs = []
for _, w in new.iterrows():
    m = re.search(r'(НФБУ-\d+) от (\d\d\.\d\d\.\d{4})', str(w['Документ отгрузки']))
    docs.append(dict(date=str(w['Дата']), ym=ymd(str(w['Дата'])), num=str(w['Номер']), amount=float(w['Сумма']), base=m.group(1) if m else '', base_date=m.group(2) if m else '',
                     cab=contract.get(str(w['Номер'])), how='по договору' if str(w['Номер']) in contract else ''))
base_cab = {}
for d in docs:
    if d['cab'] and d['base']: base_cab.setdefault(d['base'], set()).add(d['cab'])
for d in docs:
    if not d['cab'] and d['base'] in base_cab and len(base_cab[d['base']]) == 1: d['cab'] = next(iter(base_cab[d['base']])); d['how'] = 'по документу-основанию'
for d in docs:
    if not d['cab']: d['cab'] = 'не определён'; d['how'] = 'договор не виден'
    d['corr'] = d['base_date'] and ymd(d['base_date']) != d['ym']

# --- банк 19.10
bank = collections.defaultdict(float); bank_rows = []
for f in glob.glob(f'{d19}/bank-*.xlsx'):
    v = pd.read_excel(f, header=None)
    for _, w in v.iloc[6:].iterrows():
        if isinstance(w[1], str) and re.match(r'\d\d\.\d\d\.\d{4}', w[1]):
            m = re.search(r'договору (\S+?)\.', str(w[7])); cab = CONTRACT.get(m.group(1) if m else '', '?')
            bank[(ymd(w[1]), cab)] += float(w[6]); bank_rows.append([ymd(w[1]), w[1], str(w[3]), cab, float(w[6])])

# --- отчёты комиссионера 19.10, кабинет по магазину
orders = [r for r in rd('orders.ndjson') if not r.get('fake')]
camp_cab = {r['campaign']: CAB[r['business']] for r in orders}
goods = collections.defaultdict(float)
for r in orders:
    if not r['service'] and r['status'] == 'DELIVERED' and r.get('fin'): goods[(r['fin'][:7], r['campaign'])] += r['p_buyer'] * r['count']
kom = []
for f in glob.glob(f'{d19}/*.xlsx'):
    v = pd.read_excel(f, header=None)
    if 'комиссионера' not in str(v.iloc[1, 0]): continue
    for _, w in v.iloc[6:].iterrows():
        if isinstance(w[1], str) and re.match(r'\d\d\.\d\d\.\d{4}', w[1]): kom.append(dict(ym=ymd(w[1]), num=str(w[3]), amount=float(w[6]), cab=None))
used = set()
for d in sorted(kom, key=lambda d: (d['ym'], -d['amount'])):
    h = [c for (m, c) in goods if m == d['ym'] and (m, c) not in used and abs(goods[(m, c)] - d['amount']) <= 2]
    if h: d['cab'] = camp_cab[h[0]]; used.add((d['ym'], h[0]))
for d in sorted(kom, key=lambda d: (d['ym'], -d['amount'])):
    if d['cab']: continue
    cands = [c for (m, c) in goods if m == d['ym'] and (m, c) not in used]
    if cands:
        c = min(cands, key=lambda c: abs(goods[(d['ym'], c)] - d['amount']))
        if abs(goods[(d['ym'], c)] - d['amount']) < 0.25 * d['amount']: d['cab'] = camp_cab[c]; used.add((d['ym'], c))
    if not d['cab']: d['cab'] = 'не определён'

# --- Маркет: реестр
net = rd('netting.ndjson')
def cancelled(r): return bool(re.match(r'^(справочно:\s*)?не будет', str(r.get('status') or '').strip(), re.I))
M = collections.defaultdict(lambda: collections.defaultdict(float))
for r in net:
    st = str(r.get('status') or '').strip()
    if cancelled(r) or not st or st.lower().startswith('справочно') or r['type'] in ('Списание', 'Возврат списания') or 'балл' in r['src'].lower(): continue
    k = (r['d'][:7], CAB.get(r['business'], '?')); a = r['amount']
    if r['type'] == 'Начисление' and r['src'] == 'Платёж покупателя': M[k]['доставка' if r['service'] == 'Доставка' else 'продажи'] += a
    elif r['type'] == 'Возврат': M[k]['доставка' if r['service'] == 'Доставка' else 'продажи'] += a
    elif r['type'] == 'Начисление': M[k]['прочие'] += a
    elif r['type'] == 'Удержание': M[k]['сборы'] += -a

# --- по месяцам
months = ['2026-%02d' % i for i in range(1, 11)]
rows = []; T = {}
for cab in ('зеркала', 'мебель', 'не определён'):
    t = collections.defaultdict(float)
    for m in months:
        rl = sum(d['amount'] for d in docs if d['ym'] == m and d['cab'] == cab); bk = bank.get((m, cab), 0.0)
        km = sum(d['amount'] for d in kom if d['ym'] == m and d['cab'] == cab); mm = M[(m, cab)]
        pay = mm['продажи'] + mm['доставка'] + mm['прочие'] - mm['сборы']
        row = [cab, m, r2(rl), r2(bk), r2(bk - rl), r2(km), r2(bk - km), r2(mm['продажи']), r2(mm['сборы']), r2(pay), r2(mm['продажи'] - rl)]
        if cab == 'не определён' and not any(row[2:]): continue
        rows.append(row)
        for i, v in enumerate(row[2:], 2): t[i] += v
    rows.append([cab, 'ИТОГО'] + [r2(t[i]) for i in range(2, 11)]); T[cab] = t
rows_d = [[d['ym'], d['date'], d['num'], d['amount'], d['cab'], d['how'], d['base'] + (' от ' + d['base_date'] if d['base_date'] else ''), 'корректировка к ' + ymd(d['base_date']) if d['corr'] else 'основной'] for d in sorted(docs, key=lambda d: (d['ym'], d['cab'], d['date']))]
bank_rows.sort()

wb = Workbook(); bold = Font(bold=True); hdr = PatternFill('solid', fgColor='DDE7F0'); grey = PatternFill('solid', fgColor='EEEEEE'); red = Font(color='B42318', bold=True)
def sheet(title, header, rows, widths, note, redcol=None):
    ws = wb.create_sheet(title); ws.cell(1, 1, note).alignment = Alignment(wrap_text=True, vertical='top'); ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(header)); ws.row_dimensions[1].height = 15 * (note.count('\n') + 2)
    for j, h in enumerate(header, 1): c = ws.cell(3, j, h); c.font = bold; c.fill = hdr; c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[3].height = 60
    for i, row in enumerate(rows, 4):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v)
            if isinstance(v, (int, float)) and not isinstance(v, bool): c.number_format = '#,##0;[Red]-#,##0'
            if row[1] == 'ИТОГО': c.font = bold; c.fill = grey
            if redcol and j in redcol and isinstance(v, (int, float)) and v > 2: c.font = red
    for j in range(1, len(header) + 1): ws.column_dimensions[get_column_letter(j)].width = widths[j - 1] if j - 1 < len(widths) else 14
    ws.freeze_panes = ws.cell(4, 3); return ws
wb.remove(wb['Sheet'])
tz, tm, tn = T['зеркала'], T['мебель'], T['не определён']
f = lambda x: f'{x:,.0f}'.replace(',', ' ')
ws = wb.create_sheet('Как читать'); ws.column_dimensions['A'].width = 200
txt = ['Сверка 1С: реализации Маркета (выгрузка 20.10, контрагент «Яндекс маркет фбс», 52 документа) против поступлений от Маркета в банк (реестры 19.10). Все цифры [ДАННЫЕ].', '',
       'Правило: Маркет перечисляет продажи за вычетом своих сборов, поэтому поступления должны быть МЕНЬШЕ проведённых продаж. Если пришло больше, чем продано, продажи проведены не полностью.',
       f'Итог январь-октябрь: зеркала - реализации {f(tz[2])}, поступило {f(tz[3])}, пришло больше продаж на {f(tz[4])}. Мебель - реализации {f(tm[2])}, поступило {f(tm[3])}, пришло больше продаж на {f(tm[4])}.',
       f'Документы без договора (кабинет не определён): {f(tn[2])}. Если отнести их все к мебели, разрыв по мебели {f(tm[4] - tn[2])}, по зеркалам без изменений.',
       '', 'Лист «По месяцам»: реализации и поступления по месяцам и кабинетам, рядом отчёты комиссионера и данные Маркета (продажи деньгами, сборы, к перечислению) для сравнения.',
       'Лист «Реализации»: все 52 документа, кабинет и как он определён (по договору из выгрузки 19.10, по документу-основанию, не виден), основной или корректировка.',
       'Лист «Поступления»: все платёжки из реестров банка 19.10 по договорам.',
       '', 'Кабинет реализации взят по договору из выгрузки 19.10 (в ней контрагент был показан как «ФИЗИЧЕСКОЕ ЛИЦО <договор>»), для корректировок - по документу-основанию. Выгрузка 20.10 договора не показывает.',
       'Реестр отчётов комиссионера 20.10 совпадает с реестром 19.10 байт в байт.']
for i, t in enumerate(txt, 1): ws.cell(i, 1, t)
sheet('По месяцам', ['кабинет', 'месяц', '1С: реализации Маркета', 'Банк: поступило от Маркета', 'ПРИШЛО БОЛЬШЕ, ЧЕМ ПРОДАНО = банк − реализации', '1С: отчёты комиссионера', 'банк − отчёты комиссионера',
       'Маркет: продажи деньгами (месяц проводки)', 'Маркет: сборы деньгами', 'Маркет: к перечислению', 'Маркет продажи − реализации 1С'],
      rows, [13, 9, 16, 16, 20, 16, 16, 18, 14, 16, 18], 'Поступления против проведённых продаж по месяцам. Красным - месяцы, где пришло больше, чем продано. Маркет - реестр взаиморасчётов, деньги без баллов, месяц проводки.', redcol={5, 7, 11})
sheet('Реализации', ['месяц', 'дата', '№', 'сумма', 'кабинет', 'как определён кабинет', 'документ-основание', 'тип'], rows_d, [9, 11, 13, 13, 13, 22, 28, 22], 'Все документы реализации из выгрузки 20.10.')
sheet('Поступления', ['месяц', 'дата', '№', 'кабинет (по договору)', 'сумма'], bank_rows, [9, 11, 8, 16, 14], 'Все поступления от Маркета по реестрам банка 19.10.')
wb.save(out)
print('ok', out)
for r in rows: print(r[:2], [int(x) for x in r[2:]])
print('не определён:', [(d['num'], d['ym'], int(d['amount']), d['base']) for d in docs if d['cab'] == 'не определён'])
