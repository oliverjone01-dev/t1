# ФЕНИКС iter3, проба C3/D1: длинное тире, «{», самооценка вида N/10, «тыс..», «около», красные строки «Все кампании» против текста
# python3 -I -B c3_text_checks.py <xlsx>
import sys, re
from openpyxl import load_workbook
wb = load_workbook(sys.argv[1]); txt = []
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str): txt.append((ws.title, c.coordinate, c.value))
cnt = lambda p: sum(len(re.findall(p, v)) for _, _, v in txt)
print('U+2014:', cnt('—'), ' «{»:', cnt(r'\{'), ' самооценка N/10:', cnt(r'\d[.,]\d+ ?/ ?10\b'), ' «тыс..»:', cnt(r'тыс\.\.'), ' «около»:', cnt(r'\bоколо\b'))
ws = wb['Все кампании']
red = sum(1 for r in ws.iter_rows(min_row=4) if r[0].fill and r[0].fill.fill_type and r[0].fill.fgColor.rgb.endswith('FFC7CE'))
green = sum(1 for r in ws.iter_rows(min_row=4) if r[0].fill and r[0].fill.fill_type and r[0].fill.fgColor.rgb.endswith('C6EFCE'))
other = sorted({r[0].fill.fgColor.rgb for r in ws.iter_rows(min_row=4) if r[0].fill and r[0].fill.fill_type} - {'00FFC7CE', 'FFFFC7CE', '00C6EFCE', 'FFC6EFCE'})
print('«Все кампании»: красных', red, 'зелёных', green, 'другие заливки', other)
for s, k, v in txt:
    if re.search(r'\b\d+ кампани', v): print(f'  {s}!{k}: ...{re.search(r".{0,40}\b\d+ кампани.{0,20}", v).group(0)}')
