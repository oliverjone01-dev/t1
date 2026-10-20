#!/usr/bin/env python3
"""Сверка закрывающего документа Маркета «Отчёт об исполнении поручения» за месяц с продажами в 1С, по заказам. Только чтение.
Запуск: python3 -I sverka-porucheniye.py <корень репо> <xlsx отчёта Маркета> <папка 1С 19.10> <папка 1С 20.10> <выход.xlsx>
Спека (все цифры [ДАННЫЕ]):
  Отчёт Маркета: лист «Получено от потребителей» = платежи покупателей по ДАТЕ ОПЛАТЫ (так написано в самом отчёте), строки товара и доставки,
      статус «Переведён по графику выплат» или «Не будет переведён из-за отмены заказа». Лист «Перечислено на ваш счёт» = платёжки месяца.
  Магазин -> кампания по названию магазина (список кампаний из ping 04.09): DBS ООО новый 148992469, FBS OOO 149001087, DBS OOO 15 дней 149002232,
      DBS 7 дней ООО 149032438, ДБС ООО 1д 149229378, GEN GROUP 66056856.
  Заказ -> статус и дата доставки из выгрузки заказов дашборда (orders.ndjson, тот же API Маркета).
  1С: отчёты комиссионера с кабинетом и магазином (пары как в sverka-otkloneniya.py); документ «есть» для (магазин, месяц доставки).
  Класс каждой строки оплаты: отменён (деньги возвращены) / доставлен в этом месяце и магазин с документом 1С за месяц / доставлен в другом месяце
      и документ за тот месяц есть / доставлен, документа по магазину за месяц доставки нет / ещё не доставлен / заказа нет в выгрузке.
"""
import json, sys, re, glob, collections
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

root, mf, d19, d20, out = sys.argv[1:6]
CAB = {'1023124': 'зеркала', '74986385': 'мебель'}
SHOP2CAMP = {'DBS ООО новый': '148992469', 'FBS OOO': '149001087', 'DBS OOO 15 дней': '149002232', 'DBS 7 дней ООО': '149032438', 'ДБС ООО 1д Москва и область (для конюксов)': '149229378', 'GEN GROUP': '66056856'}
rd = lambda f: [json.loads(l) for l in open(f'{root}/analytics-mvp/data-ym/{f}', encoding='utf-8') if l.strip()]
ymd = lambda s: s[6:10] + '-' + s[3:5]; r2 = lambda x: round(x + 0.0, 2)
orders = [r for r in rd('orders.ndjson') if not r.get('fake')]
camp_cab = {r['campaign']: CAB[r['business']] for r in orders}
goods = collections.defaultdict(float)
for r in orders:
    if not r['service'] and r['status'] == 'DELIVERED' and r.get('fin'): goods[(r['fin'][:7], r['campaign'])] += r['p_buyer'] * r['count']
# заказ -> статус, дата доставки, магазин
ost = {}
for r in orders:
    if r['service']: continue
    ost[r['order']] = dict(status=r['status'], sd=r['statusDate'], fin=r.get('fin') or '', camp=r['campaign'])
# 1С отчёты комиссионера -> (магазин, месяц)
kom = []
for f in glob.glob(f'{d19}/*.xlsx'):
    v = pd.read_excel(f, header=None)
    if 'комиссионера' not in str(v.iloc[1, 0]): continue
    for _, w in v.iloc[6:].iterrows():
        if isinstance(w[1], str) and re.match(r'\d\d\.\d\d\.\d{4}', w[1]): kom.append(dict(ym=ymd(w[1]), num=str(w[3]), amount=float(w[6]), camp=None))
used = set()
for d in sorted(kom, key=lambda d: (d['ym'], -d['amount'])):
    h = [c for (m, c) in goods if m == d['ym'] and (m, c) not in used and abs(goods[(m, c)] - d['amount']) <= 2]
    if h: d['camp'] = h[0]; used.add((d['ym'], h[0]))
for d in sorted(kom, key=lambda d: (d['ym'], -d['amount'])):
    if d['camp']: continue
    cands = [c for (m, c) in goods if m == d['ym'] and (m, c) not in used]
    if cands:
        c = min(cands, key=lambda c: abs(goods[(d['ym'], c)] - d['amount']))
        if abs(goods[(d['ym'], c)] - d['amount']) < 0.25 * d['amount']: d['camp'] = c; used.add((d['ym'], c))
doc_of = {(d['camp'], d['ym']): d for d in kom if d['camp']}
# 1С реализации 20.10 (все 52, кабинет по договору 19.10)
real = pd.read_excel(glob.glob(f'{d20}/*.xls')[0], header=0); old = pd.read_excel(glob.glob(f'{d19}/*.xls')[0], header=0)
CONTRACT = {'54641824/26': 'зеркала', '54542918/26': 'мебель'}
contract = {str(w['Номер']): CONTRACT[str(w['Контрагент'])[-11:]] for _, w in old.iterrows() if str(w['Контрагент']).startswith('ФИЗИЧЕСКОЕ ЛИЦО')}
# отчёт Маркета
head = pd.read_excel(mf, sheet_name=0, header=None)
period = re.search(r'за период (\S+)[–-](\S+)', str(head.iloc[0, 0])); biz = re.search(r'(\d+)', str(head.iloc[2, 0])).group(1); cab = CAB[biz]
mym = '20' + period.group(1).split('.')[2] + '-' + period.group(1).split('.')[1].zfill(2)
tot = {str(head.iloc[i, 0])[:60]: float(head.iloc[i, 2]) for i in range(10, 24) if pd.notna(head.iloc[i, 2])}
def sheet_df(name):
    p = pd.read_excel(mf, sheet_name=name, header=None); h = [str(c) for c in p.iloc[2]]; d = p.iloc[4:].copy(); d.columns = h; return d[d['ID бизнес-аккаунта'].notna()]
p = sheet_df('Получено от потребителей')
rows = []; cls_sum = collections.defaultdict(lambda: [0.0, 0.0, set()])
for _, w in p.iterrows():
    o = str(int(w['Номер заказа или акта об оказанных услугах'])); shop = str(w['Названия магазинов']); camp = SHOP2CAMP.get(shop, '?'); a = float(w['Сумма транзакции, ₽'])
    is_dlv = str(w['Название товара']).strip() == 'Доставка' or str(w['Ваш SKU']).strip() in ('nan', 'Доставка')
    st = str(w['Статус']); x = ost.get(o)
    if 'отмен' in st.lower(): cl = '1. заказ отменён, деньги возвращены покупателю'; dm = ''
    elif not x: cl = '6. заказа нет в выгрузке дашборда'; dm = ''
    elif x['status'] == 'DELIVERED':
        dm = (x['fin'] or x['sd'])[:7]
        if (camp, dm) in doc_of: cl = '2. доставлен в этом месяце, в 1С есть отчёт комиссионера по магазину' if dm == mym else '3. доставлен в другом месяце, в 1С есть отчёт за тот месяц'
        else: cl = '4. доставлен, в 1С нет отчёта комиссионера по магазину за месяц доставки'
    elif x['status'].startswith('CANCELLED'): cl = '1. заказ отменён, деньги возвращены покупателю'; dm = ''
    else: cl = '5. ещё не доставлен'; dm = ''
    rows.append([o, str(w['Дата оформления заказа или акта об оказанных услугах'])[:10], shop, camp, 'доставка' if is_dlv else str(w['Ваш SKU']), a, st, x['status'] if x else '', (x['sd'][:10] if x else ''), dm, doc_of[(camp, dm)]['num'] if dm and (camp, dm) in doc_of else '', cl])
    k = cls_sum[cl]; k[0 if not is_dlv else 1] += a; k[2].add(o)
# обратная сторона: доставленные в месяце отчёта по этому кабинету, оплачены раньше
back = collections.defaultdict(lambda: [0.0, set()])
paid_month = collections.defaultdict(set)
for _, w in p.iterrows(): paid_month[str(int(w['Номер заказа или акта об оказанных услугах']))].add(mym)
for r in orders:
    if r['service'] or r['status'] != 'DELIVERED' or not r.get('fin') or r['fin'][:7] != mym or CAB[r['business']] != cab: continue
    k = 'оплачен в этом месяце' if r['order'] in paid_month else 'оплачен раньше (' + r['created'][:7] + ')'
    back[k][0] += r['p_buyer'] * r['count']; back[k][1].add(r['order'])
k1 = sum(d['amount'] for d in kom if d['ym'] == mym and d['camp'] and camp_cab[d['camp']] == cab)
rl = sum(float(w['Сумма']) for _, w in real.iterrows() if ymd(str(w['Дата'])) == mym and contract.get(str(w['Номер'])) == cab)
tr = sheet_df('Перечислено на ваш счёт'); po = tr.groupby(['Дата платёжного поручения', 'Номер платёжного поручения'])['Сумма платёжного поручения или удерживаемая за услуги сумма, ₽'].first()

wb = Workbook(); bold = Font(bold=True); hdr = PatternFill('solid', fgColor='DDE7F0'); grey = PatternFill('solid', fgColor='EEEEEE')
ws = wb.active; ws.title = 'Итог'; ws.column_dimensions['A'].width = 80; ws.column_dimensions['B'].width = 18; ws.column_dimensions['C'].width = 18; ws.column_dimensions['D'].width = 12
f = lambda x: f'{x:,.0f}'.replace(',', ' ')
L = [[f'Отчёт Маркета об исполнении поручения за {mym}, кабинет {cab} ({biz}), договор {head.iloc[7, 0]}'], [],
     ['Сводка Маркета (первый лист отчёта)', 'сумма, ₽']] + [[k, v] for k, v in tot.items()] + [[],
     ['Платёжки месяца по листу «Перечислено на ваш счёт»', 'сумма, ₽']] + [[f'{d} № {n}', v] for (d, n), v in po.items()] + [[],
     ['«Получено от потребителей» по заказам: куда попали эти деньги', 'товар, ₽', 'доставка, ₽', 'заказов']] + [[k, r2(v[0]), r2(v[1]), len(v[2])] for k, v in sorted(cls_sum.items())] + [
     ['ИТОГО', r2(sum(v[0] for v in cls_sum.values())), r2(sum(v[1] for v in cls_sum.values())), len(set().union(*[v[2] for v in cls_sum.values()]))], [],
     [f'1С за {mym}, {cab}', 'сумма, ₽'], ['отчёты комиссионера (по дате доставки)', r2(k1)], ['реализации «Яндекс маркет фбс»', r2(rl)], [],
     [f'Обратная сторона: доставлено в {mym} по кабинету (это база отчётов комиссионера), когда оплачено', 'товар, ₽', '', 'заказов']] + [[k, r2(v[0]), '', len(v[1])] for k, v in sorted(back.items())] + [
     ['ИТОГО доставлено (товар деньгами)', r2(sum(v[0] for v in back.values())), '', len(set().union(*[v[1] for v in back.values()])) if back else 0]]
for i, row in enumerate(L, 1):
    for j, v in enumerate(row, 1):
        c = ws.cell(i, j, v)
        if isinstance(v, (int, float)): c.number_format = '#,##0;[Red]-#,##0'
        if row and isinstance(row[0], str) and (row[0].startswith('ИТОГО') or (len(row) > 1 and row[1] in ('сумма, ₽', 'товар, ₽'))): c.font = bold
ws2 = wb.create_sheet('Заказы'); H = ['заказ Маркета', 'дата заказа', 'магазин', 'кампания', 'SKU / доставка', 'получено, ₽', 'статус платежа у Маркета', 'статус заказа сейчас', 'дата статуса', 'месяц доставки (сбор за продажу)', '№ отчёта комиссионера 1С', 'куда попали деньги']
for j, h in enumerate(H, 1): c = ws2.cell(1, j, h); c.font = bold; c.fill = hdr; c.alignment = Alignment(wrap_text=True, vertical='top'); ws2.column_dimensions[get_column_letter(j)].width = [14, 11, 16, 11, 20, 12, 26, 20, 11, 12, 12, 60][j - 1]
ws2.row_dimensions[1].height = 45
for i, row in enumerate(sorted(rows, key=lambda r: (r[11], r[0])), 2):
    for j, v in enumerate(row, 1):
        c = ws2.cell(i, j, v)
        if isinstance(v, float): c.number_format = '#,##0;[Red]-#,##0'
ws2.freeze_panes = 'B2'
wb.save(out)
print('ok', out, mym, cab)
for row in L: print(row)
