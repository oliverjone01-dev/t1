#!/usr/bin/env python3
"""1С (отчёты комиссионера и реализации ФБС) против отчёта Маркета о реализации товаров «доставлено минус возвраты», по магазинам и месяцам. Только чтение.
Запуск: python3 -I sverka-realizacii-marketa.py <корень репо> <папка 1С 19.10> <папка 1С 20.10> <выход.xlsx>
Спека (все цифры [ДАННЫЕ]):
  Отчёт Маркета о реализации: данные дашборда по API (goods-realization): realization_orders (строки заказов с магазином, июль-сентябрь) и
      realization_monthly (по SKU, февраль-сентябрь, магазин в поле from, где он есть). Деньги строки = деньги покупателя за товар по выгрузке заказов.
      Нетто = доставлено − возвращено, месяц = месяц отчёта Маркета.
  Покрытие отчёта неполное: по месяцам и магазинам показано, сколько строк и единиц отчёт содержит против доставленных по выгрузке заказов.
  1С ДБС: отчёт комиссионера с магазином (пары: равен доставленному по месяцу сбора, равен отчёту о реализации, ближайший). 1С ФБС: реализации по дате документа.
"""
import json, sys, re, glob, collections
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

root, d19, d20, out = sys.argv[1:5]
CAB = {'1023124': 'зеркала', '74986385': 'мебель'}; CONTRACT = {'54641824/26': 'зеркала', '54542918/26': 'мебель'}
MODEL = {'149001087': 'ФБС', '66056856': 'ФБС', '149154933': 'ФБС'}
NAMES = {'148992469': 'DBS ООО новый', '149001087': 'FBS OOO', '149002232': 'DBS OOO 15 дней', '149032438': 'DBS 7 дней ООО', '149229378': 'ДБС ООО 1д', '66056856': 'GEN GROUP (старый FBS)'}
rd = lambda f: [json.loads(l) for l in open(f'{root}/analytics-mvp/data-ym/{f}', encoding='utf-8') if l.strip()]
ymd = lambda s: s[6:10] + '-' + s[3:5]; r2 = lambda x: round(x + 0.0, 2)
orders = [r for r in rd('orders.ndjson') if not r.get('fake')]; ro = rd('realization_orders.ndjson'); rm = rd('realization_monthly.ndjson')
cat = json.load(open(f'{root}/analytics-mvp/data-ym/catalog.json'))['items']
camp_cab = {r['campaign']: CAB[r['business']] for r in orders}
ok = {(r['order'], r['sku']): r for r in orders if not r['service']}
goods = collections.defaultdict(float); units = collections.Counter()
for r in orders:
    if not r['service'] and r['status'] == 'DELIVERED' and r.get('fin'): goods[(r['fin'][:7], r['campaign'])] += r['p_buyer'] * r['count']; units[(r['fin'][:7], r['campaign'])] += r['count']
# отчёт Маркета по магазинам: июль-сентябрь из строк заказов
M = collections.defaultdict(lambda: collections.defaultdict(float))
for r in ro:
    x = ok.get((r['order'], r['sku']))
    if not x: continue
    k = (r['ym'], r['campaign']); a = x['p_buyer'] * r['count']
    if r['role'] == 'delivered': M[k]['дост'] += a; M[k]['дост_шт'] += r['count']
    else: M[k]['возвр'] += a; M[k]['возвр_шт'] += r['count']
# февраль-июнь из помесячного по SKU (магазин из from, если один)
for r in rm:
    if r['ym'] >= '2026-07': continue
    c = r['from'][0] if len(r['from']) == 1 else ('?' + CAB.get(cat.get(r['sku'], {}).get('business'), '?'))
    k = (r['ym'], c); M[k]['нетто_sku'] += r['amount']; M[k]['дост_шт'] += r['sold']; M[k]['возвр_шт'] += r['ret']
# 1С
kom = []
for f in glob.glob(f'{d19}/*.xlsx'):
    v = pd.read_excel(f, header=None)
    if 'комиссионера' not in str(v.iloc[1, 0]): continue
    for _, w in v.iloc[6:].iterrows():
        if isinstance(w[1], str) and re.match(r'\d\d\.\d\d\.\d{4}', w[1]): kom.append(dict(ym=ymd(w[1]), num=str(w[3]), amount=float(w[6]), camp=None, how=''))
realn = {k: v['дост'] - v['возвр'] for k, v in M.items() if 'дост' in v}
used = set()
for pname, base in (('равен доставленному (месяц сбора)', goods), ('равен отчёту Маркета о реализации', realn)):
    for d in sorted(kom, key=lambda d: (d['ym'], -d['amount'])):
        if d['camp']: continue
        h = [c for (m, c) in base if m == d['ym'] and (m, c) not in used and abs(base[(m, c)] - d['amount']) <= 2]
        if h: d['camp'] = h[0]; d['how'] = pname; used.add((d['ym'], h[0]))
for d in sorted(kom, key=lambda d: (d['ym'], -d['amount'])):
    if d['camp']: continue
    cands = [c for (m, c) in goods if m == d['ym'] and (m, c) not in used]
    if cands:
        c = min(cands, key=lambda c: abs(goods[(d['ym'], c)] - d['amount']))
        if abs(goods[(d['ym'], c)] - d['amount']) < 0.25 * d['amount']: d['camp'] = c; d['how'] = 'ближайший по сумме'; used.add((d['ym'], c))
new = pd.read_excel(glob.glob(f'{d20}/*.xls')[0], header=0); old = pd.read_excel(glob.glob(f'{d19}/*.xls')[0], header=0)
contract = {str(w['Номер']): CONTRACT[str(w['Контрагент'])[-11:]] for _, w in old.iterrows() if str(w['Контрагент']).startswith('ФИЗИЧЕСКОЕ ЛИЦО')}
fbs1c = collections.defaultdict(float)
for _, w in new.iterrows(): fbs1c[(ymd(str(w['Дата'])), contract.get(str(w['Номер']), 'не определён'))] += float(w['Сумма'])
months = ['2026-%02d' % i for i in range(2, 10)]
rows = []
for cab in ('мебель', 'зеркала'):
    camps = sorted(c for c in camp_cab if camp_cab[c] == cab)
    for m in months:
        for c in camps + ['?' + cab]:
            k = (m, c); v = M.get(k); g = goods.get(k, 0.0); u = units.get(k, 0)
            if not v and not g: continue
            fbs = MODEL.get(c) == 'ФБС'
            if c.startswith('?'): net = v['нетто_sku']; doc1c = None; dsum = None; how = 'магазин в отчёте не указан'
            else:
                net = (v['дост'] - v['возвр']) if v and 'дост' in v else (v['нетто_sku'] if v else None)
                if fbs: dsum = fbs1c.get((m, cab), 0.0) if c != '66056856' else None; doc1c = 'реализации ФБС за месяц' if c != '66056856' else ''
                else:
                    dd = [d for d in kom if d['ym'] == m and d['camp'] == c]; dsum = sum(d['amount'] for d in dd) if dd else None; doc1c = ', '.join('№ ' + d['num'] + ' (' + d['how'] + ')' for d in dd)
                how = ''
            rows.append([cab, m, c, NAMES.get(c, '' if c.startswith('?') else 'название в API не видно'), 'ФБС' if fbs else ('' if c.startswith('?') else 'ДБС'),
                         r2(v['дост']) if v and 'дост' in v else None, r2(v['возвр']) if v and 'дост' in v else None, r2(net) if net is not None else None, int(v['дост_шт']) if v else 0, u, r2(g),
                         r2(dsum) if dsum is not None else None, doc1c, r2(dsum - net) if (dsum is not None and net is not None) else None, how])
wb = Workbook(); bold = Font(bold=True); hdr = PatternFill('solid', fgColor='DDE7F0'); red = Font(color='B42318', bold=True); green = Font(color='067647', bold=True)
ws = wb.active; ws.title = 'Как читать'; ws.column_dimensions['A'].width = 200
txt = ['1С против отчёта Маркета о реализации товаров («доставлено» минус «возвраты»), по магазинам и месяцам. Все цифры [ДАННЫЕ].', '',
       'Отчёт Маркета взят из данных дашборда по API. Покрытие неполное, и это главное ограничение: столбцы «отчёт: единиц» и «заказы: доставлено единиц» показывают, какую часть месяца отчёт покрывает.',
       '   Полностью покрыты: август (оба кабинета), сентябрь (мебель: DBS ООО новый, FBS OOO, DBS 7 дней). Июль: только FBS OOO и 149154933. Май-июнь: только FBS OOO (и 149154933 в июне). Февраль-апрель: отчёт без магазина, около половины единиц.',
       '   Там, где отчёт есть, отчёты комиссионера 1С по ДБС равны ему: август DBS OOO 15 дней (№ 74), DBS 7 дней (№ 75), сентябрь DBS ООО новый (№ 83), DBS 7 дней (№ 84) - до рубля «доставлено минус возвраты»; август DBS ООО новый (№ 73) равен «доставлено» без вычета возврата 17 013 ₽ (заказ 57177849219, возврат майского заказа).',
       '   Реализации ФБС 1С отчёту не равны ни в одном месяце: май -1 340 к документу № 36 (отчёт комиссионера на FBS OOO) и отдельно реализации 709 509; июнь -31 995; июль -12 363; август -49 947; сентябрь -61 172.',
       '', 'Чтобы сравнить все месяцы по всем магазинам, нужны сами отчёты Маркета о реализации за февраль-сентябрь по кабинету мебель (файлы из кабинета: Финансы - Отчёты - Отчёт о реализации), либо добор отчётов сборщиком (несколько прогонов по часу).']
for i, t in enumerate(txt, 1): ws.cell(i, 1, t)
ws2 = wb.create_sheet('По магазинам'); H = ['кабинет', 'месяц', 'магазин', 'название', 'модель', 'отчёт Маркета: доставлено, ₽', 'отчёт: возвраты, ₽', 'отчёт: нетто, ₽', 'отчёт: единиц', 'заказы: доставлено единиц (месяц сбора)', 'заказы: доставлено, ₽ (месяц сбора)', '1С: сумма', '1С: документ', 'разница 1С − отчёт', 'примечание']
for j, h in enumerate(H, 1): c = ws2.cell(1, j, h); c.font = bold; c.fill = hdr; c.alignment = Alignment(wrap_text=True, vertical='top'); ws2.column_dimensions[get_column_letter(j)].width = [9, 9, 11, 18, 7, 14, 12, 14, 9, 12, 14, 14, 40, 14, 26][j - 1]
ws2.row_dimensions[1].height = 60
for i, row in enumerate(rows, 2):
    for j, v in enumerate(row, 1):
        c = ws2.cell(i, j, v)
        if isinstance(v, float): c.number_format = '#,##0;[Red]-#,##0'
        if j == 14 and isinstance(v, float): c.font = green if abs(v) <= 2 else red
ws2.freeze_panes = 'C2'
wb.save(out)
print('ok', out)
for r in rows: print(r[:5], [int(x) if isinstance(x, float) else x for x in r[5:14]])
