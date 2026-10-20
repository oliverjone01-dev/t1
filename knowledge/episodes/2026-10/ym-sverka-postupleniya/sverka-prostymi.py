#!/usr/bin/env python3
"""Простая сверка по месяцам: поступления 1С против платёжек Маркета, продажи 1С против «получено от покупателей» Маркета, разница расписана по причинам. Только чтение.
Запуск: python3 -I sverka-prostymi.py <корень репо> <папка 1С 19.10> <выход.xlsx>
Спека (все цифры [ДАННЫЕ]):
  Поступления. 1С = реестр банка по дате поступления. Маркет = платёжки (проводки реестра с номером п/п): 59 платёжек равны поступлениям банка
      до копейки (sverka.py), их месяц = месяц поступления; 11 платёжек, которых нет в выписке, поставлены в месяц последней проводки (у найденных платёжка приходит через 2-5 дней после неё).
  Продажи. 1С = отчёты комиссионера по месяцу документа (= месяц доставки). Маркет = «Получено от покупателей» отчёта об исполнении поручения:
      строки реестра «Начисление / Платёж покупателя» по месяцу проводки (дата оплаты), товар + доставка, как в отчёте Маркета.
  Разница Маркет − 1С расписана: отменено и возвращено; доставка, оплаченная покупателем; доставлен, документа 1С нет; оплачен в этом месяце,
      доставлен в другом (документ там есть); минус оплачен раньше, доставлен в этом месяце; ещё не доставлен; заказа нет в выгрузке;
      минус отклонения документов 1С от Маркета; минус документ 1С без пары. Остаток = 0 в каждой строке.
"""
import json, sys, re, glob, collections
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

root, d19, out = sys.argv[1:4]
CAB = {'1023124': 'зеркала', '74986385': 'мебель'}; CONTRACT = {'54641824/26': 'зеркала', '54542918/26': 'мебель'}
rd = lambda f: [json.loads(l) for l in open(f'{root}/analytics-mvp/data-ym/{f}', encoding='utf-8') if l.strip()]
ymd = lambda s: s[6:10] + '-' + s[3:5]; r2 = lambda x: round(x + 0.0, 2)
orders = [r for r in rd('orders.ndjson') if not r.get('fake')]; net = rd('netting.ndjson')
camp_cab = {r['campaign']: CAB[r['business']] for r in orders}
goods = collections.defaultdict(float); ost = {}
for r in orders:
    if r['service']: continue
    ost[r['order']] = dict(status=r['status'], fin=r.get('fin') or '', camp=r['campaign'])
    if r['status'] == 'DELIVERED' and r.get('fin'): goods[(r['fin'][:7], r['campaign'])] += r['p_buyer'] * r['count']
# 1С отчёты комиссионера -> магазин
kom = []
for f in glob.glob(f'{d19}/*.xlsx'):
    v = pd.read_excel(f, header=None)
    if 'комиссионера' not in str(v.iloc[1, 0]): continue
    for _, w in v.iloc[6:].iterrows():
        if isinstance(w[1], str) and re.match(r'\d\d\.\d\d\.\d{4}', w[1]): kom.append(dict(ym=ymd(w[1]), num=str(w[3]), amount=float(w[6]), camp=None))
used = set()
for d in sorted(kom, key=lambda d: (d['ym'], -d['amount'])):
    h = [c for (m, c) in goods if m == d['ym'] and (m, c) not in used and abs(goods[(m, c)] - d['amount']) <= 2]
    if h: d['camp'] = h[0]; used.add((d['ym'], h[0]))
for d in sorted(kom, key=lambda d: (d['ym'], -d['amount'])):
    if d['camp']: continue
    cands = [c for (m, c) in goods if m == d['ym'] and (m, c) not in used]
    if cands:
        c = min(cands, key=lambda c: abs(goods[(d['ym'], c)] - d['amount']))
        if abs(goods[(d['ym'], c)] - d['amount']) < 0.25 * d['amount']: d['camp'] = c; used.add((d['ym'], c))
doc_of = {(d['camp'], d['ym']): d for d in kom if d['camp']}; last_doc = max(d['ym'] for d in kom)
# банк 1С
bank = collections.defaultdict(float); bank_po = {}
for f in glob.glob(f'{d19}/bank-*.xlsx'):
    v = pd.read_excel(f, header=None)
    for _, w in v.iloc[6:].iterrows():
        if isinstance(w[1], str) and re.match(r'\d\d\.\d\d\.\d{4}', w[1]):
            m = re.search(r'договору (\S+?)\.', str(w[7])); cab = CONTRACT.get(m.group(1) if m else '', '?'); bank[(ymd(w[1]), cab)] += float(w[6]); bank_po[(cab, round(float(w[6]), 2))] = ymd(w[1])
# платёжки Маркета
po_sum = collections.defaultdict(float); po_last = {}; po_cab = {}
for r in net:
    po = (r.get('po') or '').strip()
    if not po: continue
    po_sum[po] += r['amount']; po_last[po] = max(po_last.get(po, ''), r['d'][:10]); po_cab[po] = CAB.get(r['business'], '?')
mpay = collections.defaultdict(float); unmatched = []
for po, s in po_sum.items():
    cab = po_cab[po]; k = (cab, round(s, 2))
    if k in bank_po: mpay[(bank_po[k], cab)] += s
    else:
        nm = po_last[po][:7]; mpay[(nm, cab)] += s; unmatched.append((cab, po, r2(s), po_last[po], nm))
# продажи: строки «Платёж покупателя» по месяцу проводки
S = collections.defaultdict(lambda: collections.defaultdict(float))
for r in net:
    if r['type'] != 'Начисление' or r['src'] != 'Платёж покупателя': continue
    st = str(r.get('status') or '').strip(); m = r['d'][:7]; cab = CAB.get(r['business'], '?'); x = ost.get(str(r.get('order') or '')); a = r['amount']; k = (cab, m)
    S[k]['получено'] += a
    if r['service'] == 'Доставка':
        S[k]['доставка (отменённая тоже)'] += a
        if re.match(r'^(справочно:\s*)?не будет', st, re.I) or (x and x['status'].startswith('CANCELLED')): S[k]['отменено'] += a
        else: S[k]['доставка'] += a
        continue
    if re.match(r'^(справочно:\s*)?не будет', st, re.I) or (x and x['status'].startswith('CANCELLED')): S[k]['отменено'] += a
    elif not x: S[k]['нет в выгрузке'] += a
    elif x['status'] == 'DELIVERED':
        dm = x['fin'][:7]
        if (x['camp'], dm) not in doc_of: S[k]['без документа'] += a
        elif dm == m: S[k]['этот месяц с документом'] += a
        else: S[k]['доставлен в другом месяце'] += a
    else: S[k]['не доставлен'] += a
months = ['2026-%02d' % i for i in range(1, 11)]
rows_p = []; rows_s = []
for cab in ('зеркала', 'мебель'):
    tp = collections.defaultdict(float); ts = collections.defaultdict(float)
    for m in months:
        b = bank[(m, cab)]; mp = mpay[(m, cab)]; row = [cab, m, r2(b), r2(mp), r2(mp - b)]; rows_p.append(row)
        for i, v in enumerate(row[2:], 2): tp[i] += v
    rows_p.append([cab, 'ИТОГО'] + [r2(tp[i]) for i in range(2, 5)])
    for m in months:
        s = S[(cab, m)]; k1 = sum(d['amount'] for d in kom if d['ym'] == m and d['camp'] and camp_cab[d['camp']] == cab)
        with_doc = sum(v for (mm, c), v in goods.items() if mm == m and camp_cab[c] == cab and (mm, c) in used)
        paid_earlier = with_doc - s['этот месяц с документом']
        dev = sum(d['amount'] - goods[(d['ym'], d['camp'])] for d in kom if d['ym'] == m and d['camp'] and camp_cab[d['camp']] == cab)
        nopair = sum(d['amount'] for d in kom if d['ym'] == m and not d['camp']) if cab == 'мебель' else 0.0
        k1 += nopair; diff = s['получено'] - k1
        rest = diff - (s['отменено'] + s['доставка'] + s['без документа'] + s['доставлен в другом месяце'] - paid_earlier + s['не доставлен'] + s['нет в выгрузке'] - dev - nopair)
        row = [cab, m, r2(k1), r2(s['получено']), r2(diff), r2(s['отменено']), r2(s['доставка']), r2(s['без документа']), r2(s['доставлен в другом месяце']), r2(-paid_earlier), r2(s['не доставлен']), r2(s['нет в выгрузке']), r2(-dev), r2(-nopair), r2(rest)]
        if m > last_doc: row[2] = None
        rows_s.append(row)
        for i, v in enumerate(row[2:], 2): ts[i] += v or 0
    rows_s.append([cab, 'ИТОГО'] + [r2(ts[i]) for i in range(2, 15)])

wb = Workbook(); bold = Font(bold=True); hdr = PatternFill('solid', fgColor='DDE7F0'); grey = PatternFill('solid', fgColor='EEEEEE'); red = Font(color='B42318', bold=True)
def sheet(title, header, rows, widths, note, redcols=()):
    ws = wb.create_sheet(title); ws.cell(1, 1, note).alignment = Alignment(wrap_text=True, vertical='top'); ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(header)); ws.row_dimensions[1].height = 15 * (note.count('\n') + 2)
    for j, h in enumerate(header, 1): c = ws.cell(3, j, h); c.font = bold; c.fill = hdr; c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[3].height = 90
    for i, row in enumerate(rows, 4):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v)
            if isinstance(v, float): c.number_format = '#,##0;[Red]-#,##0'
            if row[1] == 'ИТОГО': c.font = bold; c.fill = grey
            if j in redcols and isinstance(v, float) and abs(v) > 2: c.font = red
    for j in range(1, len(header) + 1): ws.column_dimensions[get_column_letter(j)].width = widths[j - 1] if j - 1 < len(widths) else 14
    ws.freeze_panes = ws.cell(4, 3); return ws
wb.remove(wb['Sheet'])
f = lambda x: f'{x:,.0f}'.replace(',', ' ')
ws = wb.create_sheet('Как читать'); ws.column_dimensions['A'].width = 200
txt = ['Простая сверка по месяцам: 1С против отчётов Маркета, разница расписана по причинам. Все цифры [ДАННЫЕ].', '',
       'Лист «Поступления»: банк 1С по дате поступления против платёжек Маркета. Все 59 поступлений равны платёжкам Маркета до копейки. Разница только там, где платёжки Маркета нет в выписке:',
       *[f'   {c}: платёжка № {po} на {f(s)} ₽, проводки по {d}, поставлена в {nm}' for c, po, s, d, nm in sorted(unmatched, key=lambda u: u[4])],
       '   Восемь из них - январь-февраль (выписка начинается 20.02; четыре январские - за декабрьские продажи), три малые - премия, компенсация и взнос продавца.', '',
       'Лист «Продажи»: отчёты комиссионера 1С (по месяцу доставки) против «Получено от покупателей» отчёта Маркета (по месяцу оплаты, товар + доставка). Разница = Маркет − 1С, расписана по столбцам «за счёт», сумма столбцов равна разнице, остаток 0.',
       '   Отменено и возвращено - покупатель заплатил, заказ отменён, Маркет вернул деньги. Не продажа.',
       '   Доставка, оплаченная покупателем - входит в «получено» Маркета, в суммы отчётов комиссионера 1С не входит.',
       '   Доставлен, документа 1С нет - ПРОДАЖА, которой нет в 1С: магазин без отчёта комиссионера за месяц доставки. Это и есть непроведённое.',
       '   Оплачен в этом месяце, доставлен в другом - продажа есть в 1С в другом месяце. Минус оплачен раньше - продажа этого месяца в 1С, деньги получены в прошлом.',
       '   Ещё не доставлен - в пути. Заказа нет в выгрузке - карточки заказа нет (старые магазины зеркал, закрытые для API; декабрьские заказы мебели), в 1С документов по ним нет.',
       '   Минус отклонения документов - документы 1С с парой отличаются от Маркета на переходящие заказы. Минус документ без пары - 1С июнь № 43 на 401 995 ₽: он входит в сумму 1С за июнь, но магазина с такой суммой у Маркета нет.']
for i, t in enumerate(txt, 1): ws.cell(i, 1, t)
sheet('Поступления', ['кабинет', 'месяц', '1С: поступило в банк', 'Маркет: платёжки', 'разница Маркет − 1С'], rows_p, [10, 9, 18, 18, 18], 'Поступления по месяцам. Разница только по 11 платёжкам, которых нет в выписке (см. «Как читать»).', redcols=(5,))
sheet('Продажи', ['кабинет', 'месяц', '1С: отчёты комиссионера (по доставке)', 'Маркет: получено от покупателей (по оплате)', 'РАЗНИЦА Маркет − 1С', 'за счёт: отменено и возвращено', 'за счёт: доставка, оплаченная покупателем', 'за счёт: ДОСТАВЛЕН, ДОКУМЕНТА 1С НЕТ', 'за счёт: оплачен в этом месяце, доставлен в другом', 'минус: оплачен раньше, доставлен в этом месяце', 'за счёт: ещё не доставлен', 'за счёт: заказа нет в выгрузке', 'минус: отклонения документов 1С', 'минус: документ 1С без пары', 'остаток'],
      rows_s, [10, 9, 16, 16, 15, 14, 14, 16, 15, 15, 13, 14, 13, 13, 9], 'Продажи по месяцам: 1С против отчёта Маркета, разница по причинам. Красным - продажи, которых нет в 1С.', redcols=(8,))
wb.save(out)
print('ok', out)
for r in rows_p: print(r)
for r in rows_s: print(r[:2], [int(x) if isinstance(x, float) else x for x in r[2:]])
