# ФЕНИКС iter2 «Рост показов»: доп. пробы - задержка отмен, «плохо с проверкой» при счёте с отменами, текст книги.
# python3 -I -B probe_iter2b.py <OZON-plan.xlsx>
import sys, json, collections as C, io, contextlib, re
V7 = '/home/user/t1/knowledge/episodes/2026-10/ozon-testy-scripts/v7'
sys.path.insert(0, V7); sys.path.insert(0, V7 + '/../v5')
with contextlib.redirect_stdout(io.StringIO()):
    import cards7 as K, forecast as F
from openpyxl import load_workbook
LV = 20000; MO = 4.33; top = K.cards[:10]
mem = {s['sku']: g['main'] for g in K.G for s in g['skus']}
cards10 = {c['card'] for c in top}
print('== Q1 задержка отмен: доля отменённых штук по неделе заказа (весь магазин и 10 карточек), снимок по 12.10')
import datetime as dt
wk = C.defaultdict(lambda: [0, 0, 0, 0])
for o in F.L('orders_daily.ndjson'):
    if o['d'] < '2026-08-03': continue
    w = (dt.date.fromisoformat(o['d']) - dt.date(2026, 8, 3)).days // 7
    r = wk[w]; r[0] += o['units']; r[1] += o['units'] if o['status'] == 'cancelled' else 0
    if mem.get(o['sku']) in cards10: r[2] += o['units']; r[3] += o['units'] if o['status'] == 'cancelled' else 0
for w in sorted(wk):
    r = wk[w]; d0 = dt.date(2026, 8, 3) + dt.timedelta(days=7 * w)
    print(f'  неделя с {d0:%d.%m}: магазин {r[1]:3d}/{r[0]:3d} = {r[1]/r[0]:.0%}   10 карточек {r[3]:3d}/{r[2]:3d} = {(r[3]/r[2] if r[2] else 0):.0%}')
st = C.Counter(o['status'] for o in F.L('orders_daily.ndjson') if '2026-10-05' <= o['d'] <= '2026-10-12')
print('  статусы заказов 05-12.10:', dict(st))

print('== Q2 «плохо с проверкой», если 29.10 считают заказы с отменами (доля отмен базы по карточке)')
canc = C.Counter(); allu = C.Counter()
for o in F.L('orders_daily.ndjson'):
    cd = mem.get(o['sku'])
    if cd in cards10 and F.A <= o['d'] <= F.B:
        allu[cd] += o['units']; canc[cd] += o['units'] if o['status'] == 'cancelled' else 0
loss = 0; loss0 = 0
for c in top:
    _, _, sp, _ = K.main_effect(c, LV, 'низкий'); _, _, sp2, _ = K.min_effect(c, 'низкий'); w = sp + sp2
    k = K.threshold(c); lam = 2 * c['u_wk']; sh = canc[c['card']] / allu[c['card']]
    p = K.pois_sf(k, lam / (1 - sh)); p0 = K.pois_sf(k, lam)
    loss += w * 2 + p * w * (MO - 2); loss0 += w * 2 + p0 * w * (MO - 2)
    print(f'  {c["main_art"]:20} расход сверх, ₽/нед {w:7,.0f}  P0 {p0:.2f} -> с отменами {p:.2f}'.replace(',', ' '))
print(f'  «плохо с проверкой»: как в книге {-loss0/1e3:+.0f} тыс., с отменами {-loss/1e3:+.0f} тыс.')
print(f'  срок до 26.11 / 03.12 у ложно оставленных: ожидаемо {sum(K.pois_sf(K.threshold(c), 2*c["u_wk"]) * (K.main_effect(c, LV, "низкий")[2] + K.min_effect(c, "низкий")[2]) for c in top) * 4 / 1e3:.0f} тыс. за 4 недели 29.10-26.11')

print('== Q3 текст книги: U+2014, шаблоны «{», «около», «волна 2», «10.11», «итог», «Эластичн», «без отмен», «7 кампаний»')
wb = load_workbook(sys.argv[1]); cnt = C.Counter()
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if not isinstance(v, str): continue
            for key in (chr(0x2014), '{', 'около', 'олна 2', '10.11', 'итог', 'Эластичн', 'без отмен', '7 кампаний', 'отмен'):
                if key in v: cnt[(ws.title, key)] += 1
            if 'олна 2' in v or '10.11' in v or 'без отмен' in v or 'Эластичн' in v:
                print(f'  [{ws.title}!{c.coordinate}] {v[:170]}')
print(' ', dict(cnt))
red = sum(1 for r in wb['Все кампании'].iter_rows(values_only=True) if r[0] and 'на минимум' in str(r[0]))
print('  строк «на минимум 15.10» в «Все кампании»:', red)
