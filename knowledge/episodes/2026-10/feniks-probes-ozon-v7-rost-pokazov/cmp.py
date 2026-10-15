import sys
from openpyxl import load_workbook
a = load_workbook(sys.argv[1]); b = load_workbook(sys.argv[2])
print('sheets', a.sheetnames, b.sheetnames)
diff = 0
for sn in a.sheetnames:
    wa, wb = a[sn], b[sn]
    for r in range(1, max(wa.max_row, wb.max_row) + 1):
        for c in range(1, max(wa.max_column, wb.max_column) + 1):
            x, y = wa.cell(r, c).value, wb.cell(r, c).value
            if x != y:
                diff += 1
                if diff < 20: print('DIFF', sn, r, c, repr(x)[:80], repr(y)[:80])
print('cells differing:', diff)
if len(sys.argv) > 3:
    for sn in a.sheetnames:
        ws = a[sn]
        print('=====', sn, ws.max_row, 'x', ws.max_column)
        for row in ws.iter_rows(values_only=True):
            vals = [str(v) for v in row if v not in (None, '')]
            if vals: print(' | '.join(vals))
