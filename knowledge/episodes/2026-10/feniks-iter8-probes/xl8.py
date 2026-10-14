#!/usr/bin/env python3
"""Real September CPO export: header rows, Период line, rows per order, row date vs earliest posting date."""
import sys, json, collections
sys.path.insert(0, "/tmp/claude-0/-home-user-t1/c4f0dfeb-cd9a-55ac-b7f6-154069ba26c7/scratchpad/i8/p8/analytics-mvp/tools/cpo")
import openpyxl
from cpo_raw import read_file

X = "/tmp/claude-0/-home-user-t1/c4f0dfeb-cd9a-55ac-b7f6-154069ba26c7/scratchpad/i8/p8/analytics-mvp/data/raw/ozon-cpo/2026-09_cpo_orders.xlsx"
D = "/tmp/claude-0/-home-user-t1/c4f0dfeb-cd9a-55ac-b7f6-154069ba26c7/scratchpad/i8/p8/analytics-mvp/data"
wb = openpyxl.load_workbook(X, read_only=True, data_only=True)
ws = wb["Statistics"] if "Statistics" in wb.sheetnames else wb.worksheets[0]
rows = list(ws.iter_rows(values_only=True))
wb.close()
for i, r in enumerate(rows[:4]):
    print("hdr", i, [c for c in r if c is not None][:14])
x = read_file(X)
print("win", x["win"], "rows", len(x["rows"]), "sum sp", round(sum(r["sp"] for r in x["rows"]), 2))
by = collections.defaultdict(list)
for r in x["rows"]:
    by[r["order"]].append(r)
multi = {k: v for k, v in by.items() if len(v) > 1}
multid = {k: v for k, v in by.items() if len({r["d"] for r in v}) > 1}
neg = [r for r in x["rows"] if r["sp"] < 0]
noord = [r for r in x["rows"] if not r["order"]]
print("orders", len(by), "orders with >1 row", len(multi), "orders with >1 date", len(multid), "negative rows", len(neg), "rows without order", len(noord))
# earliest posting date per order base
post = collections.defaultdict(list)
for l in open(D + "/orders_daily.ndjson"):
    if l.strip():
        o = json.loads(l)
        k = str(o["order"])
        post[k[: k.rfind("-")]].append(str(o["d"]))
diff = collections.Counter()
ex = []
for k, v in by.items():
    rd = min(r["d"] for r in v)
    if k not in post:
        diff["no posting"] += 1
        continue
    pd = min(post[k])
    delta = (__import__("datetime").date.fromisoformat(rd) - __import__("datetime").date.fromisoformat(pd)).days
    diff[delta] += 1
    if delta != 0 and len(ex) < 5:
        ex.append((k, rd, pd))
print("row date - earliest posting date (days):", sorted(diff.items(), key=lambda t: str(t[0])))
print("examples", ex)
# orders whose earliest posting is in August but charged in September export
aug = [k for k, v in by.items() if k in post and min(post[k]) < "2026-09-01"]
print("orders in Sep export with order date before 01.09:", len(aug), "sp", round(sum(sum(r["sp"] for r in by[k]) for k in aug), 2))
