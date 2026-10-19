# ФЕНИКС iter3, проба G1: поиск фраз по всем листам книги (ячейка, лист), обрезка до 160 символов
# python3 -I -B g1_grep_book.py <xlsx> <regex> [<regex> ...]
import sys, re
from openpyxl import load_workbook
wb = load_workbook(sys.argv[1])
pats = [re.compile(p, re.I) for p in sys.argv[2:]]
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and any(p.search(c.value) for p in pats):
                print(f'{ws.title}!{c.coordinate}: {c.value[:160]}')
