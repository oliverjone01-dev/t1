#!/usr/bin/env python3
"""Independent parse of the CPO cabinet xlsx (tolerant parser) vs data/cpo_card_orders.ndjson and main's cpo_orders.ndjson."""
import json, sys, collections, re, datetime
import openpyxl
root = sys.argv[1]
f = root + "/data/raw/ozon-cpo/2026-09_cpo_orders.xlsx"
wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
print("sheets", wb.sheetnames)
ws = wb.worksheets[0]
rows = list(ws.iter_rows(values_only=True))
print("rows total", len(rows))
for i, r in enumerate(rows[:4]):
    print("row", i, [c for c in r][:20])
hi = next(i for i, r in enumerate(rows[:6]) if r and "Номер заказа" in [str(c or "") for c in r])
H = [str(h or "") for h in rows[hi]]
print("header idx", hi, H)
types = collections.defaultdict(collections.Counter)
data = []
for r in rows[hi + 1:]:
    if not r or not any(r):
        continue
    rec = dict(zip(H, r))
    for k, v in rec.items():
        types[k][type(v).__name__] += 1
    data.append(rec)
print("data rows", len(data))
for k in H:
    print("  type hist", k, dict(types[k]))


def num(v):
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(" ", "").replace(" ", "").replace(" ", "").replace(",", ".")
    s = re.sub(r"[^0-9.\-]", "", s)
    return float(s) if s else 0.0


# rows without order number
noord = [d for d in data if not d.get("Номер заказа")]
print("rows without order number", len(noord), [ (d.get("Дата"), d.get("Расход, ₽"), d.get("Стоимость продажи, ₽")) for d in noord[:5]])
tot_rev = sum(num(d.get("Стоимость продажи, ₽")) for d in data if d.get("Номер заказа"))
tot_sp = sum(num(d.get("Расход, ₽")) for d in data if d.get("Номер заказа"))
tot_n = sum(num(d.get("Количество")) for d in data if d.get("Номер заказа"))
tot_sp_all = sum(num(d.get("Расход, ₽")) for d in data)
orders = {str(d.get("Номер заказа")) for d in data if d.get("Номер заказа")}
print("xlsx (rows with order): rev %.2f sp %.2f n %d orders %d | sp incl rows w/o order %.2f" % (tot_rev, tot_sp, tot_n, len(orders), tot_sp_all))
dates = sorted({str(d.get("Дата")) for d in data})
print("dates", dates[:3], "...", dates[-3:], len(dates))
# by date window 1-29
def dkey(v):
    s = str(v)
    if "." in s:
        dd, mm, yy = s.split(".")
        return f"{yy}-{mm}-{dd}"
    return s[:10]
w29 = [d for d in data if d.get("Номер заказа") and dkey(d.get("Дата")) <= "2026-09-29"]
print("1-29.09: rev %.2f sp %.2f n %d" % (sum(num(d.get("Стоимость продажи, ₽")) for d in w29), sum(num(d.get("Расход, ₽")) for d in w29), sum(num(d.get("Количество")) for d in w29)))
# compare with converter output
nd = [json.loads(l) for l in open(root + "/data/cpo_card_orders.ndjson", encoding="utf-8") if l.strip()]
print("ndjson rows", len(nd), "rev %.2f sp %.2f n %d" % (sum(x["rev"] for x in nd), sum(x["sp"] for x in nd), sum(x["n"] for x in nd)))
# multiset compare
A = collections.Counter((dkey(d.get("Дата")), str(d.get("Номер заказа")), str(d.get("SKU")), round(num(d.get("Стоимость продажи, ₽")), 2), int(num(d.get("Количество")))) for d in data if d.get("Номер заказа"))
B = collections.Counter((x["d"], x["order"], x["sku"], x["rev"], x["n"]) for x in nd)
print("multiset equal xlsx vs ndjson:", A == B, "A-B", list((A - B).items())[:5], "B-A", list((B - A).items())[:5])
# duplicates inside xlsx (same row twice)
dup = [k for k, c in A.items() if c > 1]
print("identical duplicate rows inside xlsx:", len(dup), dup[:5])
by_ord = collections.Counter(str(d.get("Номер заказа")) for d in data if d.get("Номер заказа"))
print("orders with >1 row:", sum(1 for v in by_ord.values() if v > 1))
# main's cpo_orders.ndjson (if present)
try:
    co = [json.loads(l) for l in open(root + "/data/cpo_orders.ndjson", encoding="utf-8") if l.strip()]
    sep = [x for x in co if str(x.get("d", x.get("date", "")))[:7] == "2026-09"]
    print("cpo_orders.ndjson keys", list(co[0].keys()), "rows", len(co), "sep rows", len(sep))
    spk = [k for k in co[0].keys() if k in ("sp", "cpo", "amount", "sum", "spend")]
    print("  sep sum by", spk, [round(sum(float(x.get(k) or 0) for x in sep), 2) for k in spk])
except FileNotFoundError:
    print("no cpo_orders.ndjson")
