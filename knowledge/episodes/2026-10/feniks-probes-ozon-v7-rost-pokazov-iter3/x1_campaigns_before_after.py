# ФЕНИКС iter3, проба X1: команды «Все кампании» до склейки (лист «Клик» из plan7_xlsx) и после (merge_xlsx), по номеру кампании
# python3 -I -B x1_campaigns_before_after.py <v1.xlsx из plan7_xlsx> <поданный.xlsx>
import sys
from openpyxl import load_workbook
def rows(ws):
    out = {}; tail = []
    for row in ws.iter_rows():
        b = str(row[1].value or '') if len(row) > 1 else ''
        fill = row[0].fill.fgColor.rgb[-6:] if row[0].fill and row[0].fill.fill_type else '-'
        if b.isdigit():
            out.setdefault(b, []).append(dict(r=row[0].row, a=str(row[0].value), fill=fill, j=row[9].value if len(row) > 9 else None,
                                              k=str(row[10].value) if len(row) > 10 else '', c=row[2].value, e=row[4].value, g=row[6].value))
        elif row[0].value is not None:
            tail.append((row[0].row, fill, str(row[0].value)[:150]))
    return out, tail
before, tb = rows(load_workbook(sys.argv[1])['Клик'])
after, ta = rows(load_workbook(sys.argv[2])['Все кампании'])
print('нецифровые строки ДО склейки:')
for t in tb: print('  ', t)
print('нецифровые строки ПОСЛЕ:')
for t in ta: print('  ', t)
print('кампаний до:', len(before), 'после:', len(after))
print('только до (удалены):', sorted(set(before) - set(after)))
print('только после:', sorted(set(after) - set(before)))
print('изменилась команда:')
for cid in before:
    for i, b in enumerate(before[cid]):
        a = after.get(cid, [None] * (i + 1))[i] if cid in after and i < len(after[cid]) else None
        ba = b['a'][:45]; aa = a['a'][:45] if a else 'УДАЛЕНА'
        if not a or ba != aa or b['fill'] != a['fill']:
            print(f"  {cid} {b['c']} r{b['r']}: [{b['fill']}] {ba} | J={b['j']} | K={b['k'][:70]}")
            print(f"  {'':>8} -> " + (f"r{a['r']}: [{a['fill']}] {aa} | J={a['j']}" if a else 'УДАЛЕНА'))
dup = [cid for cid, v in before.items() if len(v) > 1]
print('кампании, встречающиеся в листе до склейки больше одного раза:', dup)
