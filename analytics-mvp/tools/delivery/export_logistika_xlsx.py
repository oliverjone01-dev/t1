#!/usr/bin/env python3
# Выгрузка блока «Доставка по городам» в Excel: один лист на месяц.
# Запуск из analytics-mvp:  python3 tools/delivery/export_logistika_xlsx.py 2026-07 2026-08
# Без аргументов берёт июль и август 2026.
#
# Считает ровно то же, что блок на странице «Деньги» (ctDraw в build-katya.ts), и тем же правилом:
#   итог = доход с покупателя − счёт нашего перевозчика − сбор Маркета за логистику,
# по ЗАКАЗУ, потому что город и перевозка - свойства заказа, а не позиции. Режим заказа один:
# либо везёт Маркет (есть его сбор за логистику), либо везём мы (платёж покупателя нам или наш
# счёт перевозчика). Проверено: строк с обоими признаками в снимке нет, смешанных заказов нет.
#
# Сверка выгрузки со сводом (снимок 22.09.2026): июль 200 заказов, доход 323 199, перевозчик
# 366 991, сбор 255 058, итог −298 849; август 158 / 351 600 / 378 514 / 126 966 / −153 880.
# Оба листа сходятся с svod_orders.json до рубля.
import json, sys, calendar
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

DEL=["Доставка покупателю","Доставка (средняя миля)","Доставка невыкупов и возвратов"]
sv=json.load(open('data-ym/svod_orders.json',encoding='utf-8'))
months=sv.get('months',sv)

def collect(frm,to):
    byo={}
    for m in months:
        for r in m.get('rows',[]):
            d=str(r.get('d',''))
            if d<frm or d>to: continue
            k=r.get('order','—')
            o=byo.setdefault(k,{'d':d,'city':'','inc':0.0,'our':0.0,'fee':0.0,'kn':False,'mode':''})
            if d<o['d']: o['d']=d
            if not o['city'] and r.get('region'): o['city']=r['region']
            o['inc']+=r.get('ship_buyer',0) or 0
            o['our']+=r.get('ship_our',0) or 0
            led=sum(((r.get('svc') or {}).get(x,0) or 0)+((r.get('svc_pts') or {}).get(x,0) or 0) for x in DEL)
            o['fee']+=led
            o['kn']=o['kn'] or bool(r.get('ship_known'))
            md='mk' if led>0 else ('own' if ((r.get('ship_buyer',0) or 0)>0 or (r.get('ship_our',0) or 0)>0 or r.get('ship_known')) else '')
            if md=='mk': o['mode']='mk'
            elif md=='own' and o['mode']!='mk': o['mode']='own'
    return byo

HDR=['Город','Кто везёт','Заказов','Доход с покупателя','Наш перевозчик','Сбор Маркета','Итог по доставке','На заказ','Заказов без счёта перевозчика']
HFILL=PatternFill('solid',fgColor='1F2937'); HFONT=Font(color='FFFFFF',bold=True,size=10)
TFONT=Font(bold=True); RED=Font(color='B00020'); GRN=Font(color='0B6B3A')
THIN=Side(style='thin',color='D0D5DD'); BRD=Border(bottom=THIN)
MODE={'mk':'везёт Маркет','own':'везём мы','':'не определено'}

def sheet(wb,title,frm,to):
    ws=wb.create_sheet(title)
    byo=collect(frm,to)
    cities=defaultdict(lambda: defaultdict(lambda:{'inc':0.0,'our':0.0,'fee':0.0,'ord':0,'noVed':0}))
    tot=defaultdict(lambda:{'inc':0.0,'our':0.0,'fee':0.0,'ord':0,'noVed':0})
    nocity=0
    for o in byo.values():
        b=o['mode'] or ''
        t=tot[b]; t['inc']+=o['inc']; t['our']+=o['our']; t['fee']+=o['fee']; t['ord']+=1
        if b=='own' and not o['kn']: t['noVed']+=1
        if not o['city']: nocity+=1; continue
        c=cities[o['city']][b]
        c['inc']+=o['inc']; c['our']+=o['our']; c['fee']+=o['fee']; c['ord']+=1
        if b=='own' and not o['kn']: c['noVed']+=1
    res=lambda x: x['inc']-x['our']-x['fee']
    ws.append([f'Логистика Маркета, {title}. Итог = доход с покупателя − счёт нашего перевозчика − сбор Маркета за логистику.'])
    ws.append(['Там, где везёт Маркет, дохода у нас нет: покупатель платит площадке, а сбор за логистику она удерживает с нас.'])
    ws.append([f'Заказов без города: {nocity}. Строки ИТОГО их учитывают, разбивка по городам — нет.'])
    ws.append([])
    r0=ws.max_row+1
    ws.append(HDR)
    for c in range(1,len(HDR)+1):
        cell=ws.cell(row=r0,column=c); cell.fill=HFILL; cell.font=HFONT
        cell.alignment=Alignment(wrap_text=True,vertical='center',horizontal='center' if c>2 else 'left')
    allm={'inc':sum(t['inc'] for t in tot.values()),'our':sum(t['our'] for t in tot.values()),
          'fee':sum(t['fee'] for t in tot.values()),'ord':sum(t['ord'] for t in tot.values()),
          'noVed':tot['own']['noVed']}
    def row(name,mode,x,bold=False):
        r=res(x)
        ws.append([name,mode,x['ord'],round(x['inc']),round(x['our']),round(x['fee']),round(r),
                   round(r/x['ord']) if x['ord'] else 0, x['noVed'] if mode!='везёт Маркет' else None])
        i=ws.max_row
        for c in range(1,len(HDR)+1):
            cell=ws.cell(row=i,column=c); cell.border=BRD
            if bold: cell.font=TFONT
            if c in (4,5,6,7,8): cell.number_format='# ##0'
        cell=ws.cell(row=i,column=7)
        cell.font=Font(bold=bold,color='B00020' if r<0 else '0B6B3A')
        return i
    row('ИТОГО','все',allm,True)
    for b in ('mk','own',''):
        if tot[b]['ord']: row('  из них','  '+MODE[b],tot[b],True)
    rows=[]
    for city,byb in cities.items():
        agg={'inc':sum(x['inc'] for x in byb.values()),'our':sum(x['our'] for x in byb.values()),
             'fee':sum(x['fee'] for x in byb.values()),'ord':sum(x['ord'] for x in byb.values()),
             'noVed':byb['own']['noVed'] if 'own' in byb else 0}
        rows.append((res(agg),city,byb,agg))
    rows.sort()
    for _,city,byb,agg in rows:
        modes=[b for b in ('mk','own','') if b in byb and byb[b]['ord']]
        if len(modes)==1:
            row(city,MODE[modes[0]],agg)
        else:
            row(city,'все режимы',agg,True)
            for b in modes: row('  '+city,'  '+MODE[b],byb[b])
    for i,w in enumerate([26,18,10,18,16,15,17,12,16],start=1):
        ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes=ws.cell(row=r0+1,column=1)
    ws.auto_filter.ref=f"A{r0}:{get_column_letter(len(HDR))}{ws.max_row}"
    ws.cell(row=1,column=1).font=Font(bold=True,size=11)
    return len(rows), nocity

RU=['январь','февраль','март','апрель','май','июнь','июль','август','сентябрь','октябрь','ноябрь','декабрь']
args=sys.argv[1:] or ['2026-07','2026-08']
wb=Workbook(); wb.remove(wb.active)
info=[]
for ym in args:
    y,mo=(int(x) for x in ym.split('-'))
    title=RU[mo-1].capitalize()+' '+str(y)
    frm=f'{y:04d}-{mo:02d}-01'; to=f'{y:04d}-{mo:02d}-{calendar.monthrange(y,mo)[1]:02d}'
    n,nc=sheet(wb,title,frm,to); info.append((title,n,nc))
out='exports/logistika_goroda_'+'_'.join(args)+'.xlsx'
wb.save(out)
for t,n,nc in info: print(f'{t}: городов {n}, заказов без города {nc}')
print('файл:', out)
