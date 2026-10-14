# ФЕНИКС iter2: дамп всех листов книги в текст + сравнение двух книг по ячейкам
# python3 -I -B dump.py <a.xlsx> [<b.xlsx>]
import sys
from openpyxl import load_workbook
def cells(p):
    wb = load_workbook(p); out = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None: out[(ws.title, c.coordinate)] = c.value
    return wb.sheetnames, out
na, a = cells(sys.argv[1])
if len(sys.argv) == 2:
    print(na)
    cur = None
    for (s, k), v in a.items():
        if s != cur: print('=====', s); cur = s
        print(f'{k}\t{v}')
else:
    nb, b = cells(sys.argv[2])
    print('листы', na == nb, na, nb)
    diff = [(k, a.get(k), b.get(k)) for k in sorted(set(a) | set(b)) if a.get(k) != b.get(k)]
    print('различающихся ячеек', len(diff))
    for d in diff[:40]: print(d)
