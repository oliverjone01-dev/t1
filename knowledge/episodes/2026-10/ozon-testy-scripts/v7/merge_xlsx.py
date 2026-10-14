# v7: один файл для Ивана (14.10 «мне не нужно два файла»): книга «Рост показов» + листы «Цена вниз» и «Клик» из книги варианта 1.
# python3 -I merge_xlsx.py <рост.xlsx> <вариант1.xlsx> <out.xlsx>
import sys, copy
from openpyxl import load_workbook
A, Bf, OUT = sys.argv[1:4]
wa = load_workbook(A); wb = load_workbook(Bf)
for src, dst in (('Цена вниз', 'Цена вниз'), ('Клик', 'Все кампании')):
    s = wb[src]; d = wa.create_sheet(dst, index=3 if dst == 'Цена вниз' else 4)
    for row in s.iter_rows():
        for c in row:
            n = d.cell(row=c.row, column=c.column, value=c.value)
            if c.has_style:
                n.font = copy.copy(c.font); n.fill = copy.copy(c.fill); n.alignment = copy.copy(c.alignment); n.border = copy.copy(c.border)
    for k, dim in s.column_dimensions.items(): d.column_dimensions[k].width = dim.width
    d.freeze_panes = s.freeze_panes
# ---------- «Все кампании» сводим с «Карточками» (ФЕНИКС iter2 Б5): главный товар 10 карточек - повышаем; >2 500 ₽/нед не на главном - минимум; остальное - оставляем
import os as _os
sys.path.insert(0, _os.path.dirname(_os.path.abspath(__file__))); sys.path.insert(0, _os.path.dirname(_os.path.abspath(__file__)) + '/../v5')
import cards7 as K, plan7 as P, forecast as F
from openpyxl.styles import PatternFill as _PF
LVL = 20000
MAINS = {c['main']: c for c in K.cards[:10]}
camp_main = {}
for x in F.L('ads_sku_daily.ndjson'):
    if '2026-10-08' <= x['d'] <= F.B and x['sku'] in MAINS and x['sp'] > 0: camp_main.setdefault(x['cid'], set()).add(x['sku'])
d0 = wa['Все кампании']
cut = None
for row in d0.iter_rows(min_row=2):
    v = row[0].value
    if isinstance(v, str) and v.startswith('Волна 2'): cut = row[0].row; break
if cut: d0.delete_rows(cut - 1, d0.max_row)
n_min = 0
for row in d0.iter_rows(min_row=2):
    cid = str(row[1].value or '')
    if not cid.isdigit(): continue
    kind, t = K.campaign_action(cid, LVL)
    if kind == 'main':
        act, nw, why, col = ('главный товар: без изменений', f'{t:,.0f}'.replace(',', ' '), 'уже не ниже 20 000', 'C6EFCE') if t >= LVL else \
                            ('главный товар: клик до 20 000 с 15.10', '20 000', 'лист «Карточки»', 'C6EFCE')
    elif kind == 'min':
        n_min += 1
        extra = ' Часть расхода шла на главный ' + MAINS[next(iter(camp_main[cid]))]['main_art'] + ' - он получает свои 20 000.' if cid in camp_main else ''
        act, nw, why, col = 'на минимум 15.10', '2 000 (8-12 ₽ за клик)', 'не главный товар, расход больше минимума, заказов мало.' + extra, 'FFC7CE'
    else:
        act, nw, why, col = 'оставляем', f'{t:,.0f}'.replace(',', ' ') if t else 'как есть', 'до 2 500 ₽ в неделю', None
    row[0].value = act; row[9].value = nw; row[10].value = why
    for c in row[:11]: c.fill = _PF('solid', fgColor=col) if col else _PF(fill_type=None)
print('на минимум', n_min)
# строка в «Все кампании»: решения по главным товарам и минимальному клику - на листах «Карточки» и «Минимальный клик»
d = wa['Все кампании']; d.insert_rows(1)
d.cell(1, 1, f'Решения по всем идущим кампаниям: зелёные - главные товары 10 карточек (клик до 20 000); красные - {n_min} кампаний на минимум 15.10; белые - оставляем. Новый минимальный клик 8 ₽ - лист «Минимальный клик».')
e = wa['Этапы']; e.insert_rows(2, 3)
for i, r in enumerate((['15.10 чт, до 10:00', 'Цена вниз, часть 1', '16 артикулов (лист «Цена вниз», дата 15.10)', 'кабинет', 'цена на сайте ниже на 5-10%'],
                       ['22.10 чт', 'Цена вниз, часть 2', '10 артикулов (лист «Цена вниз», дата 22.10)', 'кабинет', 'если тест 4 продлят - 29.10'],
                       ['каждый день', 'Стоп-правило цены', 'артикулы со сниженной ценой', 'Катя', 'выпал из «Эластичного бустинга» - вернуть цену'])):
    for j, v in enumerate(r): e.cell(2 + i, 1 + j, v)
# ---------- единые цвета (Иван 14.10: «выдели чётко разными цветами заголовки, смешано и непонятно»)
from openpyxl.styles import PatternFill, Font
fill = lambda c: PatternFill('solid', fgColor=c)
HEAD, SECT = fill('1F4E78'), fill('D9D9D9')
GREEN, RED, YEL, PUR, BLUE = fill('C6EFCE'), fill('FFC7CE'), fill('FFEB9C'), fill('E4DFEC'), fill('BDD7EE')
old = lambda c: (c.fill.fgColor.rgb or '')[-6:] if c.fill and c.fill.fill_type == 'solid' else ''
LEG = {
    'Цель': 'Цвета: тёмно-синий - заголовок таблицы; серый - раздел.',
    'Этапы': 'Цвета: тёмно-синий - заголовок; фиолетовый - цена; зелёный - реклама вверх; красный - реклама на минимум; жёлтый - проверка и решения.',
    'Карточки': 'Цвета: тёмно-синий - заголовок; зелёный - карточки первой волны (с 15.10), белые - следующие.',
    'Цена вниз': 'Цвета: тёмно-синий - заголовок; фиолетовый - снижаем 15.10; голубой - снижаем 22.10.',
    'Все кампании': 'Цвета: тёмно-синий - заголовок; зелёный - главный товар, клик вверх; красный - на минимум 15.10; белый - оставляем как есть.',
    'Минимальный клик': 'Цвета: тёмно-синий - заголовок; зелёный - включаем 15.10 (первая волна); белые - следующая волна, только если первая прошла порог.',
    'Прогноз': 'Цвета: тёмно-синий - заголовок; серый - раздел.',
    'Пути к 20 млн': 'Цвета: тёмно-синий - заголовок таблицы; серый - раздел.',
}
for ws in wa.worksheets:
    for row in ws.iter_rows():
        o = [old(c) for c in row]
        vals = [c.value for c in row if c.value not in (None, '')]
        if 'DDEBF7' in o:                                        # заголовок таблицы
            for c in row: c.fill = HEAD; c.font = Font(bold=True, color='FFFFFF')
        for c, oc in zip(row, o):
            if oc == 'E2EFDA': c.fill = PUR if ws.title == 'Цена вниз' else GREEN
            elif oc == 'FFF2CC': c.fill = BLUE if ws.title == 'Цена вниз' else YEL
            elif oc == 'FCE4D6': c.fill = RED
    # разделы: строка с одним текстом, жирным
    for row in ws.iter_rows():
        vals = [c for c in row if c.value not in (None, '')]
        if len(vals) == 1 and vals[0].font and vals[0].font.bold and old(vals[0]) != '1F4E78':
            for c in row[:6]: c.fill = SECT
    if ws.title == 'Этапы':
        for row in ws.iter_rows(min_row=2):
            t = str(row[1].value or '')
            f_ = PUR if 'Цена' in t or 'Стоп-правило' in t else GREEN if ('клик' in t.lower() or 'Главным' in t) and 'минимум' not in t and 'Итог' not in t else RED if 'минимум' in t else YEL if t and t != 'Распродажа' else None
            if f_:
                for c in row: c.fill = f_
    ws.insert_rows(1); ws.cell(1, 1, LEG.get(ws.title, '')); ws.cell(1, 1).font = Font(italic=True, bold=True, color='1F4E78')
wa.save(OUT); print(wa.sheetnames)
